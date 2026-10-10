# tests/test_p2_regression/test_v502_docx_image_extract.py
"""
v5.0.2 回归：import_from_docx 应提取 DOCX 内嵌图片。

背景：
  原 import_from_docx 只读段落文本，忽略内嵌图片。
  从 Word 导入后图片全丢。

修复：
  新增 _extract_para_images，支持：
    - DrawingML（现代 Word）：<w:drawing>/<a:blip r:embed="rIdX"/>
    - VML（旧版 Word）：<w:pict>/<v:imagedata r:id="rIdX"/>
  图片保存到 <docx同名>_media/ 目录，Markdown 使用相对路径引用。

测试策略：
  1. 用 python-docx 造一个含内嵌图片的真实 docx
  2. 调 import_from_docx
  3. 断言 _media/ 目录被创建、图片被提取、Markdown 含 ![](path)
"""
import pytest


pytestmark = pytest.mark.docx


# ---------- 依赖检测 ----------
try:
    from docx import Document
    from docx.shared import Inches
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False


@pytest.fixture
def sample_image(tmp_path):
    """造一张小图片。"""
    if not HAS_PIL:
        pytest.skip("Pillow 未安装")
    img_path = tmp_path / "src.png"
    img = Image.new("RGB", (100, 50), color="red")
    img.save(str(img_path))
    return img_path


@pytest.fixture
def docx_with_image(tmp_path, sample_image):
    """造一个含内嵌图片的 docx。"""
    if not HAS_DOCX:
        pytest.skip("python-docx 未安装")

    docx_path = tmp_path / "with_image.docx"
    doc = Document()
    doc.add_heading("测试标题", level=1)
    doc.add_paragraph("图片前的一段文字。")
    doc.add_picture(str(sample_image), width=Inches(1.5))
    doc.add_paragraph("图片后的一段文字。")
    doc.save(str(docx_path))
    return docx_path


def test_import_from_docx_creates_media_dir(docx_with_image, tmp_path):
    """v5.0.2：import_from_docx 应创建 <docx同名>_media/ 目录。"""
    from beeoffice.plugins.bee_export import import_from_docx

    import_from_docx(str(docx_with_image))

    media_dir = tmp_path / "with_image_media"
    assert media_dir.exists(), (
        "v5.0.2：未创建 _media/ 目录。"
        "import_from_docx 的图片提取功能可能被回退了"
    )
    assert media_dir.is_dir()


def test_import_from_docx_extracts_image(docx_with_image, tmp_path):
    """v5.0.2：_media/ 目录应含至少 1 张提取的图片。"""
    from beeoffice.plugins.bee_export import import_from_docx

    import_from_docx(str(docx_with_image))

    media_dir = tmp_path / "with_image_media"
    images = sorted(media_dir.glob("image_*"))
    assert len(images) >= 1, (
        f"v5.0.2：_media/ 目录里没有图片。目录内容: "
        f"{list(media_dir.iterdir()) if media_dir.exists() else '目录不存在'}"
    )
    # 图片应有非零大小
    assert images[0].stat().st_size > 0, (
        f"提取的图片 {images[0].name} 是空文件"
    )


def test_import_from_docx_markdown_contains_image_ref(docx_with_image):
    """v5.0.2：导入后的 Markdown 应含 ![](path) 图片语法。"""
    from beeoffice.plugins.bee_export import import_from_docx

    md = import_from_docx(str(docx_with_image))

    assert "![" in md, (
        f"v5.0.2：Markdown 未含图片语法。"
        f"实际内容前 300 字符: {md[:300]!r}"
    )
    # 相对路径应包含 _media
    assert "_media" in md, (
        f"v5.0.2：Markdown 里的图片引用未使用 _media/ 相对路径。"
        f"实际内容: {md[:300]!r}"
    )


def test_import_from_docx_preserves_text(docx_with_image):
    """v5.0.2：图片提取不应影响文本提取。"""
    from beeoffice.plugins.bee_export import import_from_docx

    md = import_from_docx(str(docx_with_image))

    assert "测试标题" in md
    assert "图片前的一段文字" in md
    assert "图片后的一段文字" in md


def test_import_from_docx_no_image_no_media_dir(tmp_path):
    """v5.0.2：无图片的 docx 不应创建含图片的 _media/ 目录。"""
    if not HAS_DOCX:
        pytest.skip("python-docx 未安装")

    from beeoffice.plugins.bee_export import import_from_docx

    docx_path = tmp_path / "no_image.docx"
    doc = Document()
    doc.add_heading("纯文本", level=1)
    doc.add_paragraph("没有任何图片。")
    doc.save(str(docx_path))

    import_from_docx(str(docx_path))

    media_dir = tmp_path / "no_image_media"
    # _media 目录要么不存在，要么存在但没有 image_* 文件
    if media_dir.exists():
        assert list(media_dir.glob("image_*")) == [], (
            "无图片的 docx 不应提取出任何图片"
        )