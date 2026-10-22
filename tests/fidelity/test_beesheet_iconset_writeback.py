# tests/fidelity/test_beesheet_iconset_writeback.py
"""
蜜算条件格式「图标集」写回 xlsx 回归（2026-10-22，细节 4.2）。

## 修的是什么
GUI 的「条件格式 → 图标集」只在视图模型里记录（`_conditional_formats[col]`），
而导出链 `_apply_conditional` 原先只处理 data_bars/color_scale，
**图标集被 continue 跳过** → 用户能看到 🔴🟡🟢，保存 xlsx 后规则消失。

现在写回 Excel 原生 `3TrafficLights1` 图标集（数值阈值 = 视图的 min/mid/max），
保存后用 openpyxl 能读回 iconSet 规则；用 Excel/WPS 打开可见原生图标。
"""
from openpyxl import Workbook, load_workbook

from beeoffice.plugins.beesheet import BeeSheetView
from beeoffice.plugins.beesheet_style_export import _apply_conditional


def _icon_rules(ws):
    rules = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            if getattr(rule, "type", None) == "iconSet":
                rules.append(rule)
    return rules


def test_apply_conditional_writes_icon_set_rule():
    wb = Workbook()
    ws = wb.active
    applied = _apply_conditional(
        ws, {0: {"type": "icon_set", "min": 0, "max": 9}}, 3
    )
    assert applied == 1
    rules = _icon_rules(ws)
    assert len(rules) == 1
    assert rules[0].iconSet.iconSet == "3TrafficLights1"


def test_icon_set_rule_is_in_saved_xlsx(tmp_path):
    view = BeeSheetView()
    view.model.set_full_dataframe(
        __import__("pandas").DataFrame({"分数": [1, 5, 9]})
    )
    view._add_icon_set(0)
    assert 0 in view._conditional_formats

    path = tmp_path / "iconset.xlsx"
    assert view._write_xlsx(str(path)) is True

    wb = load_workbook(path)
    ws = wb.active
    rules = _icon_rules(ws)
    assert rules, "保存后的 xlsx 里没有 iconSet 规则（导出链又跳过了？）"
    assert any(
        str(getattr(rule.iconSet, "iconSet", "")) == "3TrafficLights1"
        for rule in rules
    )


def test_data_bars_and_color_scale_still_written(tmp_path):
    view = BeeSheetView()
    view.model.set_full_dataframe(
        __import__("pandas").DataFrame({"值": [1, 2, 3]})
    )
    view._add_data_bars(0)
    view._add_color_scale(0)          # 同列覆盖为色阶，只验证能写
    path = tmp_path / "cf.xlsx"
    assert view._write_xlsx(str(path)) is True
    wb = load_workbook(path)
    kinds = {
        rule.type
        for cf in wb.active.conditional_formatting
        for rule in cf.rules
    }
    assert "colorScale" in kinds or "dataBar" in kinds
