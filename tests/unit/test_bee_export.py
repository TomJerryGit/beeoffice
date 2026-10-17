# tests/test_bee_export.py
"""
bee_export.py 测试

覆盖：
  - export_to_md        : 规范化 Markdown
  - export_to_html      : HTML 导出
  - export_to_docx      : DOCX 导出（依赖 python-docx）
  - import_from_docx    : DOCX 导入
  - generate_toc        : 目录生成
  - compare_documents   : 文档对比
  - batch_convert_docx_to_md : 批量 DOCX → MD
  - batch_convert_md_to_docx : 批量 MD → DOCX

设计：
  - 依赖 beeoffice.plugins.bee_export
  - 每个测试用 tmp_path（pytest 内置）隔离文件系统
  - 缺依赖（如 python-docx）时用 pytest.skip 跳过
"""
import pytest
from pathlib import Path


# ==================== 示例 Markdown ====================
# 注意：三引号内不要出现连续三个双引号
SAMPLE_MD = """# 测试标题

这是第一段正文，包含**粗体**和*斜体*。

## 小节一

- 列表项 1
- 列表项 2
- 列表项 3

## 小节二

| 姓名 | 收入 |
|------|-----:|
| 张三 | 200 |
| 李四 | 300 |

> 这是引用文字。

结尾段落。
"""


# ==================== export_to_md ====================

class TestExportToMd:
    """export_to_md 测试"""

    def test_basic(self):
        """基础导出"""
        from beeoffice.plugins.bee_export import export_to_md
        result = export_to_md("# 标题\n\n正文")
        assert "# 标题" in result
        assert "正文" in result

    def test_none_input(self):
        """None 输入返回空字符串"""
        from beeoffice.plugins.bee_export import export_to_md
        assert export_to_md(None) == ""

    def test_normalize_crlf(self):
        """CRLF 换行应被规范化为 LF"""
        from beeoffice.plugins.bee_export import export_to_md
        result = export_to_md("a\r\nb\r\nc")
        assert "\r" not in result
        assert "a\nb\nc" in result

    def test_trailing_whitespace_removed(self):
        """末尾多余空行应被清理，保留一个 \\n"""
        from beeoffice.plugins.bee_export import export_to_md
        result = export_to_md("内容\n\n\n\n")
        assert result == "内容\n"

    def test_empty_string(self):
        """空字符串输入"""
        from beeoffice.plugins.bee_export import export_to_md
        result = export_to_md("")
        assert isinstance(result, str)

    def test_unicode_content(self):
        """中文 / emoji 保留"""
        from beeoffice.plugins.bee_export import export_to_md
        result = export_to_md("# 标题 🐝\n\n正文")
        assert "🐝" in result
        assert "标题" in result


# ==================== export_to_html ====================

class TestExportToHtml:
    """export_to_html 测试"""

    def test_basic(self, tmp_path):
        """基础导出"""
        from beeoffice.plugins.bee_export import export_to_html
        out = tmp_path / "test.html"
        ok = export_to_html(SAMPLE_MD, str(out))
        assert ok is True
        assert out.exists()

        content = out.read_text(encoding="utf-8")
        assert "测试标题" in content
        assert "<html" in content.lower() or "<!doctype" in content.lower()

    def test_creates_parent_dirs(self, tmp_path):
        """自动创建父目录"""
        from beeoffice.plugins.bee_export import export_to_html
        out = tmp_path / "sub" / "deep" / "test.html"
        ok = export_to_html("# 标题", str(out))
        assert ok is True
        assert out.exists()

    def test_contains_table(self, tmp_path):
        """表格应被渲染"""
        from beeoffice.plugins.bee_export import export_to_html
        out = tmp_path / "test.html"
        export_to_html(SAMPLE_MD, str(out))
        content = out.read_text(encoding="utf-8")
        assert "<table" in content.lower()

    def test_title_extracted(self, tmp_path):
        """从首行 # 提取标题"""
        from beeoffice.plugins.bee_export import export_to_html
        out = tmp_path / "test.html"
        export_to_html("# 我的标题\n\n正文", str(out))
        content = out.read_text(encoding="utf-8")
        assert "我的标题" in content

    def test_contains_code_block(self, tmp_path):
        """代码块应被渲染为 <pre>"""
        from beeoffice.plugins.bee_export import export_to_html
        out = tmp_path / "test.html"
        md = "# 标题\n\n```python\nprint('hello')\n```\n"
        export_to_html(md, str(out))
        content = out.read_text(encoding="utf-8")
        assert "<pre" in content.lower() or "<code" in content.lower()

    def test_contains_blockquote(self, tmp_path):
        """引用应被渲染为 <blockquote>"""
        from beeoffice.plugins.bee_export import export_to_html
        out = tmp_path / "test.html"
        export_to_html("> 引用内容\n", str(out))
        content = out.read_text(encoding="utf-8")
        assert "<blockquote" in content.lower()


# ==================== export_to_docx ====================

class TestExportToDocx:
    """export_to_docx 测试"""

    def test_basic(self, tmp_path):
        """基础导出"""
        from beeoffice.plugins.bee_export import export_to_docx, HAS_DOCX
        if not HAS_DOCX:
            pytest.skip("python-docx 未安装")

        out = tmp_path / "test.docx"
        ok = export_to_docx(SAMPLE_MD, str(out))
        assert ok is True
        assert out.exists()
        assert out.stat().st_size > 0

    def test_empty_input(self, tmp_path):
        """空内容也应生成文档"""
        from beeoffice.plugins.bee_export import export_to_docx, HAS_DOCX
        if not HAS_DOCX:
            pytest.skip("python-docx 未安装")

        out = tmp_path / "empty.docx"
        ok = export_to_docx("", str(out))
        assert ok is True
        assert out.exists()

    def test_content_preserved(self, tmp_path):
        """导出后内容应保留"""
        from beeoffice.plugins.bee_export import export_to_docx, HAS_DOCX
        if not HAS_DOCX:
            pytest.skip("python-docx 未安装")

        from docx import Document
        out = tmp_path / "content.docx"
        export_to_docx("# 我的文档\n\n这是正文。", str(out))

        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs]
        full_text = "\n".join(texts)
        assert "我的文档" in full_text
        assert "这是正文" in full_text

    def test_table_preserved(self, tmp_path):
        """表格应被导出"""
        from beeoffice.plugins.bee_export import export_to_docx, HAS_DOCX
        if not HAS_DOCX:
            pytest.skip("python-docx 未安装")

        from docx import Document
        out = tmp_path / "table.docx"
        md = (
            "# 标题\n\n"
            "| 姓名 | 收入 |\n"
            "|------|-----:|\n"
            "| 张三 | 200 |\n"
            "| 李四 | 300 |\n"
        )
        export_to_docx(md, str(out))

        doc = Document(str(out))
        assert len(doc.tables) >= 1
        table = doc.tables[0]
        header_cells = [c.text for c in table.rows[0].cells]
        assert "姓名" in header_cells or "收入" in header_cells


# ==================== import_from_docx ====================

class TestImportFromDocx:
    """import_from_docx 测试"""

    def test_roundtrip(self, tmp_path):
        """导出后重新导入"""
        from beeoffice.plugins.bee_export import (
            export_to_docx, import_from_docx, HAS_DOCX
        )
        if not HAS_DOCX:
            pytest.skip("python-docx 未安装")

        docx_path = tmp_path / "roundtrip.docx"
        export_to_docx(SAMPLE_MD, str(docx_path))

        md = import_from_docx(str(docx_path))
        assert isinstance(md, str)
        assert "测试标题" in md

    def test_nonexistent_file(self, tmp_path):
        """不存在的文件应报错"""
        from beeoffice.plugins.bee_export import import_from_docx, HAS_DOCX
        if not HAS_DOCX:
            pytest.skip("python-docx 未安装")

        with pytest.raises(Exception):
            import_from_docx(str(tmp_path / "nonexistent.docx"))

    def test_heading_converted(self, tmp_path):
        """标题应转换为 # 语法"""
        from beeoffice.plugins.bee_export import (
            export_to_docx, import_from_docx, HAS_DOCX
        )
        if not HAS_DOCX:
            pytest.skip("python-docx 未安装")

        docx_path = tmp_path / "heading.docx"
        export_to_docx("# 一级标题\n\n## 二级标题\n\n正文", str(docx_path))

        md = import_from_docx(str(docx_path))
        # 至少应有一个 # 开头
        assert "#" in md


# ==================== generate_toc ====================

class TestGenerateToc:
    """generate_toc 测试"""

    def test_basic(self):
        """生成目录"""
        from beeoffice.plugins.bee_export import generate_toc
        toc = generate_toc(SAMPLE_MD)
        assert isinstance(toc, str)
        assert "目录" in toc
        assert "测试标题" in toc
        assert "小节一" in toc
        assert "小节二" in toc

    def test_no_headings(self):
        """没有标题时的处理"""
        from beeoffice.plugins.bee_export import generate_toc
        toc = generate_toc("这是纯文本，没有标题。")
        assert isinstance(toc, str)

    def test_empty_input(self):
        """空输入"""
        from beeoffice.plugins.bee_export import generate_toc
        result = generate_toc("")
        assert isinstance(result, str)

    def test_nested_headings(self):
        """多级标题应正确缩进"""
        from beeoffice.plugins.bee_export import generate_toc
        md = "# H1\n\n## H2\n\n### H3\n"
        toc = generate_toc(md)
        assert "H1" in toc
        assert "H2" in toc
        assert "H3" in toc


# ==================== compare_documents ====================

class TestCompareDocuments:
    """compare_documents 测试"""

    def test_identical(self):
        """相同内容无差异"""
        from beeoffice.plugins.bee_export import compare_documents
        result = compare_documents("相同内容", "相同内容")
        assert result["additions"] == 0
        assert result["deletions"] == 0
        assert result["total_changes"] == 0

    def test_addition(self):
        """新增行"""
        from beeoffice.plugins.bee_export import compare_documents
        result = compare_documents("第 1 行", "第 1 行\n第 2 行")
        assert result["additions"] >= 1
        assert result["total_changes"] >= 1

    def test_deletion(self):
        """删除行"""
        from beeoffice.plugins.bee_export import compare_documents
        result = compare_documents("第 1 行\n第 2 行", "第 1 行")
        assert result["deletions"] >= 1

    def test_diff_text_returned(self):
        """diff 字段应存在"""
        from beeoffice.plugins.bee_export import compare_documents
        result = compare_documents("原文本", "新文本")
        assert "diff" in result
        assert isinstance(result["diff"], str)
        assert len(result["diff"]) > 0

    def test_modified_lines(self):
        """修改行统计"""
        from beeoffice.plugins.bee_export import compare_documents
        result = compare_documents("原始内容", "修改内容")
        # 修改会产生 additions + deletions
        assert result["total_changes"] >= 1


# ==================== 批量转换 ====================

class TestBatchConvert:
    """批量转换测试"""

    def test_docx_to_md(self, tmp_path):
        """批量 DOCX → MD"""
        from beeoffice.plugins.bee_export import (
            export_to_docx, batch_convert_docx_to_md, HAS_DOCX
        )
        if not HAS_DOCX:
            pytest.skip("python-docx 未安装")

        input_dir = tmp_path / "input"
        output_dir = tmp_path / "output"
        input_dir.mkdir()

        export_to_docx("# 文档一\n\n内容一", str(input_dir / "a.docx"))
        export_to_docx("# 文档二\n\n内容二", str(input_dir / "b.docx"))

        results = batch_convert_docx_to_md(str(input_dir), str(output_dir))
        assert len(results) == 2
        assert all(r["status"] == "✅" for r in results)
        assert (output_dir / "a.md").exists()
        assert (output_dir / "b.md").exists()

    def test_md_to_docx(self, tmp_path):
        """批量 MD → DOCX"""
        from beeoffice.plugins.bee_export import (
            batch_convert_md_to_docx, HAS_DOCX
        )
        if not HAS_DOCX:
            pytest.skip("python-docx 未安装")

        input_dir = tmp_path / "input"
        output_dir = tmp_path / "output"
        input_dir.mkdir()

        (input_dir / "a.md").write_text("# 文档一\n\n内容一", encoding="utf-8")
        (input_dir / "b.md").write_text("# 文档二\n\n内容二", encoding="utf-8")

        results = batch_convert_md_to_docx(str(input_dir), str(output_dir))
        assert len(results) == 2
        assert all(r["status"] == "✅" for r in results)
        assert (output_dir / "a.docx").exists()
        assert (output_dir / "b.docx").exists()

    def test_nonexistent_input_dir(self, tmp_path):
        """输入目录不存在时返回空列表"""
        from beeoffice.plugins.bee_export import batch_convert_docx_to_md
        results = batch_convert_docx_to_md(
            str(tmp_path / "nonexistent"),
            str(tmp_path / "output")
        )
        assert results == []

    def test_empty_input_dir(self, tmp_path):
        """输入目录为空时返回空列表"""
        from beeoffice.plugins.bee_export import batch_convert_docx_to_md
        input_dir = tmp_path / "input"
        input_dir.mkdir()
        results = batch_convert_docx_to_md(
            str(input_dir),
            str(tmp_path / "output")
        )
        assert results == []