# tests/fidelity/test_docx_header_footer_5_1.py
"""5.1：页眉页脚编辑（对话框 + 视图字段 + DOCX 导出文字）。"""
from docx import Document
from PySide6.QtWidgets import QLineEdit

from beeoffice.plugins.bee_export import apply_header_footer
from beeoffice.plugins.beewrite_doc_mixin import DocMixin
from beeoffice.plugins.beewrite_page_extras import HeaderFooterDialog


class _Fake:
    def __init__(self):
        self._docx_header_text = ""
        self._docx_footer_text = ""
        self._docx_section_parts = [
            {"header_linked": True, "footer_linked": True,
             "header_xml": "<w:p/>", "footer_xml": "<w:p/>"},
        ]
        self._status = []
        self.contentChanged = type(
            "S", (), {"emit": staticmethod(lambda: None)}
        )()

    def _set_status(self, text, *_a):
        self._status.append(text)


def test_dialog_values(qapp):
    dlg = HeaderFooterDialog("旧页眉", "旧页脚")
    dlg._header_edit.setText("新页眉")
    dlg._footer_edit.setText("新页脚")
    assert dlg.values() == ("新页眉", "新页脚")
    assert isinstance(dlg._header_edit, QLineEdit)


def test_apply_updates_first_section(qapp):
    fake = _Fake()
    DocMixin._apply_header_footer_texts(fake, "H", "F")
    assert fake._docx_header_text == "H"
    assert fake._docx_footer_text == "F"
    entry = fake._docx_section_parts[0]
    assert entry["header_text"] == "H"
    assert entry["footer_text"] == "F"
    assert entry["header_xml"] == ""
    assert entry["footer_linked"] is False
    assert any("页眉/页脚" in s for s in fake._status)


def test_dialog_first_and_odd_even_options(qapp):
    dlg = HeaderFooterDialog()
    dlg._first_check.setChecked(True)
    dlg._first_header_edit.setText("首页眉")
    dlg._odd_even_check.setChecked(True)
    dlg._even_footer_edit.setText("偶数页脚")
    opts = dlg.options()
    assert opts["different_first"] is True
    assert opts["first_header"] == "首页眉"
    assert opts["odd_even"] is True
    assert opts["even_footer"] == "偶数页脚"


def test_docx_export_first_and_even_header_footer(tmp_path):
    doc = Document()
    options = {
        "different_first": True, "odd_even": True,
        "first_header": "首页页眉", "first_footer": "首页页脚",
        "even_header": "偶数页眉", "even_footer": "偶数页脚",
    }
    assert apply_header_footer(doc, header_footer_options=options)
    out = tmp_path / "hf_first_even.docx"
    doc.save(str(out))
    reopened = Document(str(out))
    section = reopened.sections[0]
    assert section.different_first_page_header_footer is True
    assert reopened.settings.odd_and_even_pages_header_footer is True
    assert section.first_page_header.paragraphs[0].text == "首页页眉"
    assert section.first_page_footer.paragraphs[0].text == "首页页脚"
    assert section.even_page_header.paragraphs[0].text == "偶数页眉"
    assert section.even_page_footer.paragraphs[0].text == "偶数页脚"


def test_docx_export_writes_header_footer_text(tmp_path):
    doc = Document()
    assert apply_header_footer(doc, header_text="页眉文字", footer_text="页脚文字")
    out = tmp_path / "hf.docx"
    doc.save(str(out))
    reopened = Document(str(out))
    section = reopened.sections[0]
    assert section.header.paragraphs[0].text == "页眉文字"
    assert section.footer.paragraphs[0].text == "页脚文字"
