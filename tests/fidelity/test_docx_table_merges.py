# tests/fidelity/test_docx_table_merges.py
"""
保真度阶段 A 最后一项：DOCX **表格合并单元格**往返。

问题
----
Markdown 表格语法**表达不了合并**，因此导入→导出后合并必然丢失。
实测基线：`merged_cells 1 → 0`，并连带 `table_cells 8→9`、
`paragraphs 14→15`、`text_chars 102→104`（一个合并格被拆回两格）。

修法：与页面设置/页眉页脚同样的思路——**绕开 Markdown** 传递。
导入时 `_read_table_merges()` 记下"第几张表 + 网格行列 + 跨度 + 合并后文字"，
导出时 `apply_table_merges()` 用 python-docx 的 `merge()` 还原。

**注意 merge() 会把两个单元格的文字拼接起来**，所以合并后必须用记录里的原文
覆盖，否则会出现"合计合计"。
"""
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document                                  # noqa: E402

from beeoffice.plugins.bee_export import (                 # noqa: E402
    _read_table_merges,
    apply_table_merges,
    export_to_docx,
    import_from_docx,
    read_docx_extras,
)


def _make_merged_docx(path: Path, *, vertical=False) -> Path:
    """构造含合并单元格的 DOCX。"""
    from docx import Document as D

    doc = D()
    doc.add_paragraph("表格测试")
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for ci, text in enumerate(["指标", "本期", "上期"]):
        table.cell(0, ci).text = text
    table.cell(1, 0).text = "营收"
    table.cell(1, 1).text = "1,234"
    table.cell(1, 2).text = "1,000"

    if vertical:
        # 纵向合并第 0 列的第 1~2 行（先把两行的文字写好）
        table.cell(2, 0).text = "整季"
        table.cell(2, 1).text = "2,234"
        table.cell(2, 2).text = "2,000"
        table.cell(1, 0).merge(table.cell(2, 0))
    else:
        # 横向合并最后一行的前两列
        table.cell(2, 0).text = "合计"      # 忘了这一步就会出现"空合并"，我第一版就漏了
        table.cell(2, 2).text = "3,234"
        table.cell(2, 0).merge(table.cell(2, 1))
    doc.save(str(path))
    return path


def _roundtrip(src: Path, tmp_path: Path) -> Path:
    out = tmp_path / "out.docx"
    extras = read_docx_extras(src)
    assert export_to_docx(
        import_from_docx(str(src)), str(out),
        page_settings=extras, base_dir=str(tmp_path),
        header_text=extras.get("header_text"),
        footer_text=extras.get("footer_text"),
        table_merges=extras.get("table_merges"),
    ) is True
    return out


# ==================== 1. 读取合并信息 ====================

def test_reads_horizontal_merge(tmp_path):
    src = _make_merged_docx(tmp_path / "h.docx")

    merges = _read_table_merges(Document(str(src)))

    assert len(merges) == 1, merges
    record = merges[0]
    assert record["table"] == 0
    assert record["row"] == 2
    assert record["col"] == 0
    assert record["colspan"] == 2
    assert record["rowspan"] == 1
    assert record["text"] == "合计" or record["text"] == "营收", record


def test_merge_record_carries_text(tmp_path):
    """合并后的文字要记下来（否则还原时会出现重复文字）。"""
    from docx import Document as D

    src = tmp_path / "t.docx"
    doc = D()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "甲"
    table.cell(0, 1).text = "乙"
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.text = "合并后"
    doc.save(str(src))

    merges = _read_table_merges(Document(str(src)))

    assert merges[0]["text"] == "合并后", merges


def test_reads_vertical_merge_rowspan(tmp_path):
    """
    纵向合并（w:vMerge）必须累计出 rowspan。

    这段逻辑是"盲写"的（夹具里没有纵向合并），所以单独验证。
    """
    src = _make_merged_docx(tmp_path / "v.docx", vertical=True)

    merges = _read_table_merges(Document(str(src)))

    assert len(merges) == 1, merges
    record = merges[0]
    assert record["row"] == 1, record
    assert record["col"] == 0, record
    assert record["rowspan"] == 2, (
        f"纵向合并应跨 2 行，实际 {record['rowspan']}：{record}"
    )
    assert record["colspan"] == 1, record


def test_no_merge_returns_empty(tmp_path):
    from docx import Document as D

    src = tmp_path / "plain.docx"
    doc = D()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    doc.save(str(src))

    assert _read_table_merges(Document(str(src))) == []


def test_read_docx_extras_includes_table_merges(tmp_path):
    src = _make_merged_docx(tmp_path / "h.docx")

    extras = read_docx_extras(src)

    assert extras.get("table_merges"), extras.keys()


# ==================== 2. 应用合并 ====================

def test_apply_restores_merge(tmp_path):
    from tools.ooxml_probe import probe

    src = _make_merged_docx(tmp_path / "h.docx")
    out = _roundtrip(src, tmp_path)

    before = probe(src)["features"]
    after = probe(out)["features"]
    assert before["merged_cells"] == 1, before
    assert after["merged_cells"] == 1, (
        f"合并没有还原：{before['merged_cells']} → {after['merged_cells']}"
    )


def test_apply_does_not_duplicate_text(tmp_path):
    """
    `merge()` 会拼接文字；必须用记录里的原文覆盖，否则合并格内部会变成
    "合计合计"。

    ⚠️ 不要用 `row.cells` 拼接来判断：python-docx 对合并格会**按网格列
    重复返回同一个单元格**，于是 `['合计', '合计', '3,234']` 看起来像重复，
    实际是同一格出现两次（我第一版就是这么误判的）。要看单元格**自身**的文字。
    """
    src = _make_merged_docx(tmp_path / "h.docx")
    out = _roundtrip(src, tmp_path)

    table = Document(str(out)).tables[0]
    anchor = table.cell(2, 0)

    assert anchor.text.strip() == "合计", (
        f"合并格文字被重复了: {anchor.text!r}"
    )
    # 合并格覆盖的两个网格位置应当指向同一个单元格
    assert table.cell(2, 1).text.strip() == "合计"


def test_apply_skips_out_of_range_records(tmp_path):
    """越界的合并记录应被跳过，而不是抛异常中断导出。"""
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    doc.tables[0].cell(0, 0).text = "A"

    applied = apply_table_merges(doc, [
        {"table": 0, "row": 5, "col": 0, "colspan": 2, "rowspan": 1,
         "text": "越界"},
        {"table": 9, "row": 0, "col": 0, "colspan": 2, "rowspan": 1,
         "text": "表号越界"},
    ])

    assert applied == 0, f"越界记录不应被应用: {applied}"


def test_apply_ignores_single_cell_records(tmp_path):
    """跨度为 1×1 的记录没有意义，应跳过。"""
    doc = Document()
    doc.add_table(rows=2, cols=2)
    assert apply_table_merges(doc, [
        {"table": 0, "row": 0, "col": 0, "colspan": 1, "rowspan": 1,
         "text": "x"},
    ]) == 0


def test_apply_empty_is_noop():
    doc = Document()
    assert apply_table_merges(doc, None) == 0
    assert apply_table_merges(doc, []) == 0


# ==================== 3. 端到端：整表零差异 ====================

@pytest.mark.parametrize("vertical", [False, True])
def test_full_roundtrip_has_no_feature_loss(tmp_path, vertical):
    """
    **核心断言**：整份文档往返后不应有任何特征减少。

    这条覆盖 DOCX 全部五刀（超链接、页面设置、页眉页脚、段落边界、表格合并）。
    """
    from tools.ooxml_probe import compare, probe

    src = _make_merged_docx(tmp_path / "src.docx", vertical=vertical)
    out = _roundtrip(src, tmp_path)

    rows = compare(probe(src), probe(out))
    lost = [r for r in rows if r["status"] == "lost"]
    assert not lost, "\n".join(
        f"  {r['feature']}: {r['before']} → {r['after']}" for r in lost
    )


# ==================== 4. 编辑器与产品接线 ====================

def test_editor_carries_table_merges(beewrite_editor):
    records = [{"table": 0, "row": 2, "col": 0, "colspan": 2,
                "rowspan": 1, "text": "合计"}]

    beewrite_editor.apply_docx_extras({"table_merges": records})

    assert beewrite_editor.get_docx_extras()["table_merges"] == records


def test_editor_defaults_to_no_merges(beewrite_editor):
    assert beewrite_editor.get_docx_extras()["table_merges"] == []


def test_write_docx_restores_merges(beewrite_editor, tmp_path):
    """产品级：编辑器带着合并信息导出后，文件里必须真有合并。"""
    from tools.ooxml_probe import probe

    beewrite_editor.set_content(
        "| 指标 | 本期 | 上期 |\n|---|---|---|\n"
        "| 营收 | 1 | 2 |\n| 合计 | 合计 | 3 |\n"
    )
    beewrite_editor.apply_docx_extras({"table_merges": [
        {"table": 0, "row": 2, "col": 0, "colspan": 2,
         "rowspan": 1, "text": "合计"},
    ]})

    out = tmp_path / "merged.docx"
    assert beewrite_editor._write_docx(str(out)) is True

    assert probe(out)["features"]["merged_cells"] >= 1, (
        "导出的文件里没有合并单元格"
    )
