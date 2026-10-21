# tests/fidelity/test_large_xlsx_fast_read.py
"""
v5.0.17：大 xlsx（>= 8 MB）的快速读值路径（python-calamine）。

## 背景（用户 2026-10-21 放入 20/25/30 万行订单表）
- openpyxl/pandas `read_excel` 实测：20 万行 **165 s**、25 万行 **276 s**、
  30 万行 >10 分钟（未跑完）；
- calamine（Rust）实测：20 万行 33 s（首跑含加载）、25 万行 11.6 s、
  30 万行 13.1 s；
- 策略：**只对大文件**（>=8 MB）启用 calamine；小文件保持 openpyxl 路径，
  避免类型/日期解析细节变化影响既有保真测试；未安装 calamine 时自动回退。

## 本文件锁住
1. 大文件走 calamine（且公式仍由 `_read_formulas` 保留）；
2. 小文件仍走 openpyxl（行为不漂移）；
3. 未安装 calamine / calamine 抛错时自动回退，不阻断打开。
"""
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from beeoffice.plugins import beesheet_workbook as bw


def _make_workbook(path: Path, formula_sheet2: bool = False):
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["名称", "数量"])
    for i in range(3):
        ws.append([f"项{i}", i + 1])
    ws2 = wb.create_sheet("公式")
    ws2.append(["x", "y", "z"])       # 第 1 行是表头（df 坐标约定）
    ws2.append([1, 2, "=A1+B1"])      # C2 才是数据区公式
    if not formula_sheet2:
        ws2["C2"] = "无公式"
    wb.save(path)


@pytest.fixture
def plain_xlsx(tmp_path):
    path = tmp_path / "plain.xlsx"
    _make_workbook(path, formula_sheet2=False)
    return path


def test_large_file_uses_calamine_engine(plain_xlsx, monkeypatch):
    calls = []
    real_read_excel = pd.read_excel

    def _spy(*args, **kwargs):
        calls.append(kwargs.get("engine"))
        return real_read_excel(*args, **kwargs)

    monkeypatch.setattr(bw, "_CALAMINE_MIN_BYTES", 0)
    monkeypatch.setattr(bw, "_calamine_available", lambda: True)
    monkeypatch.setattr(bw.pd, "read_excel", _spy)
    wb = bw.read_workbook(plain_xlsx)
    assert "calamine" in calls
    assert wb.sheets[0].df.iloc[0, 0] == "项0"
    assert wb.sheets[0].original_shape == (3, 2)


def test_small_file_keeps_openpyxl_engine(plain_xlsx, monkeypatch):
    calls = []
    real_read_excel = pd.read_excel

    def _spy(*args, **kwargs):
        calls.append(kwargs.get("engine"))
        return real_read_excel(*args, **kwargs)

    monkeypatch.setattr(bw, "_calamine_available", lambda: True)
    monkeypatch.setattr(bw.pd, "read_excel", _spy)
    assert plain_xlsx.stat().st_size < bw._CALAMINE_MIN_BYTES
    bw.read_workbook(plain_xlsx)
    assert calls == ["openpyxl"]


def test_missing_calamine_falls_back(plain_xlsx, monkeypatch):
    calls = []
    real_read_excel = pd.read_excel

    def _spy(*args, **kwargs):
        calls.append(kwargs.get("engine"))
        return real_read_excel(*args, **kwargs)

    monkeypatch.setattr(bw, "_CALAMINE_MIN_BYTES", 0)
    monkeypatch.setattr(bw, "_calamine_available", lambda: False)
    monkeypatch.setattr(bw.pd, "read_excel", _spy)
    wb = bw.read_workbook(plain_xlsx)
    assert "calamine" not in calls
    assert wb.sheets[0].df.iloc[0, 0] == "项0"


def test_calamine_error_falls_back(plain_xlsx, monkeypatch):
    real_read_excel = pd.read_excel

    def _spy(*args, **kwargs):
        if kwargs.get("engine") == "calamine":
            raise ValueError("模拟 calamine 失败")
        return real_read_excel(*args, **kwargs)

    monkeypatch.setattr(bw, "_CALAMINE_MIN_BYTES", 0)
    monkeypatch.setattr(bw, "_calamine_available", lambda: True)
    monkeypatch.setattr(bw.pd, "read_excel", _spy)
    wb = bw.read_workbook(plain_xlsx)
    assert wb.sheets[0].df.iloc[0, 0] == "项0"


def test_row_limited_preview_uses_openpyxl(plain_xlsx, monkeypatch):
    """阶段 1（只读前 500 行）不能走 calamine：calamine 会先解析整份文件。"""
    calls = []
    real_read_excel = pd.read_excel

    def _spy(*args, **kwargs):
        calls.append((kwargs.get("engine"), kwargs.get("nrows")))
        return real_read_excel(*args, **kwargs)

    monkeypatch.setattr(bw, "_CALAMINE_MIN_BYTES", 0)
    monkeypatch.setattr(bw, "_calamine_available", lambda: True)
    monkeypatch.setattr(bw.pd, "read_excel", _spy)
    wb = bw.read_workbook(plain_xlsx, max_rows=bw._CALAMINE_MAX_SKIP_ROWS)
    assert wb.sheets
    assert calls and calls[0][0] == "openpyxl", calls
    assert calls[0][1] == bw._CALAMINE_MAX_SKIP_ROWS


def test_no_formula_sheets_skip_openpyxl_load(plain_xlsx, monkeypatch):
    """全部工作表预检无公式时，不能为了找公式再加载一遍 openpyxl 工作簿。"""
    import beeoffice.plugins.beesheet_workbook as bw_mod

    def _boom(*args, **kwargs):
        raise AssertionError("无公式的工作簿不应再 load_workbook")

    monkeypatch.setattr(bw_mod, "load_workbook", _boom)
    wb = bw_mod.read_workbook(plain_xlsx)
    assert wb.formula_count == 0
    assert wb.sheets[0].df.iloc[0, 0] == "项0"


def test_sparse_large_sheet_uses_calamine(plain_xlsx, monkeypatch):
    """磁盘小但声明行数很大的稀疏表，也应走 calamine（不能只看文件大小）。"""
    calls = []
    real_read_excel = pd.read_excel

    def _spy(*args, **kwargs):
        calls.append(kwargs.get("engine"))
        return real_read_excel(*args, **kwargs)

    monkeypatch.setattr(bw, "_calamine_available", lambda: True)
    monkeypatch.setattr(bw, "_declared_row_count", lambda _p: 200_000)
    monkeypatch.setattr(bw.pd, "read_excel", _spy)
    bw.read_workbook(plain_xlsx)
    assert "calamine" in calls, calls


def test_small_sheet_still_uses_openpyxl(plain_xlsx, monkeypatch):
    calls = []
    real_read_excel = pd.read_excel

    def _spy(*args, **kwargs):
        calls.append(kwargs.get("engine"))
        return real_read_excel(*args, **kwargs)

    monkeypatch.setattr(bw, "_calamine_available", lambda: True)
    monkeypatch.setattr(bw, "_declared_row_count", lambda _p: 10)
    monkeypatch.setattr(bw.pd, "read_excel", _spy)
    bw.read_workbook(plain_xlsx)
    assert calls and calls[0] == "openpyxl", calls


def test_probe_xlsx_shape_for_skipped_large_table(tmp_path):
    """大表被 --max-size-mb 跳过时，报告仍应给出声明维度（不读值）。"""
    from tools.fidelity_report import _probe_xlsx_shape

    path = tmp_path / "shape.xlsx"
    _make_workbook(path, formula_sheet2=True)
    shape = _probe_xlsx_shape(path)
    assert shape == {"rows": 3, "cols": 2, "sheets": 2}


def test_calamine_path_still_keeps_formulas(tmp_path, monkeypatch):
    path = tmp_path / "f.xlsx"
    _make_workbook(path, formula_sheet2=True)
    monkeypatch.setattr(bw, "_CALAMINE_MIN_BYTES", 0)
    monkeypatch.setattr(bw, "_calamine_available", lambda: True)
    wb = bw.read_workbook(path)
    formulas = wb.get("公式")
    assert formulas is not None
    assert any(v == "=A1+B1" for v in formulas.formulas.values())
