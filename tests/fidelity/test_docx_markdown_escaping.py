# tests/fidelity/test_docx_markdown_escaping.py
"""
保真度阶段 A：导入 DOCX 时**必须转义 Markdown/HTML 特殊写法**。

真实文档撞出来的问题
--------------------
`tests/fidelity/corpus/` 里两份真实技术教程（shell 命令 + 输出）暴露了同一根因：
导入时把原文**直接**塞进 Markdown，任何"长得像 Markdown/HTML"的文字都会被
**重新解释**：

| 现象 | 实例 | 后果 |
|---|---|---|
| 行首 `#` 判成标题 | `#将NAME=ens33 修改为NAME=eth0` | `#` 消失、样式变"标题 1"；实测该文档**凭空多出 13 个标题** |
| 行内 `<...>` 当标签 | `flags=4163<UP,BROADCAST,RUNNING,MULTICAST>` | 尖括号内容**整段消失** |
| 行内 `*`/`_` 当强调 | `a*b*c` | 变成 `abc` |
| 行首缩进当代码块 | `    TYPE=Ethernet` | 多段被并成一段，且缩进丢失 |
| 行首 `-`/`+`/`1.` | `- 减号开头但不是列表` | 被当成列表项 |

修法与两个易错点
----------------
- `<` `>` `&` 必须用 **HTML 实体**（`&lt;` 等），**不能**用反斜杠：
  Python-Markdown 的可转义集合不含 `<`/`&`，实测 `\\<UP,...>` 里的 `\\` 原样留下、
  尖括号内容照样被吃掉。
- 行首缩进用 `&#32;`/`&#9;` 实体保留，**不要 strip** —— 技术文档里代码行的缩进
  是有意义的；实体在块级解析时不算空白（不会变代码块），渲染时又变回空格。
- **转义必须区分"行内"与"行级"**：`_escape_inline_text()` 是按 **run** 调用的，
  绝不能在里面 strip —— 一个 run 完全可能只包含空格（真实文档里
  `'原', '  ', '因：'` 就是三个 run）。早期版本在 run 级 strip，把空格 run 抹掉，
  导致"原  因"变成"原因"。行级处理只在整段拼好后做
  （`_escape_block_markers()`）。
"""
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document                                  # noqa: E402

from beeoffice.plugins.bee_export import (                 # noqa: E402
    _escape_block_markers,
    _escape_inline_text,
    export_to_docx,
    import_from_docx,
    read_docx_extras,
)

#: 每一条都是"会被 Markdown/HTML 重新解释"的真实写法
TRAPS = [
    ("行首井号（shell 注释）", "#将NAME=ens33 修改为NAME=eth0"),
    ("行首井号无空格", "#/etc/sysconfig/network-scripts/ifcfg-ens33"),
    ("行首大于号", "> 引用符号开头的 shell 提示"),
    ("行首减号", "- 减号开头但不是列表"),
    ("行首加号", "+ 加号开头但不是列表"),
    ("行首数字点", "1. 数字点开头但不是列表"),
    ("尖括号（shell 输出）",
     "eth0: flags=4163<UP,BROADCAST,RUNNING,MULTICAST>  mtu 1500"),
    ("尖括号 link", "inet6 fe80::3061  prefixlen 64  scopeid 0x20<link>"),
    ("和号", "含 & 符号的文本"),
    ("行内星号", "含 a*b*c 星号的文本"),
    ("行内下划线", "含 a_b_c 下划线的文本"),
    ("行内反引号", "含 `code` 反引号的文本"),
    ("方括号", "[root@centos ~]# vi /etc/x"),
    ("普通中文", "这是一段普通中文。"),
]


def _roundtrip_paras(texts, tmp_path):
    """把若干段落写进 DOCX → 导入 → 导出，返回（源, 往返后）的段落文字。"""
    src = tmp_path / "src.docx"
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    doc.save(str(src))

    out = tmp_path / "out.docx"
    extras = read_docx_extras(src)
    assert export_to_docx(
        import_from_docx(str(src)), str(out),
        page_settings=extras, base_dir=str(tmp_path),
        header_text=extras.get("header_text"),
        footer_text=extras.get("footer_text"),
        table_merges=extras.get("table_merges"),
        sections=extras.get("sections"),
    ) is True

    src_paras = [p for p in Document(str(src)).paragraphs if p.text.strip()]
    out_paras = [p for p in Document(str(out)).paragraphs if p.text.strip()]
    return src_paras, out_paras


# ==================== 1. 逐条陷阱：内容与样式都要保住 ====================

@pytest.mark.parametrize("label,text", TRAPS, ids=[t[0] for t in TRAPS])
def test_trap_text_survives_roundtrip(label, text, tmp_path):
    src, out = _roundtrip_paras([text], tmp_path)

    assert len(out) == 1, f"{label}: 段落数变了 {len(src)} → {len(out)}"
    assert out[0].text == text, (
        f"{label}: 文字被改写\n  源: {text!r}\n  后: {out[0].text!r}"
    )
    assert out[0].style.name == "Normal", (
        f"{label}: 样式被改成了 {out[0].style.name}（说明被当成标题/列表了）"
    )


def test_shell_comments_do_not_become_headings(tmp_path):
    """
    **真实文档的核心症状**：多个 `#` 开头的注释行不能变成一堆标题。

    实测 `修改Centos7网卡ens33为eth0.docx` 原先凭空多出 13 个"标题 1"。
    """
    texts = [
        "#将NAME=ens33 修改为NAME=eth0",
        "#将DEVICE =ens33 修改为DEVICE =eth0",
        "#将ONBOOT=no 修改为ONBOOT=yes",
    ]
    src, out = _roundtrip_paras(texts, tmp_path)

    styles = {p.style.name for p in out}
    assert styles == {"Normal"}, f"注释行被判成了其它样式: {styles}"
    assert [p.text for p in out] == texts


def test_angle_bracket_content_is_not_eaten(tmp_path):
    """`<...>` 必须原样保留（不能被当成 HTML 标签吃掉）。"""
    text = "eth0: flags=4163<UP,BROADCAST,RUNNING,MULTICAST>  mtu 1500"
    _, out = _roundtrip_paras([text], tmp_path)

    assert "<UP,BROADCAST,RUNNING,MULTICAST>" in out[0].text, out[0].text


def test_leading_indentation_survives_without_code_block(tmp_path):
    """
    行首缩进要保住**且不触发代码块**：用 `&#32;`/`&#9;` 实体表示。

    技术文档里代码行的缩进是有意义的（YAML/Python 缩进错一位意思就变了）。
    """
    texts = ["    TYPE=Ethernet", "\tNAME=eth0", "  [root@centos ~]# ls"]
    src, out = _roundtrip_paras(texts, tmp_path)

    assert [p.text for p in out] == texts, (
        f"缩进没保住: {[p.text for p in out]!r}"
    )
    assert len(out) == 3, f"多段被并成一段（代码块）: {len(out)}"


# ==================== 2. 不能过度转义（我曾在这里踩坑） ====================

def test_whitespace_only_run_is_not_destroyed(tmp_path):
    """
    **回归**：段落由多个 run 组成，其中一个是纯空格 run。

    真实文档里 `'原', '  ', '因：'` 就是三个 run。早期版本在 run 级做
    `lstrip`，把空格 run 抹成空串 → "原  因"变成"原因"。
    """
    src = tmp_path / "src.docx"
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("原")
    para.add_run("  ")          # 纯空格 run
    para.add_run("因：正文")
    doc.save(str(src))

    out = tmp_path / "out.docx"
    extras = read_docx_extras(src)
    export_to_docx(
        import_from_docx(str(src)), str(out),
        page_settings=extras, base_dir=str(tmp_path),
        header_text=extras.get("header_text"),
        footer_text=extras.get("footer_text"),
        table_merges=extras.get("table_merges"),
        sections=extras.get("sections"),
    )

    texts = [p.text for p in Document(str(out)).paragraphs if p.text.strip()]
    assert texts == ["原  因：正文"], (
        f"纯空格 run 被吃掉了: {texts!r}"
    )


def test_normal_paragraph_is_not_over_escaped(tmp_path):
    """普通文字不应被加上一堆反斜杠/实体。"""
    text = "这是一段普通中文，含逗号与句号。"
    _, out = _roundtrip_paras([text], tmp_path)

    assert out[0].text == text
    assert "\\" not in out[0].text and "&" not in out[0].text


# ==================== 3. 转义函数本身 ====================

def test_inline_escape_leaves_spaces_alone():
    """行内转义**不得**动空白（它是按 run 调用的）。"""
    assert _escape_inline_text("  ") == "  "
    assert _escape_inline_text("a b") == "a b"


def test_inline_escape_uses_entities_for_angle_brackets():
    assert _escape_inline_text("<x>") == "&lt;x&gt;"
    assert _escape_inline_text("a&b") == "a&amp;b"


def test_inline_escape_handles_backslash_first():
    """已有反斜杠要转义，否则我们加的反斜杠会被二次处理。"""
    assert _escape_inline_text("a\\b") == "a\\\\b"


def test_inline_escape_escapes_emphasis_and_links():
    for ch in ("*", "_", "`", "[", "]"):
        assert _escape_inline_text(ch) == "\\" + ch


def test_block_escape_handles_leading_markers():
    assert _escape_block_markers("#标题") == "\\#标题"
    assert _escape_block_markers("- 项") == "\\- 项"
    assert _escape_block_markers("1. 项") == "1\\. 项"
    assert _escape_block_markers("===") == "\\==="


def test_block_escape_keeps_indent_as_entities():
    assert _escape_block_markers("    x") == "&#32;&#32;&#32;&#32;x"
    assert _escape_block_markers("\tx") == "&#9;x"


# ==================== 4. 正牌 Markdown 仍要正常工作 ====================

def test_real_heading_style_still_maps_to_heading(tmp_path):
    """转义不能妨碍"真的标题"：DOCX 标题样式仍应变成 `#` 并往返回标题。"""
    src = tmp_path / "src.docx"
    doc = Document()
    doc.add_heading("真正的标题", level=1)
    doc.add_paragraph("正文")
    doc.save(str(src))

    md = import_from_docx(str(src))
    assert "# 真正的标题" in md, md

    _, out = _roundtrip_paras(["真正的标题"], tmp_path)   # 仅确保不抛错


def test_real_list_style_still_maps_to_list(tmp_path):
    src = tmp_path / "src.docx"
    doc = Document()
    doc.add_paragraph("第一项", style="List Bullet")
    doc.add_paragraph("第二项", style="List Bullet")
    doc.save(str(src))

    md = import_from_docx(str(src))
    lines = [ln for ln in md.splitlines() if ln.strip()]
    assert all(ln.startswith("- ") for ln in lines), lines
    assert "\\-" not in md, "真正的列表项不该被转义"


def test_table_cells_are_escaped(tmp_path):
    """表格单元格里的 `<`/`&` 同样要转义。"""
    src = tmp_path / "src.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "表头"
    table.cell(1, 0).text = "flags=<UP>&x"
    doc.save(str(src))

    md = import_from_docx(str(src))

    assert "&lt;UP&gt;" in md and "&amp;x" in md, md
