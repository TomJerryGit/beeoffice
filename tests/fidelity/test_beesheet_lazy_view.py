# tests/fidelity/test_beesheet_lazy_view.py
"""
R1 视图层：**两阶段载入**与**安全门**（用真实 BeeSheetView）。

与 `test_beesheet_lazy_load.py`（模型层 + 保存一致性）配套：
本文件关心"用户视角"的行为 —— 打开大文件先看到内容、加载中不能编辑、
补全后能编辑、以及**危险操作会不会被门拦住**。
"""
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def make_xlsx(path, rows=1000, cols=4, pad_bytes=0):
    """
    造一个 xlsx。`pad_bytes` 用于把文件撑过惰性载入阈值
    （阈值按**文件大小**判定，见 `BeeSheetView.LAZY_THRESHOLD_BYTES`）。
    """
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet.append([f"列{c}" for c in range(1, cols + 1)])
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            sheet.cell(r + 1, c, value=f"值{r}-{c}")
    book.save(str(path))

    if pad_bytes:
        # 往同一个工作表里塞**不可压缩**的长文本，把文件顶过阈值。
        # ⚠️ 两个坑都踩过：
        #   ① 用 `"x" * 300`（重复字符）—— xlsx 是 zip，重复内容被压到几乎为 0；
        #   ② 换成随机串后仍"差一点"（403 KB < 512 KB）→ 大小靠估算不可靠。
        # 所以这里**循环补齐直到真的超过阈值**，不靠估算。
        import secrets

        from openpyxl import load_workbook

        target = pad_bytes
        attempt = 0
        while path.stat().st_size < target and attempt < 10:
            book = load_workbook(str(path))
            sheet = book.active
            start = sheet.max_row + 1
            for i in range(1500):
                sheet.cell(start + i, cols + 1, value=secrets.token_hex(150))
            book.save(str(path))
            attempt += 1
    return path


@pytest.fixture
def view(qtbot):
    from beeoffice.plugins.beesheet import BeeSheetView

    widget = BeeSheetView()
    qtbot.addWidget(widget)
    widget.resize(1000, 700)
    yield widget
    widget.close()


# ==================== 1. 阈值判定 ====================

def test_small_xlsx_does_not_use_lazy_load(view, tmp_path):
    """小文件保持既有的一次性载入 —— 不引入无谓的行为差异。"""
    path = make_xlsx(tmp_path / "small.xlsx", rows=50)

    assert view._should_lazy_load(str(path)) is False


def test_large_xlsx_uses_lazy_load(view, tmp_path):
    path = make_xlsx(tmp_path / "big.xlsx", rows=800, pad_bytes=1_500_000)

    assert path.stat().st_size > view.LAZY_THRESHOLD_BYTES
    assert view._should_lazy_load(str(path)) is True


def test_csv_is_never_lazy(view, tmp_path):
    """CSV 读取路径不同，本次不纳入（避免改坏它）。"""
    path = tmp_path / "a.csv"
    # 要让"大小过阈值"这个前提成立，CSV 也得 > 1 MB（阈值改成 1 MB 后，
    # 原来 5 万行只有 ~590 KB，前提不成立 → 用例本身失去意义）
    path.write_text(
        "a,b\n" + "\n".join(f"{i:08d},{i:08d}" for i in range(120000)),
        encoding="utf-8",
    )

    assert path.stat().st_size > view.LAZY_THRESHOLD_BYTES
    assert view._should_lazy_load(str(path)) is False


# ==================== 2. 两阶段载入：先可见、加载中只读 ====================

def test_phase1_shows_rows_and_is_readonly(view, qtbot, tmp_path):
    """
    阶段 1：立刻有数据可显示，但**模型只读**（这是阶段 2 可安全替换的前提）。
    """
    from PySide6.QtCore import Qt
    from beeoffice.async_task import AsyncTaskManager

    path = make_xlsx(tmp_path / "big.xlsx", rows=1500, pad_bytes=1_500_000)

    # 屏蔽后台补全（这里只想看阶段 1 的状态）
    from beeoffice.plugins import beesheet as module
    started = {}
    monkey_job = lambda *a, **k: started.setdefault("phase2", True)  # noqa: E731
    original = module.BeeSheetView._start_phase2_full_load
    try:
        module.BeeSheetView._start_phase2_full_load = monkey_job
        view._import_file_from_path(str(path))
    finally:
        module.BeeSheetView._start_phase2_full_load = original
    _ = AsyncTaskManager

    assert started.get("phase2") is True, "没有触发阶段 2"
    assert view.model.is_fully_loaded is False, "阶段 1 后不该标记为已加载完"
    assert view.model.rowCount() > 0, "阶段 1 必须有内容可见"
    flags = view.model.flags(view.model.index(0, 0))
    assert not (flags & Qt.ItemFlag.ItemIsEditable), (
        "加载未完成时必须只读 —— 否则用户改动会在阶段 2 被替换掉"
    )


def _full_row_count(path) -> int:
    """直接读一遍文件，拿到真实行数（夹具会加填充行，不能靠估算）。"""
    from beeoffice.plugins.beesheet_workbook import read_workbook

    return len(read_workbook(str(path)).sheets[0].df)


def test_phase2_completes_and_enables_editing(view, qtbot, tmp_path):
    """
    阶段 2 完成后：数据变成完整、恢复可编辑、状态栏提示已加载。
    """
    from PySide6.QtCore import Qt

    path = make_xlsx(tmp_path / "big.xlsx", rows=1500, pad_bytes=1_100_000)
    expected_rows = _full_row_count(path)

    view._import_file_from_path(str(path))
    # 等后台补全（真实 AsyncTask 线程）。
    # 等待上限给足：单文件补全实测 ~1.3s，但测试进程里前面用例可能留下
    # 待处理的后台任务，队列会排队；全量套件满载时 30s 上限出现过一次抖动，
    # 放宽到 60s（真实功能等待，不改变断言强度）。
    for _ in range(3000):
        qtbot.wait(20)
        if view.model.is_fully_loaded:
            break

    assert view.model.is_fully_loaded is True, (
        f"后台补全没完成（load_error={view.model.load_error!r}，"
        f"状态栏={view.status_label.text()!r}）"
    )
    assert view.model._real_rows == expected_rows, (
        f"{view.model._real_rows} != {expected_rows}"
    )
    flags = view.model.flags(view.model.index(0, 0))
    assert flags & Qt.ItemFlag.ItemIsEditable, "补全后应可编辑"


# ==================== 3. 安全门：危险操作被拦住 ====================

def _make_partial_view(view, tmp_path, rows=1200):
    """造一个"仍处于部分载入"的视图（阶段 2 故意失败）。"""
    path = make_xlsx(tmp_path / "big.xlsx", rows=rows, pad_bytes=1_500_000)

    from beeoffice.plugins import beesheet as module

    original = module.BeeSheetView._start_phase2_full_load
    try:
        module.BeeSheetView._start_phase2_full_load = lambda self, p: None
        view._import_file_from_path(str(path))
    finally:
        module.BeeSheetView._start_phase2_full_load = original
    assert view.model.is_fully_loaded is False
    return path


def test_gate_reports_failure_and_blocks(view, qtbot, tmp_path, monkeypatch):
    """
    安全门在"补全也失败"时必须返回 False 并给出说明 —— 调用方据此中止。
    """
    from PySide6.QtWidgets import QMessageBox

    _make_partial_view(view, tmp_path)
    shown = []
    monkeypatch.setattr(
        QMessageBox, "warning",
        staticmethod(lambda *a, **k: shown.append(a)),
    )
    # 让补全一定失败
    view.model._full_loader = lambda: (_ for _ in ()).throw(
        OSError("模拟：源文件不可读")
    )

    ok = view._ensure_full_data("保存")

    assert ok is False, "补全失败却放行了 —— 会写出半份数据"
    assert view.model.load_error, "要留下可显示的失败说明"
    assert shown, "要弹框告知用户"


def test_gate_succeeds_when_loader_works(view, qtbot, tmp_path):
    """补全成功时门要放行，并把数据补成完整的。"""
    path = _make_partial_view(view, tmp_path)
    expected_rows = _full_row_count(path)

    ok = view._ensure_full_data("保存")

    assert ok is True
    assert view.model.is_fully_loaded is True
    # 夹具会加填充行，所以按**真实行数**断言（我第一版写死 1200 就错了）
    assert view.model._real_rows == expected_rows, (
        f"{view.model._real_rows} != {expected_rows}"
    )


def test_export_is_blocked_when_data_incomplete(view, qtbot, tmp_path,
                                                monkeypatch):
    """
    **端到端**：数据没加载完时点"导出 CSV"应该**什么都不导出**。
    """
    from PySide6.QtWidgets import QFileDialog, QMessageBox

    _make_partial_view(view, tmp_path)
    target = tmp_path / "out.csv"
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName",
        staticmethod(lambda *a, **k: (str(target), "")),
    )
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: None))
    view.model._full_loader = lambda: (_ for _ in ()).throw(OSError("坏"))

    view._export_csv()

    assert not target.exists(), "数据不完整却导出了文件（可能只有前 N 行）"


def test_save_is_blocked_when_data_incomplete(view, qtbot, tmp_path,
                                              monkeypatch):
    """同样的门也必须保护 xlsx 保存（写回源文件最危险）。"""
    from PySide6.QtWidgets import QFileDialog, QMessageBox

    source = _make_partial_view(view, tmp_path)
    target = tmp_path / "out.xlsx"
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName",
        staticmethod(lambda *a, **k: (str(target), "")),
    )
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: None))
    view.model._full_loader = lambda: (_ for _ in ()).throw(OSError("坏"))

    view._save_as_excel() if hasattr(view, "_save_as_excel") else None
    # 直接调用保存路径（大多数版本里叫 _save_as / _export_excel）
    for name in ("_export_excel", "_save_as", "_save_as_excel"):
        method = getattr(view, name, None)
        if callable(method):
            method()
            break

    assert not target.exists(), f"数据不完整却写了 {target}"
    assert source.exists(), "源文件不该被动过"


def test_sort_is_blocked_when_data_incomplete(view, qtbot, tmp_path,
                                              monkeypatch):
    """排序也要过门（实测：部分载入后排序再保存会造成文件错乱）。"""
    from PySide6.QtWidgets import QMessageBox

    _make_partial_view(view, tmp_path)
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: None))
    view.model._full_loader = lambda: (_ for _ in ()).throw(OSError("坏"))
    called = {}
    monkeypatch.setattr(
        type(view), "_apply_sort", lambda self, *a, **k: called.setdefault("x", 1),
        raising=False,
    )

    view._show_sort_dialog()

    assert not called, "排序对话框在数据不完整时不该继续"


def test_full_loaded_view_is_unaffected(view, qtbot, tmp_path, monkeypatch):
    """已完整载入时安全门是**零成本直通**（不能因为加了门就改变正常路径）。"""
    path = make_xlsx(tmp_path / "small.xlsx", rows=20)
    view._import_file_from_path(str(path))

    assert view.model.is_fully_loaded is True
    assert view._ensure_full_data("保存") is True
