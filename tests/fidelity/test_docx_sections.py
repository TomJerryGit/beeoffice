# tests/fidelity/test_docx_sections.py
"""
保真度阶段 A：DOCX **多分节**往返。

缺陷（实测）
------------
`docx_sections.docx` 有 2 个分节（第二节横向 A4、页边距不同）。原先只读
`doc.sections[0]`，往返后 `sections 2 → 1` —— 用户可见后果是
**横向的那一节变成纵向**，页边距也随第一节。

两个关键实现点
--------------
1. OOXML 里每个分节的 `w:sectPr` 挂在**该节最后一个段落**的 `w:pPr`；
   最后一节的挂在 body 上。
2. **不能用段落索引定位**：承载 `sectPr` 的段落常常是**空段落**
   （python-docx 的 `add_section()` 就是新建一个空段承载它），而空段落在
   Markdown 里会丢（实测顶层段落 5 → 4）。所以改用"该节最后一个**非空**
   段落的文字"作锚点，导出时按文字定位 —— 对索引偏移免疫。
"""
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document                                  # noqa: E402

from beeoffice.plugins.bee_export import (                 # noqa: E402
    _read_sections,
    apply_sections,
    export_to_docx,
    import_from_docx,
    read_docx_extras,
)


def _make_two_section_docx(path: Path) -> Path:
    """第一节纵向 Letter，第二节横向 A4 且页边距不同。"""
    from docx import Document as D
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.shared import Mm

    doc = D()
    doc.add_heading("第一节（纵向）", level=1)
    doc.add_paragraph("这一段属于第一节。")

    second = doc.add_section(WD_SECTION.NEW_PAGE)
    second.orientation = WD_ORIENT.LANDSCAPE
    second.page_width = Mm(297)
    second.page_height = Mm(210)
    second.left_margin = Mm(15)
    second.right_margin = Mm(15)
    doc.add_heading("第二节（横向）", level=1)
    doc.add_paragraph("这一段属于第二节。")
    doc.save(str(path))
    return path


def _roundtrip(src: Path, tmp_path: Path) -> Path:
    out = tmp_path / "out.docx"
    extras = read_docx_extras(src)
    assert export_to_docx(
        import_from_docx(str(src)), str(out),
        page_settings=extras, base_dir=str(tmp_path),
        header_text=extras.get("header_text"),
        footer_text=extras.get("footer_text"),
        table_merges=extras.get("table_merges"),
        sections=extras.get("sections"),
    ) is True
    return out


# ==================== 1. 读取 ====================

def test_reads_all_sections(tmp_path):
    src = _make_two_section_docx(tmp_path / "s.docx")

    sections = _read_sections(Document(str(src)))

    assert len(sections) == 2, sections


def test_reads_per_section_geometry(tmp_path):
    src = _make_two_section_docx(tmp_path / "s.docx")

    sections = _read_sections(Document(str(src)))

    assert sections[0]["orientation"] == "portrait"
    assert sections[0]["page_width_mm"] == pytest.approx(216, abs=1)
    assert sections[1]["orientation"] == "landscape"
    assert sections[1]["page_width_mm"] == pytest.approx(297, abs=1)
    assert sections[1]["page_height_mm"] == pytest.approx(210, abs=1)
    assert sections[1]["margin_left"] == pytest.approx(15, abs=0.5)


def test_section_anchor_is_last_non_empty_paragraph(tmp_path):
    """
    承载 sectPr 的段落是空的，所以锚点应是**往前找的最后一个非空段落**。
    """
    src = _make_two_section_docx(tmp_path / "s.docx")

    sections = _read_sections(Document(str(src)))

    assert sections[0]["after_text"] == "这一段属于第一节。", sections[0]
    # 最后一节挂在 body 上，没有锚点
    assert sections[-1]["end_paragraph"] is None


def test_single_section_document(tmp_path):
    from docx import Document as D

    src = tmp_path / "one.docx"
    doc = D()
    doc.add_paragraph("只有一节")
    doc.save(str(src))

    sections = _read_sections(Document(str(src)))

    assert len(sections) == 1
    assert sections[0]["end_paragraph"] is None


def test_read_docx_extras_includes_sections(tmp_path):
    src = _make_two_section_docx(tmp_path / "s.docx")

    extras = read_docx_extras(src)

    assert len(extras.get("sections", [])) == 2


# ==================== 2. 重建 ====================

def test_apply_recreates_second_section(tmp_path):
    """**核心**：重建后必须仍是 2 个分节，且第二节是横向。"""
    src = _make_two_section_docx(tmp_path / "s.docx")
    sections = _read_sections(Document(str(src)))

    doc = Document()
    doc.add_paragraph("第一节（纵向）")
    doc.add_paragraph("这一段属于第一节。")
    doc.add_paragraph("第二节（横向）")
    doc.add_paragraph("这一段属于第二节。")

    applied = apply_sections(doc, sections)

    assert applied == 1, f"应重建 1 个中间分节，实际 {applied}"
    assert len(doc.sections) == 2
    out = tmp_path / "rebuilt.docx"
    doc.save(str(out))

    again = Document(str(out))
    assert len(again.sections) == 2
    assert again.sections[1].page_width.mm == pytest.approx(297, abs=1)


def test_anchor_survives_missing_empty_paragraph(tmp_path):
    """
    导出文档里**没有**那个空段落（Markdown 丢了它），锚点仍应定位正确。

    这正是"按文字定位"而非"按索引定位"的原因：索引会偏，文字不会。
    """
    src = _make_two_section_docx(tmp_path / "s.docx")
    sections = _read_sections(Document(str(src)))

    # 模拟导出文档：没有承载 sectPr 的那个空段
    doc = Document()
    doc.add_paragraph("第一节（纵向）")
    doc.add_paragraph("这一段属于第一节。")     # ← 分界应在这里
    doc.add_paragraph("第二节（横向）")
    doc.add_paragraph("这一段属于第二节。")

    apply_sections(doc, sections)

    out = tmp_path / "anchored.docx"
    doc.save(str(out))

    # 分节边界应落在"这一段属于第一节。"这一段上
    from docx.oxml.ns import qn

    holder = None
    for idx, para in enumerate(Document(str(out)).paragraphs):
        pPr = para._p.find(qn("w:pPr"))
        if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
            holder = idx
            break
    assert holder == 1, f"分节边界落错了段落: {holder}"


def test_falls_back_to_index_when_anchor_missing(tmp_path):
    """锚点文字找不到时应退回索引，而不是静默不建分节。"""
    sections = [{
        "end_paragraph": 0, "after_text": "这段文字不存在",
        "page_width_mm": 200, "page_height_mm": 280,
        "orientation": "portrait", "margin_top": 20,
    }, {
        "end_paragraph": None, "after_text": "",
        "page_width_mm": 210, "page_height_mm": 297,
        "orientation": "portrait",
    }]

    doc = Document()
    doc.add_paragraph("第一段")
    doc.add_paragraph("第二段")

    assert apply_sections(doc, sections) == 1


def test_empty_sections_is_noop():
    doc = Document()
    assert apply_sections(doc, None) == 0
    assert apply_sections(doc, []) == 0
    assert len(doc.sections) == 1


# ==================== 3. 端到端 ====================

def test_roundtrip_preserves_section_count(tmp_path):
    """**用户可见的核心断言**：往返后仍是 2 个分节。"""
    from tools.ooxml_probe import probe

    src = _make_two_section_docx(tmp_path / "s.docx")
    out = _roundtrip(src, tmp_path)

    before = probe(src)["features"]["sections"]
    after = probe(out)["features"]["sections"]
    assert before == 2 and after == 2, f"{before} → {after}"


def test_roundtrip_preserves_landscape_section(tmp_path):
    """第二节必须仍是横向 —— 否则用户打印出来就不对了。"""
    src = _make_two_section_docx(tmp_path / "s.docx")
    out = _roundtrip(src, tmp_path)

    sections = Document(str(out)).sections
    assert len(sections) == 2
    assert sections[0].page_width.mm < sections[0].page_height.mm, "第一节应纵向"
    assert sections[1].page_width.mm > sections[1].page_height.mm, (
        "第二节应横向"
    )


def test_roundtrip_preserves_per_section_margins(tmp_path):
    src = _make_two_section_docx(tmp_path / "s.docx")
    out = _roundtrip(src, tmp_path)

    sections = Document(str(out)).sections
    assert sections[0].left_margin.mm == pytest.approx(31.75, abs=1)
    assert sections[1].left_margin.mm == pytest.approx(15, abs=1)


def test_content_order_is_preserved(tmp_path):
    """重建分节不能打乱正文顺序。"""
    src = _make_two_section_docx(tmp_path / "s.docx")
    out = _roundtrip(src, tmp_path)

    texts = [p.text.strip() for p in Document(str(out)).paragraphs
             if p.text.strip()]
    assert texts == [
        "第一节（纵向）", "这一段属于第一节。",
        "第二节（横向）", "这一段属于第二节。",
    ], texts


# ==================== 4. 编辑器与产品接线 ====================

def test_editor_carries_sections(beewrite_editor):
    records = [{"end_paragraph": None, "after_text": "",
                "page_width_mm": 210, "page_height_mm": 297,
                "orientation": "portrait"}]

    beewrite_editor.apply_docx_extras({"sections": records})

    assert beewrite_editor.get_docx_extras()["sections"] == records


def test_editor_defaults_to_no_sections(beewrite_editor):
    assert beewrite_editor.get_docx_extras()["sections"] == []


def test_write_docx_recreates_sections(beewrite_editor, tmp_path):
    """产品级：编辑器带着分节信息导出后，文件里必须真有 2 个分节。"""
    from tools.ooxml_probe import probe

    beewrite_editor.set_content(
        "# 第一节（纵向）\n\n这一段属于第一节。\n\n"
        "## 第二节（横向）\n\n这一段属于第二节。\n"
    )
    beewrite_editor.apply_docx_extras({"sections": [
        {
            "end_paragraph": 1, "after_text": "这一段属于第一节。",
            "page_width_mm": 215.9, "page_height_mm": 279.4,
            "orientation": "portrait",
            "margin_top": 25, "margin_bottom": 25,
            "margin_left": 31, "margin_right": 31,
        },
        {
            "end_paragraph": None, "after_text": "",
            "page_width_mm": 297, "page_height_mm": 210,
            "orientation": "landscape",
            "margin_top": 15, "margin_bottom": 15,
            "margin_left": 15, "margin_right": 15,
        },
    ]})

    out = tmp_path / "sections.docx"
    assert beewrite_editor._write_docx(str(out)) is True

    assert probe(out)["features"]["sections"] == 2, (
        "导出的文件里没有 2 个分节"
    )
