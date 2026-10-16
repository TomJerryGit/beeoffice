# tests/fidelity/test_p1_beewrite_remnants.py
"""
第十批：蜜写剩余三条（审计报告 §2.13）。

1. **Markdown 表格双击编辑抹掉列对齐**：`parse_md_table()` 把分隔行
   **整个丢弃**，`render_md_table()` 又**硬编码** `---` →
   `|:---|---:|` 编辑一次就被抹成 `|---|---|`。
2. **富文本→DOCX：超链接与单元格多段落丢失**：run 导出只处理图片与文本，
   链接 URL **整条丢掉**；单元格用 `docx_cell.text = cell_block.text()`
   **只取首段**。
3. **`set_markdown` 忽略嵌套列表缩进**：正则捕获了 `lm.group(1)`（缩进）
   却**从未使用** → 多级列表被压成一级。
"""
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


# ==================== 1. Markdown 表格列对齐 ====================

def test_parse_and_render_preserve_alignment():
    from beeoffice.plugins.beewrite_utils import (
        parse_md_table, parse_md_table_aligns, render_md_table,
    )

    lines = ["| 名称 | 数量 | 备注 |", "|:---|:---:|---:|", "| 甲 | 1 | 说明 |"]

    rows = parse_md_table(lines)
    aligns = parse_md_table_aligns(lines)

    assert aligns == ["left", "center", "right"]
    out = render_md_table(rows, aligns)
    assert out.split("\n")[1] == "|:---|:---:|---:|"


def test_render_without_aligns_stays_backward_compatible():
    """不传对齐时保持旧行为（`|---|---|`）—— 不能破坏其它调用方。"""
    from beeoffice.plugins.beewrite_utils import parse_md_table, render_md_table

    rows = parse_md_table(["| A | B |", "|---|---|", "| 1 | 2 |"])
    assert render_md_table(rows).split("\n")[1] == "|---|---|"


def test_no_alignment_columns_render_plain():
    from beeoffice.plugins.beewrite_utils import parse_md_table_aligns

    assert parse_md_table_aligns(["| a | b |", "|---|---|"]) == [None, None]
    # 非表格输入不炸
    assert parse_md_table_aligns(["随便一行"]) == []


@pytest.fixture
def editor(qtbot):
    from beeoffice.plugins.beewrite import BeeWriteEditor

    widget = BeeWriteEditor()
    qtbot.addWidget(widget)
    widget._set_status = lambda *a, **k: None
    return widget


def test_editor_remembers_alignment_across_rerender(editor):
    """
    **核心回归**：解析过一张表之后，后续任何整表重渲染都要带上对齐。

    原来双击编辑（或增删行列）会调用 `_render_md_table(rows)`，
    只渲染内容 → 对齐丢失。
    """
    editor._switch_mode("markdown")
    lines = ["| 名称 | 数量 |", "|:---:|---:|", "| 甲 | 1 |"]

    editor._parse_md_table(lines)            # 解析时记住对齐
    rendered = editor._render_md_table([["名称", "数量"], ["甲", "1"]])

    assert rendered.split("\n")[1] == "|:---:|---:|", rendered


def test_editor_render_without_prior_parse_is_plain(editor):
    """没解析过就渲染：按旧行为输出（不凭空造对齐）。"""
    editor._md_table_aligns = []
    rendered = editor._render_md_table([["A", "B"]])
    assert rendered.split("\n")[1] == "|---|---|"


# ==================== 2. DOCX 超链接与单元格多段落 ====================

def test_hyperlink_is_exported_as_real_word_link(editor, tmp_path):
    editor._switch_mode("rich_text")
    editor.rich_editor.setHtml(
        '<p>正文</p><p><a href="https://example.com/x">示例链接</a></p>'
    )

    target = tmp_path / "链接.docx"
    assert editor._write_docx(str(target)) is True

    with zipfile.ZipFile(target) as archive:
        rels = [n for n in archive.namelist() if n.endswith(".rels")]
        joined = "\n".join(
            archive.read(n).decode("utf-8", "ignore") for n in rels
        )
        document = archive.read("word/document.xml").decode("utf-8", "ignore")

    assert "example.com/x" in joined, "URL 没有写成外部关系（链接点不动）"
    assert "<w:hyperlink" in document, "没有写成 Word 超链接"
    assert "示例链接" in document, "链接文字丢了"

    from docx import Document

    assert [p.text for p in Document(str(target)).paragraphs] == [
        "正文", "示例链接"
    ]


def test_cell_with_multiple_paragraphs_keeps_them_all(editor, tmp_path):
    """**核心回归**：单元格里的多段落不能被压成一段。"""
    from PySide6.QtGui import QTextTableFormat

    editor._switch_mode("rich_text")
    editor.rich_editor.setPlainText("")
    cursor = editor.rich_editor.textCursor()
    table = cursor.insertTable(1, 1, QTextTableFormat())
    cell = table.cellAt(0, 0)
    cell_cursor = cell.firstCursorPosition()
    cell_cursor.insertText("第一段")
    cell_cursor.insertBlock()
    cell_cursor.insertText("第二段")

    target = tmp_path / "单元格.docx"
    assert editor._write_docx(str(target)) is True

    from docx import Document

    document = Document(str(target))
    assert len(document.tables) == 1
    texts = [p.text for p in document.tables[0].cell(0, 0).paragraphs]
    assert texts == ["第一段", "第二段"], texts


# ==================== 3. 嵌套列表缩进 ====================

def test_nested_list_indent_is_kept(editor):
    """**核心回归**：多级列表不能被压平（`lm.group(1)` 从被忽略到被使用）。"""
    editor._switch_mode("rich_text")
    editor.rich_editor.set_markdown("- 一级\n  - 二级\n    - 三级")

    document = editor.rich_editor.document()
    block = document.firstBlock()
    levels = []
    while block.isValid():
        text_list = block.textList()
        levels.append((block.text(), text_list.format().indent()
                       if text_list else None))
        block = block.next()

    levels = [(t, lv) for t, lv in levels if t]
    assert [t for t, _ in levels] == ["一级", "二级", "三级"]
    assert [lv for _, lv in levels] == [1, 2, 3], levels


def test_tab_indent_counts_as_one_level(editor):
    editor._switch_mode("rich_text")
    editor.rich_editor.set_markdown("- 甲\n\t- 乙")

    document = editor.rich_editor.document()
    block = document.firstBlock()
    levels = []
    while block.isValid():
        text_list = block.textList()
        if block.text():
            levels.append(text_list.format().indent() if text_list else None)
        block = block.next()

    assert levels == [1, 2], levels
