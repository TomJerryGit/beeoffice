# tests/fidelity/test_beecmd_merge.py
"""
`beecmd merge`（v0.2 多表汇总）的回归。

## 为什么需要
"每月一张表 → 汇总成一张"是小公司/普通人最高频的表格工作之一，
也是《路线图》P1-2 三个自动化模板里的第 ② 个（Excel 多表汇总）。

## 本文件锁住的行为
* `append`：按列名纵向追加 + 「来源文件」列；列不一致时取并集、缺列留空
* `sheets`：每个输入文件的每个工作表进入同一工作簿（名冲突自动去重）
* 错误输入（<2 个文件、错误输出格式、缺文件）→ 干净报错（退出码 1，不甩 traceback）
* `batch` 支持 merge 任务并写审计流水
* **零 Qt**：merge 路径不得 import PySide6（服务器/容器可用）
"""
import json
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def _run(args):
    from beeoffice import beecmd

    return beecmd.main(args)


def _write_csv(path: Path, text: str) -> Path:
    path.write_text(text, encoding="utf-8-sig")
    return path


# ==================== 1. append：同列纵向追加 ====================

def test_merge_append_same_columns(tmp_path):
    a = _write_csv(tmp_path / "1月.csv", "产品,数量\nA,10\nB,20\n")
    b = _write_csv(tmp_path / "2月.csv", "产品,数量\nC,30\n")
    out = tmp_path / "汇总.xlsx"

    assert _run(["merge", str(a), str(b), "-o", str(out)]) == 0
    assert out.is_file()

    from openpyxl import load_workbook

    wb = load_workbook(str(out))
    ws = wb["汇总"]
    rows = list(ws.values)
    wb.close()
    assert rows[0][0] == "来源文件"
    assert rows[0][1:] == ("产品", "数量")
    assert len(rows) == 4, rows                      # 表头 + 3 行数据
    sources = {r[0] for r in rows[1:]}
    assert sources == {"1月.csv", "2月.csv"}
    assert rows[1][2] == 10


# ==================== 2. append：列不一致取并集 ====================

def test_merge_append_union_columns(tmp_path):
    a = _write_csv(tmp_path / "a.csv", "x\n1\n2\n")
    b = _write_csv(tmp_path / "b.csv", "x,y\n3,9\n")
    out = tmp_path / "u.xlsx"

    assert _run(["merge", str(a), str(b), "-o", str(out)]) == 0

    from openpyxl import load_workbook

    wb = load_workbook(str(out))
    rows = list(wb["汇总"].values)
    wb.close()
    header = rows[0]
    assert set(header) == {"来源文件", "x", "y"}, header
    # a 的两行 y 为空，b 的一行 y=9
    y_index = header.index("y")
    values = [r[y_index] for r in rows[1:]]
    assert sum(1 for v in values if v not in (None, "")) == 1


# ==================== 3. sheets：每个表一个工作表 ====================

def test_merge_sheets_mode(tmp_path):
    import pandas as pd

    a = tmp_path / "一月.xlsx"
    b = tmp_path / "二月.xlsx"
    with pd.ExcelWriter(a) as writer:
        pd.DataFrame({"x": [1]}).to_excel(writer, sheet_name="销售", index=False)
        pd.DataFrame({"y": [2]}).to_excel(writer, sheet_name="库存", index=False)
    with pd.ExcelWriter(b) as writer:
        pd.DataFrame({"x": [3]}).to_excel(writer, sheet_name="销售", index=False)

    out = tmp_path / "合并.xlsx"
    assert _run(["merge", str(a), str(b), "-o", str(out),
                 "--mode", "sheets"]) == 0

    from openpyxl import load_workbook

    wb = load_workbook(str(out))
    names = wb.sheetnames
    wb.close()
    assert len(names) == 3, names
    assert len(set(names)) == 3, "工作表名必须去重"


# ==================== 4. 错误输入要干净失败 ====================

def test_merge_requires_two_files(tmp_path, capsys):
    a = _write_csv(tmp_path / "a.csv", "x\n1\n")
    assert _run(["merge", str(a), "-o", str(tmp_path / "o.xlsx")]) == 1
    assert "至少需要 2 个" in capsys.readouterr().err


def test_merge_rejects_non_xlsx_output(tmp_path, capsys):
    a = _write_csv(tmp_path / "a.csv", "x\n1\n")
    b = _write_csv(tmp_path / "b.csv", "x\n2\n")
    assert _run(["merge", str(a), str(b), "-o", str(tmp_path / "o.csv")]) == 1
    assert ".xlsx" in capsys.readouterr().err


def test_merge_missing_file(tmp_path, capsys):
    a = _write_csv(tmp_path / "a.csv", "x\n1\n")
    assert _run(["merge", str(a), str(tmp_path / "nope.csv"),
                 "-o", str(tmp_path / "o.xlsx")]) == 1
    assert "不存在" in capsys.readouterr().err


# ==================== 5. batch 支持 merge ====================

def test_batch_supports_merge(tmp_path):
    a = _write_csv(tmp_path / "a.csv", "x\n1\n")
    b = _write_csv(tmp_path / "b.csv", "x\n2\n")
    out = tmp_path / "out.xlsx"
    audit = tmp_path / "audit.jsonl"
    jobs = tmp_path / "jobs.json"
    jobs.write_text(json.dumps({
        "jobs": [
            {"op": "merge", "srcs": [str(a), str(b)], "out": str(out)}
        ]
    }), encoding="utf-8-sig")

    assert _run(["batch", str(jobs), "--audit", str(audit)]) == 0
    lines = [json.loads(x) for x in
             audit.read_text(encoding="utf-8").splitlines() if x.strip()]
    assert len(lines) == 1
    assert lines[0]["op"] == "merge" and lines[0]["ok"] is True
    assert out.is_file()


# ==================== 6. 零 Qt 证明 ====================

def test_merge_path_does_not_import_qt(tmp_path):
    """
    **架构约束**：多表汇总必须能在没有 Qt 的服务器/容器里跑。

    ⚠️ 沙箱限制：子进程不能用管道捕获输出，所以让子进程把结果写文件。
    """
    a = _write_csv(tmp_path / "a.csv", "x\n1\n")
    b = _write_csv(tmp_path / "b.csv", "x\n2\n")
    out = tmp_path / "out.xlsx"
    result_file = tmp_path / "qt_probe.txt"
    script = (
        "import sys\n"
        "from beeoffice import beecmd\n"
        f"code = beecmd.main(['merge', {str(a)!r}, {str(b)!r},"
        f" '-o', {str(out)!r}])\n"
        "loaded = [m for m in sys.modules if m.startswith(('PySide6', 'shiboken'))]\n"
        f"open({str(result_file)!r}, 'w', encoding='utf-8').write("
        "f'code={code};qt={\",\".join(loaded)}')\n"
    )
    completed = subprocess.run(
        [sys.executable, "-c", script], cwd=str(ROOT),
        env={**__import__("os").environ, "PYTHONPATH": str(ROOT)},
        timeout=180,
    )
    assert completed.returncode == 0, "子进程 merge 失败"
    payload = result_file.read_text(encoding="utf-8")
    assert payload.startswith("code=0;"), payload
    assert payload.endswith("qt="), f"merge 路径不该加载 Qt：{payload}"
    assert out.is_file()
