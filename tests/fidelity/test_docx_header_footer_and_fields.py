# tests/fidelity/test_docx_header_footer_and_fields.py
"""
保真度阶段 A：**页眉页脚部件**与**域**的往返。

两个真实文档撞出来的问题
------------------------
### 1. 页脚整个丢失（`footer_parts 1 → 0`，3/7 份）

根因比想象的深：python-docx 的 `section.footer.paragraphs` **只返回 1 个空段落**，
而那 3 份文档的页码其实在**嵌套的内容控件**里：

    ftr
    ├── sdt → sdtContent → sdt → sdtContent → p   ← 真正的页码段落（含 PAGE/NUMPAGES 域）
    └── p                                          ← 空段落（python-docx 只看到这个）

**python-docx 不进入 `w:sdt`**，所以 `.text` 是空串 → 现有的"只搬文字"逻辑
判定"没有页脚" → 整个部件都不写。

修法：**整段搬运页眉/页脚部件的 XML**（保留域与内容控件）。
只在部件含**关系引用**（图片，`r:id`/`r:embed`）时退回文字路径 ——
那种内容不能只靠搬 XML（关系 id 在新文档里不存在，会变成悬空引用）。

### 2. "域"其实是 **HYPERLINK 域**

`fields: 6 → 0` 看着像"目录/页码域丢失"，实测指令是：

    ' HYPERLINK "https://blog.csdn.net/…"'

也就是**以域形式存储的超链接**（Word 的另一种存法）。之前我的超链接修复只覆盖了
`w:hyperlink` 元素，没覆盖域形式 —— 于是链接文字还在、**链接目标丢了**。

修法：遍历段落时用状态机识别复杂域（`w:fldChar` begin/separate/end +
`w:instrText`）与简单域（`w:fldSimple`），把 `HYPERLINK` 指令里的 URL 附到
域的**结果区** run 上，输出成 `[文字](url)`。

结果是 `fields 6 → 0` 但 `hyperlinks 0 → 6` —— **信息没丢，只是换了更好的表示法**
（真 `w:hyperlink` 元素）。
"""
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document                                  # noqa: E402
from docx.oxml import OxmlElement                          # noqa: E402
from docx.oxml.ns import qn                                # noqa: E402

from beeoffice.plugins.bee_export import (                 # noqa: E402
    _paragraph_text,
    _part_has_external_refs,
    _url_from_field_instruction,
    export_to_docx,
    import_from_docx,
    read_docx_extras,
)

W = qn("w:t")


# ==================== 1. 页脚部件（含域与内容控件）====================

def _make_docx_with_field_footer(path: Path) -> Path:
    """
    构造一个"页码页脚"，并把内容包进**内容控件 `w:sdt`**。

    这样 `section.footer.paragraphs` 只会看到外层空段落 —— 与真实文档一致。
    """
    doc = Document()
    doc.add_paragraph("正文")
    section = doc.sections[0]
    section.footer.is_linked_to_previous = False

    # 造一个含 PAGE 域的段落
    para = OxmlElement("w:p")
    run = OxmlElement("w:r")
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run.append(fld)
    para.append(run)

    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    instr_run.append(instr)
    para.append(instr_run)

    sep_run = OxmlElement("w:r")
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run.append(sep)
    para.append(sep_run)

    text_run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "7"
    text_run.append(t)
    para.append(text_run)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    para.append(end_run)

    # 把该段落包进两层 sdt（与真实文档的嵌套结构一致）
    inner_sdt = OxmlElement("w:sdt")
    inner_content = OxmlElement("w:sdtContent")
    inner_content.append(para)
    inner_sdt.append(inner_content)
    outer_sdt = OxmlElement("w:sdt")
    outer_content = OxmlElement("w:sdtContent")
    outer_content.append(inner_sdt)
    outer_sdt.append(outer_content)

    footer = section.footer._element
    footer.insert(0, outer_sdt)

    doc.save(str(path))
    return path


def _roundtrip(src: Path, tmp_path: Path) -> Path:
    out = tmp_path / "out.docx"
    extras = read_docx_extras(src)
    assert export_to_docx(
        import_from_docx(str(src)), str(out),
        page_settings=extras, base_dir=str(tmp_path),
        header_text=extras.get("header_text"),
        footer_text=extras.get("footer_text"),
        header_xml=extras.get("header_xml", ""),
        footer_xml=extras.get("footer_xml", ""),
        table_merges=extras.get("table_merges"),
        sections=extras.get("sections"),
        bookmarks=extras.get("bookmarks"),
    ) is True
    return out


def test_footer_in_sdt_is_serialized(tmp_path):
    """内容控件里的页脚必须能被序列化出来（否则整部件都会丢）。"""
    src = _make_docx_with_field_footer(tmp_path / "a.docx")

    extras = read_docx_extras(src)

    assert extras.get("footer_xml"), "页脚部件 XML 为空 —— 又是 p.text 那种漏读"
    assert "PAGE" in extras["footer_xml"]


def test_footer_part_is_restored(tmp_path):
    """**核心**：带域的页脚往返后部件仍在（原先 `footer_parts 1 → 0`）。"""
    from tools.ooxml_probe import probe

    src = _make_docx_with_field_footer(tmp_path / "a.docx")
    out = _roundtrip(src, tmp_path)

    assert probe(src)["features"]["footer_parts"] == 1
    assert probe(out)["features"]["footer_parts"] == 1, "页脚部件又丢了"


def test_footer_field_instruction_survives(tmp_path):
    """
    不只是"部件在" —— **PAGE 域必须还是域**。

    若只搬文字，页码会变成写死的 "7"，换页后就是错的（比丢页脚更糟）。
    """
    src = _make_docx_with_field_footer(tmp_path / "a.docx")
    out = _roundtrip(src, tmp_path)

    doc = Document(str(out))
    footer = doc.sections[0].footer._element
    instr = [i.text for i in footer.iter(qn("w:instrText"))]

    assert "PAGE" in instr, f"PAGE 域没了（退化成写死文字）: {instr}"
    assert len(footer.findall(f".//{qn('w:sdt')}")) >= 1, "内容控件没保住"


def test_part_with_image_ref_falls_back(tmp_path):
    """含图片引用的部件不能靠搬 XML（关系 id 会悬空），应退回文字。"""
    assert _part_has_external_refs('<w:p><a:blip r:embed="rId3"/></w:p>')
    assert not _part_has_external_refs("<w:p><w:t>纯文字</w:t></w:p>")


# ==================== 2. HYPERLINK 域 → 真超链接 ====================

def test_url_from_field_instruction():
    assert _url_from_field_instruction(
        ' HYPERLINK "https://a.example/x" '
    ) == "https://a.example/x"
    assert _url_from_field_instruction(" HYPERLINK https://b.example") \
        == "https://b.example"
    assert _url_from_field_instruction(" PAGE ") == ""
    assert _url_from_field_instruction("") == ""


def _make_docx_with_hyperlink_field(path: Path, url="https://x.example/doc"):
    """构造一个用**域**存超链接的文档（Word 的另一种存法）。"""
    doc = Document()
    para = doc.add_paragraph()

    def _run(child):
        r = OxmlElement("w:r")
        r.append(child)
        para._p.append(r)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    _run(begin)

    instr = OxmlElement("w:instrText")
    instr.text = f' HYPERLINK "{url}" '
    _run(instr)

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    _run(separate)

    t = OxmlElement("w:t")
    t.text = "点这里"
    _run(t)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    _run(end)

    doc.save(str(path))
    return path


def test_hyperlink_field_becomes_markdown_link(tmp_path):
    src = _make_docx_with_hyperlink_field(tmp_path / "a.docx")

    md = import_from_docx(str(src))

    assert "[点这里](https://x.example/doc)" in md, md


def test_hyperlink_field_roundtrips_as_real_hyperlink(tmp_path):
    """
    **信息不丢**：域形式的超链接往返后变成真 `w:hyperlink`。

    指标上 `fields` 会归零、`hyperlinks` 会增加 —— 这是**表示法的改善**，
    不是丢失（用户看到的是同样的可点击链接）。
    """
    from tools.ooxml_probe import probe

    src = _make_docx_with_hyperlink_field(tmp_path / "a.docx")
    assert probe(src)["features"]["fields"] >= 1, "前置条件：源里有域"
    assert probe(src)["features"]["hyperlinks"] == 0, "前置条件：源里没有 w:hyperlink"

    out = _roundtrip(src, tmp_path)

    assert probe(out)["features"]["hyperlinks"] == 1, "超链接丢了"
    doc = Document(str(out))
    targets = []
    for para in doc.paragraphs:
        for link in para._p.findall(qn("w:hyperlink")):
            r_id = link.get(qn("r:id"))
            if r_id:
                targets.append(para.part.rels[r_id].target_ref)
    assert "https://x.example/doc" in targets, targets


def test_fldSimple_hyperlink_also_converted(tmp_path):
    """简单域（`w:fldSimple`）形式的超链接同样要转换。"""
    doc = Document()
    para = doc.add_paragraph()
    simple = OxmlElement("w:fldSimple")
    simple.set(qn("w:instr"), ' HYPERLINK "https://simple.example" ')
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "简单域链接"
    run.append(t)
    simple.append(run)
    para._p.append(simple)
    src = tmp_path / "a.docx"
    doc.save(str(src))

    md = import_from_docx(str(src))

    assert "[简单域链接](https://simple.example)" in md, md


# ==================== 3. 段落文字读取：含嵌套内容 ====================

def test_paragraph_text_sees_nested_runs(tmp_path):
    """
    `_paragraph_text` 必须能看到**内容控件 `w:sdt`** 里的 run。

    ⚠️ 我第一版拿"超链接"当前置条件，写错了：这个 python-docx 版本的
    `para.text` **能**看到 `w:hyperlink` 里的文字（本轮早先已实测过）。
    真正看不到的是 `w:sdt` 里的内容 —— 而那正是页脚页码被漏读的直接原因。
    """
    doc = Document()
    para = doc.add_paragraph()

    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "控件里的文字"
    run.append(t)
    content.append(run)
    sdt.append(content)
    para._p.append(sdt)

    assert para.text == "", "前置条件：p.text 看不到 w:sdt 里的文字"
    assert _paragraph_text(para) == "控件里的文字"


def test_paragraph_text_ignores_field_instructions(tmp_path):
    """域指令（PAGE/HYPERLINK…）本身不是可见文字，不应混进段落文字。"""
    src = _make_docx_with_hyperlink_field(tmp_path / "a.docx")

    text = _paragraph_text(Document(str(src)).paragraphs[0])

    assert text == "点这里", f"域指令混进来了: {text!r}"
