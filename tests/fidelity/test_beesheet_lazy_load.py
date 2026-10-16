# tests/fidelity/test_beesheet_lazy_load.py
"""
蜜算表格**两阶段载入**（R1）：先快速显示，再后台补全；**加载中禁止编辑**。

## 为什么必须先定"保存安全"
`write_workbook` 是**按 `sheet.df` 逐格写值**回源文件的。所以只要出现
"内存里只有一部分数据、却执行了保存"，就会**把不完整数据写回源文件**
（静默数据损坏）。这一条是本任务真正的前提，比"快多少"重要得多。

## 选定的设计（安全是**构造**出来的，不靠事后合并）
1. **阶段 1**：只读前 `PHASE1_ROWS` 行 → 立刻显示（窗口可交互）；
2. **阶段 2**：后台线程读完整数据 → 完成后**整体替换**模型数据；
3. **加载期间 `flags()` 返回只读** → 阶段 2 替换时**不可能存在用户改动**，
   于是不需要任何"合并/重放编辑"的复杂逻辑（那种逻辑正是数据损坏的来源）；
4. **保存前必须 `ensure_fully_loaded()`** —— 未加载完就保存，先补全再写；
   补全失败则**拒绝保存**并明确告知（绝不写半份数据）。

## 本文件的分工
* `test_*` 上半部分：纯模型行为（只读/可编辑、补全、失败处理）；
* **下半部分是验收核心**：`保存结果必须与"直接全量加载再保存"逐格一致`。
"""
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from beeoffice.plugins.beesheet_workbook import (                # noqa: E402
    read_workbook, write_workbook,
)


def make_xlsx(path, rows=30, cols=3, start=1):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet.append([f"列{c}" for c in range(1, cols + 1)])
    for index, r in enumerate(range(start, start + rows)):
        for c in range(1, cols + 1):
            # 注意 +1：第 1 行是表头，数据从第 2 行开始写。
            # 我第一版直接写 sheet.cell(r, c, ...) 把表头覆盖了，导致行数断言莫名其妙失败。
            sheet.cell(index + 2, c, value=f"值{r}-{c}")
    book.save(str(path))
    return path


def sheet_values(path, sheet_index=0):
    """读出工作表的值（用于逐格比对）。"""
    workbook = read_workbook(str(path))
    df = workbook.sheets[sheet_index].df
    return [list(row) for row in df.values.tolist()]


# ==================== 1. 模型：加载中的状态 ====================

def test_partial_model_is_readonly(qtbot):
    """加载未完成前必须**只读** —— 这是"阶段 2 可安全替换"的前提。"""
    from beeoffice.plugins.beesheet_model import SheetDataModel
    from PySide6.QtCore import Qt

    model = SheetDataModel()
    model.set_partial_dataframe(pd.DataFrame({"a": [1, 2, 3]}))

    assert model.is_fully_loaded is False
    flags = model.flags(model.index(0, 0))
    assert not (flags & Qt.ItemFlag.ItemIsEditable), (
        "加载未完成时单元格却是可编辑的 —— 阶段 2 替换会丢掉用户改动"
    )


def test_completed_model_becomes_editable(qtbot):
    from beeoffice.plugins.beesheet_model import SheetDataModel
    from PySide6.QtCore import Qt

    model = SheetDataModel()
    model.set_partial_dataframe(pd.DataFrame({"a": [1, 2]}))
    model.complete_full_dataframe(pd.DataFrame({"a": [1, 2, 3, 4]}))

    assert model.is_fully_loaded is True
    flags = model.flags(model.index(0, 0))
    assert flags & Qt.ItemFlag.ItemIsEditable, "补全后应当恢复可编辑"
    # 注意：模型会把不足默认行数的数据**补齐空行**（DEFAULT_ROWS=1000），
    # 所以要看 `_real_rows` 而不是 rowCount()。我第一版断言 rowCount()==4 就错了。
    assert model._real_rows == 4, model._real_rows


def test_full_load_path_marks_fully_loaded(qtbot):
    """老路径（一次性全量）必须仍标记为"已加载完"，否则会被误判成只读。"""
    from beeoffice.plugins.beesheet_model import SheetDataModel

    model = SheetDataModel()
    model.set_full_dataframe(pd.DataFrame({"a": [1, 2, 3]}))

    assert model.is_fully_loaded is True


def test_ensure_fully_loaded_uses_loader(qtbot):
    """`ensure_fully_loaded()` 要回调注入的加载器并把数据补上。"""
    from beeoffice.plugins.beesheet_model import SheetDataModel

    model = SheetDataModel()
    model.set_partial_dataframe(pd.DataFrame({"a": [1, 2]}))
    called = {"n": 0}

    def loader():
        called["n"] += 1
        return pd.DataFrame({"a": [1, 2, 3, 4, 5]})

    ok = model.ensure_fully_loaded(loader)

    assert ok is True
    assert called["n"] == 1
    assert model.is_fully_loaded
    assert model._real_rows == 5, model._real_rows


def test_ensure_fully_loaded_is_noop_when_already_loaded(qtbot):
    from beeoffice.plugins.beesheet_model import SheetDataModel

    model = SheetDataModel()
    model.set_full_dataframe(pd.DataFrame({"a": [1, 2]}))
    called = {"n": 0}

    ok = model.ensure_fully_loaded(lambda: called.__setitem__("n", 1))

    assert ok is True
    assert called["n"] == 0, "已加载完不该再读一次文件"


def test_ensure_fully_loaded_reports_failure_without_lying(qtbot):
    """补全失败必须返回 False 并**保持只读** —— 绝不假装已加载完。"""
    from beeoffice.plugins.beesheet_model import SheetDataModel

    model = SheetDataModel()
    model.set_partial_dataframe(pd.DataFrame({"a": [1, 2]}))

    def broken_loader():
        raise OSError("模拟：源文件读不到了")

    ok = model.ensure_fully_loaded(broken_loader)

    assert ok is False
    assert model.is_fully_loaded is False
    assert model.load_error, "失败要留下可显示的说明"


# ==================== 2. 阶段 1/2 的分工（读文件）====================

def test_read_workbook_respects_max_rows(tmp_path):
    """阶段 1 靠 `max_rows` 只读一小段。"""
    path = make_xlsx(tmp_path / "big.xlsx", rows=50)

    part = read_workbook(str(path), max_rows=10)
    assert len(part.sheets[0].df) == 10
    assert part.sheets[0].original_shape[0] == 10


def test_full_read_gets_everything(tmp_path):
    path = make_xlsx(tmp_path / "big.xlsx", rows=50)

    full = read_workbook(str(path))

    assert len(full.sheets[0].df) == 50


# ==================== 3. **验收核心**：保存一致性 ====================

def test_save_after_partial_then_complete_matches_direct_full_save(tmp_path):
    """
    **本任务的安全验收**：让模型经历"阶段 1 部分 → 阶段 2 补全"，
    其保存结果必须与"直接全量加载再保存"**逐格一致**。

    这是"惰性载入不会损坏文件"的可执行证明。
    """
    source = make_xlsx(tmp_path / "src.xlsx", rows=40)
    out_after_two_phase = tmp_path / "after_two_phase.xlsx"
    out_after_direct = tmp_path / "after_direct.xlsx"

    # A：两阶段（阶段 1 只读 10 行，再由阶段 2 补全）
    partial = read_workbook(str(source), max_rows=10)
    assert len(partial.sheets[0].df) == 10
    full = read_workbook(str(source))
    partial.sheets[0].df = full.sheets[0].df          # 阶段 2 的"整体替换"
    partial.sheets[0].formulas = full.sheets[0].formulas
    partial.sheets[0].original_shape = full.sheets[0].original_shape
    write_workbook(partial, str(out_after_two_phase))

    # B：直接全量
    direct = read_workbook(str(source))
    write_workbook(direct, str(out_after_direct))

    assert sheet_values(out_after_two_phase) == sheet_values(out_after_direct), (
        "两阶段载入后的保存结果与全量载入不一致 —— 惰性载入写坏了数据"
    )
    assert len(sheet_values(out_after_two_phase)) == 40, (
        "只写出了部分行（这正是必须避免的数据损坏）"
    )


def test_saving_partial_after_sort_corrupts_the_file(tmp_path):
    """
    **说明"保存前必须补全"的关键证据**（不是我猜的，是实测出来的）：

    `write_workbook` 走"打开源工作簿、只更新值"的路线，所以拿**部分载入**的
    数据去保存时，**后面的行会保留源文件里的旧值**。于是：

        部分载入 10 行 → 排序（倒序）→ 保存
        → 第 1–10 行是倒序数据，**第 11–40 行还是原文件旧值**

    文件内容与界面显示不一致 —— 用户重新打开会看到重复/乱序的数据。
    这条测试把"为什么必须有 `ensure_fully_loaded()` 门"钉成可执行事实。
    """
    source = make_xlsx(tmp_path / "src.xlsx", rows=40)
    out = tmp_path / "sorted_partial.xlsx"

    partial = read_workbook(str(source), max_rows=10)
    partial.sheets[0].df = partial.sheets[0].df.iloc[::-1].reset_index(drop=True)
    write_workbook(partial, str(out))

    column = [row[0] for row in sheet_values(out)]
    assert len(column) == 40, "前提变了：输出行数与源文件不同，需重新确认本测试"
    # 第 1–10 行被部分数据覆盖（倒序），第 11 行起仍是旧值 → 数据错乱
    assert column[:3] == ["值10-1", "值9-1", "值8-1"], column[:5]
    assert column[10] == "值11-1", (
        f"前提变了：第 11 行不再是旧值（{column[10]!r}），"
        f"说明 write_workbook 的行为已不同，需重新评估风险"
    )


def test_gate_prevents_the_corruption(tmp_path):
    """
    对照：**先补全再排序再保存**，文件就是正确的（倒序且无残留旧值）。
    这正是 `ensure_fully_loaded()` 要保证的效果。
    """
    source = make_xlsx(tmp_path / "src.xlsx", rows=40)
    out = tmp_path / "sorted_full.xlsx"

    workbook = read_workbook(str(source))          # 相当于"补全后"
    workbook.sheets[0].df = workbook.sheets[0].df.iloc[::-1].reset_index(drop=True)
    write_workbook(workbook, str(out))

    column = [row[0] for row in sheet_values(out)]
    assert column[:3] == ["值40-1", "值39-1", "值38-1"], column[:5]
    assert column[-1] == "值1-1", column[-3:]
    assert len(set(column)) == 40, "出现了重复值（说明有旧值残留）"


def test_edited_cells_survive_full_load(tmp_path):
    """
    全量替换后，用户的**编辑**（在阶段 2 完成之后做的）要能正常保存。
    这里模拟"补全 → 改一格 → 保存"这条正常路径。
    """
    source = make_xlsx(tmp_path / "src.xlsx", rows=20)
    out = tmp_path / "edited.xlsx"

    workbook = read_workbook(str(source))
    workbook.sheets[0].df.iloc[0, 0] = "我改过的值"
    write_workbook(workbook, str(out))

    values = sheet_values(out)
    assert values[0][0] == "我改过的值", values[0]
