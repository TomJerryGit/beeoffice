# tests/fidelity/test_fidelity_report.py
"""
P1-A：**导出前保真度提示**。

为什么要有它
------------
有一批 OOXML 构造在导入→导出之间**必然**丢失或降级（脚注、批注、文本框、
公式、嵌套表格、宏…），而以往是**静默**发生的 —— 用户要等导出后自己发现。
竞品普遍也不提示。

本文件的两条验收线
------------------
1. **不能漏报**：每一项已实测会丢的构造都要被检出（用 `_docx_features.py`
   造的夹具逐项验证）。
2. **不能误报**：对我们已验证"往返零丢失"的真实文档，检测器**一个字都不能报**
   —— 报了就是吓唬用户。（第一版就是在这里翻车的：按"部件存在"判定脚注，
   而**每个 DOCX 默认都带 footnotes.xml/endnotes.xml**，于是 7 份真实文档里
   3 份被误报。）
"""
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests" / "fidelity"))

from _docx_features import build_feature_docx        # noqa: E402

from beeoffice.plugins.fidelity_report import (                # noqa: E402
    HIGH, MEDIUM, FidelityReport, LossItem, analyze_file,
    format_report, should_prompt_before_export,
)

REAL_DOCS = sorted((ROOT / "tests" / "fidelity" / "corpus").glob("*.docx"))
TEST_FILES = sorted(
    p for p in (ROOT / "test+files").iterdir()
    if p.suffix.lower() in (".docx", ".xlsx", ".pptx")
) if (ROOT / "test+files").is_dir() else []


# ==================== 1. 不能误报 ====================

@pytest.mark.parametrize("path", REAL_DOCS + TEST_FILES,
                         ids=lambda p: p.name[:40])
def test_verified_lossless_documents_produce_no_warning(path):
    """
    **最重要的一条**：已实测零丢失的真实文档，检测器必须**零输出**。

    这些文件都是用户提供的真实教程，我们逐项对比过往返结果。
    如果检测器在这里报东西，说明规则有误报 —— 会直接损害提示的可信度。
    """
    report = analyze_file(path)

    assert not report.items, (
        f"{path.name} 被误报: "
        f"{[(i.key, i.count) for i in report.items]}"
    )
    assert not report.note, f"{path.name} 被误报说明文字: {report.note}"


def test_plain_docx_without_any_feature_is_clean(tmp_path):
    """连"普通文档"都不该报警。"""
    from docx import Document

    path = tmp_path / "plain.docx"
    doc = Document()
    doc.add_heading("标题", level=1)
    doc.add_paragraph("普通正文")
    doc.add_table(rows=2, cols=2)
    doc.save(str(path))

    assert analyze_file(path).items == []


# ==================== 2. 不能漏报（逐项夹具）====================

@pytest.mark.parametrize("feature,expected_key", [
    ("textbox", "textbox"),
    ("equation", "equation"),
    ("footnote", "footnote"),
    ("comment", "comment"),
    ("legacy_vml", "legacy_graphics"),
])
def test_content_losing_features_are_detected(tmp_path, feature, expected_key):
    """这些是**内容真会丢**的（high），必须检出。"""
    path = tmp_path / "f.docx"
    build_feature_docx(path, only=[feature])

    report = analyze_file(path)
    keys = {i.key: i for i in report.items}

    assert expected_key in keys, f"{feature} 没被检出: {sorted(keys)}"
    assert keys[expected_key].severity == HIGH
    assert keys[expected_key].count >= 1


@pytest.mark.parametrize("feature,expected_key", [
    ("toc_field", "toc_field"),
    ("content_control", "content_control"),
    ("tracked_change", "revision"),
    # R4（2026-10-14）之后嵌套表格**降级为 medium**：结构与文字已能往返保真
    # （见 test_docx_nested_tables.py），只有内层格式不保 ——
    # 继续把它说成"整块丢失"就是**误报**，会吓到用户。
    ("nested_table", "nested_table"),
])
def test_capability_degradations_are_detected_as_medium(
        tmp_path, feature, expected_key):
    """
    这些**文字/结构会保留**，只是格式或能力丢了（medium）—— 必须检出，但**不该弹窗**。
    """
    path = tmp_path / "f.docx"
    build_feature_docx(path, only=[feature])

    report = analyze_file(path)
    keys = {i.key: i for i in report.items}

    assert expected_key in keys, f"{feature} 没被检出: {sorted(keys)}"
    assert keys[expected_key].severity == MEDIUM
    assert not report.has_high_losses
    assert not should_prompt_before_export(report), (
        "只是能力降级（文字都在）不该打断用户导出"
    )


def test_all_features_together_are_all_reported(tmp_path):
    """一个文档里塞满各构造时，应尽量都报出来（而不是只报第一个）。"""
    path = tmp_path / "all.docx"
    build_feature_docx(path)

    keys = {i.key for i in analyze_file(path).items}

    assert {"textbox", "equation", "nested_table", "legacy_graphics",
            "toc_field", "content_control", "revision"} <= keys, sorted(keys)


def test_nested_table_wording_matches_reality(tmp_path):
    """
    R4 之后，嵌套表格的**措辞必须与实现一致**：结构与文字保留、只有格式不保。

    这条是"文档别撒谎"的守卫：如果哪天 R4 的实现回退了，
    要么改回 HIGH，要么这条测试会提醒"措辞与事实不符"。
    """
    path = tmp_path / "f.docx"
    build_feature_docx(path, only=["nested_table"])

    item = next(i for i in analyze_file(path).items
                if i.key == "nested_table")

    assert item.severity == MEDIUM, "嵌套表格现在不该再按'内容丢失'报"
    assert "结构" in item.effect and "文字" in item.effect, item.effect
    assert "不保" in item.effect or "不保留" in item.effect, item.effect


def test_footnote_reference_count_matches_document(tmp_path):
    """计数要跟着实际引用数走（不是"有部件就报 1"）。"""
    path = tmp_path / "f.docx"
    build_feature_docx(path, only=["footnote"])

    item = next(i for i in analyze_file(path).items if i.key == "footnote")

    assert item.count == 1


# ==================== 3. 其它格式与异常输入 ====================

def test_legacy_format_gets_explanation_not_crash(tmp_path):
    path = tmp_path / "old.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)

    report = analyze_file(path)

    assert report.note, "旧格式应给出说明"
    assert should_prompt_before_export(report), "旧格式应提示用户"


def test_pptx_gets_no_import_explanation(tmp_path):
    import zipfile

    path = tmp_path / "deck.pptx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("ppt/presentation.xml", "<x/>")

    report = analyze_file(path)

    assert "PPTX" in report.note or "pptx" in report.note


def test_missing_file_does_not_raise(tmp_path):
    """分析失败必须静默降级，绝不能因为它让打开/导出失败。"""
    report = analyze_file(tmp_path / "不存在.docx")

    assert report.items == []


def test_corrupt_zip_does_not_raise(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"this is not a zip")

    report = analyze_file(path)

    assert report.items == []
    assert report.note


def test_xlsx_macro_is_reported(tmp_path):
    """蜜算：宏（.xlsm）应提示"文件里保留但蜜算不显示/不能运行"。"""
    import openpyxl

    path = tmp_path / "book.xlsx"
    openpyxl.Workbook().save(str(path))
    # 手工塞一个 vbaProject.bin 进去
    import zipfile

    temp = tmp_path / "tmp.xlsx"
    with zipfile.ZipFile(path) as src, \
            zipfile.ZipFile(temp, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/vbaProject.bin", b"fake vba")
    temp.replace(path)

    report = analyze_file(path)

    keys = {i.key for i in report.items}
    assert "macro" in keys, sorted(keys)
    assert not report.has_high_losses, "宏在文件里是保留的，不该算内容丢失"


# ==================== 4. 提示文本与"是否弹窗"的纯逻辑 ====================

def test_should_prompt_only_for_high_or_note():
    assert not should_prompt_before_export(None)
    assert not should_prompt_before_export(FidelityReport())

    medium = FidelityReport(items=[
        LossItem("toc_field", "目录域", 1, "文字保留", MEDIUM)])
    assert not should_prompt_before_export(medium)

    high = FidelityReport(items=[
        LossItem("footnote", "脚注", 2, "会丢失", HIGH)])
    assert should_prompt_before_export(high)

    noted = FidelityReport(note="加密文档")
    assert should_prompt_before_export(noted)


def test_format_report_lists_high_first():
    report = FidelityReport(path="x.docx", items=[
        LossItem("toc_field", "目录域", 1, "文字保留", MEDIUM),
        LossItem("footnote", "脚注", 2, "内容与标记都不会保留", HIGH),
    ])

    text = format_report(report, "x.docx")

    assert "脚注" in text and "目录域" in text
    assert text.index("脚注") < text.index("目录域"), "内容丢失的应排在前面"
    assert "2 处" in text
    assert "x.docx" in text


def test_format_report_says_clean_when_no_losses():
    text = format_report(FidelityReport(path="a.docx"), "a.docx")
    assert "未发现" in text


def test_summary_mentions_counts():
    report = FidelityReport(items=[
        LossItem("footnote", "脚注", 1, "x", HIGH),
        LossItem("toc_field", "目录域", 1, "y", MEDIUM),
    ])
    summary = report.summary()
    assert "1 类" in summary or "1" in summary


# ==================== 5. UI 接线：守卫要真的拦得住 ====================

def _editor(qtbot):
    from beeoffice.plugins.beewrite import BeeWriteEditor

    widget = BeeWriteEditor()
    qtbot.addWidget(widget)
    return widget


def test_export_guard_skips_dialog_when_no_report(qtbot, monkeypatch):
    """没有报告（普通文档）时不该弹窗，直接放行。"""
    from PySide6.QtWidgets import QMessageBox

    called = {"n": 0}
    monkeypatch.setattr(
        QMessageBox, "warning",
        staticmethod(lambda *a, **k: called.__setitem__("n", called["n"] + 1)),
    )

    editor = _editor(qtbot)
    assert editor._confirm_fidelity_losses("导出") is True
    assert called["n"] == 0, "无损失时不该弹窗"


def test_export_guard_skips_dialog_for_medium_only(qtbot, monkeypatch):
    """只有"能力降级"（文字都在）时不打断用户。"""
    from PySide6.QtWidgets import QMessageBox

    called = {"n": 0}
    monkeypatch.setattr(
        QMessageBox, "warning",
        staticmethod(lambda *a, **k: called.__setitem__("n", called["n"] + 1)),
    )

    editor = _editor(qtbot)
    editor._fidelity_report = FidelityReport(path="a.docx", items=[
        LossItem("toc_field", "目录域", 1, "文字保留，不能更新", MEDIUM)])

    assert editor._confirm_fidelity_losses("导出") is True
    assert called["n"] == 0


def test_export_guard_prompts_and_can_be_cancelled(qtbot, monkeypatch):
    """
    **关键接线**：有内容会丢时必须弹窗；用户点"否"要真的拦住导出。
    """
    from PySide6.QtWidgets import QMessageBox

    seen = {}

    def fake_warning(parent, title, text, *args, **kwargs):
        seen["title"] = title
        seen["text"] = text
        return QMessageBox.No

    monkeypatch.setattr(QMessageBox, "warning", staticmethod(fake_warning))

    editor = _editor(qtbot)
    editor._fidelity_report = FidelityReport(path="有脚注.docx", items=[
        LossItem("footnote", "脚注", 2, "脚注内容与标记都不会保留", HIGH)])

    allowed = editor._confirm_fidelity_losses("导出")

    assert allowed is False, "用户点了\"否\"，导出必须被拦住"
    assert "脚注" in seen["text"], seen
    assert "2 处" in seen["text"], seen
    assert "有脚注.docx" in seen["text"], seen


def test_export_guard_allows_when_user_confirms(qtbot, monkeypatch):
    from PySide6.QtWidgets import QMessageBox

    monkeypatch.setattr(
        QMessageBox, "warning",
        staticmethod(lambda *a, **k: QMessageBox.Yes),
    )

    editor = _editor(qtbot)
    editor._fidelity_report = FidelityReport(path="a.docx", items=[
        LossItem("nested_table", "嵌套表格", 1, "内层会丢", HIGH)])

    assert editor._confirm_fidelity_losses("导出") is True


def test_open_attaches_report_to_widget(qtbot, tmp_path, monkeypatch):
    """打开文件时要把报告挂到文档对象上（导出前才有东西可提示）。"""
    from beeoffice.app import BeeApp

    path = tmp_path / "f.docx"
    build_feature_docx(path, only=["footnote"])

    editor = _editor(qtbot)
    window = BeeApp.__new__(BeeApp)          # 不走 __init__（避免起整套 UI）
    # _attach_fidelity_report 只用到 set_global_status，这里给个空实现
    calls = {}
    window.set_global_status = lambda msg, ms=0: calls.setdefault("msg", msg)

    window._attach_fidelity_report(editor, path)

    report = getattr(editor, "_fidelity_report", None)
    assert report is not None and report.has_high_losses
    assert "脚注" in calls.get("msg", "") or "无法保留" in calls.get("msg", "")
