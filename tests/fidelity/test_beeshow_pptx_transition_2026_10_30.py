"""蜜演阶段 B：PPTX 页面切换效果导出。"""
from pptx import Presentation

from beeoffice.plugins.beeshow_pptx_engine import add_slide_transition


def _slide():
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def test_add_fade_transition():
    _, slide = _slide()
    assert add_slide_transition(slide, "fade") is True
    xml = slide._element.xml
    assert "p:transition" in xml
    assert "p:fade" in xml


def test_add_push_cover_wipe():
    for key, tag in (("push", "p:push"), ("cover", "p:cover"),
                     ("wipe", "p:wipe")):
        _, slide = _slide()
        assert add_slide_transition(slide, key) is True
        assert tag in slide._element.xml


def test_none_transition_is_noop():
    _, slide = _slide()
    assert add_slide_transition(slide, "none") is False
    assert "p:transition" not in slide._element.xml
