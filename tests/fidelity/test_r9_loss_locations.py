# tests/fidelity/test_r9_loss_locations.py
"""
R9 兼容性收口：**损失提示要精确到"哪一处"**。

## 为什么
原先只说"脚注 2 处"——用户既不知道去文档哪一页找，也没法判断该不该处理。
R9 的定位是"**让不兼容从'打开后才发现'变成'打开前就知道'**"，
所以每一项都要尽量给出**位置线索**（第几段 + 该段文字片段）。

## 本文件守的
1. 能定位的构造给出"第 N 段：『片段』"；
2. 位置**有上限**（报告不能变成几百行），超出以"…等 N 处"概括；
3. 定位不到的（包级别的宏/图表）**不编造位置**；
4. 位置要出现在**对话框明细**与**导出的报告文本**里（用户两条路都走得到）。
"""
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests" / "fidelity"))

from _docx_features import build_feature_docx                     # noqa: E402

from beeoffice.plugins.fidelity_report import (                    # noqa: E402
    LOCATION_LIMIT, SNIPPET_CHARS, HIGH, FidelityReport, LossItem,
    analyze_file, format_report,
)


def _item(path, key):
    return next(i for i in analyze_file(path).items if i.key == key)


# ==================== 1. 能定位的构造要给出位置 ====================

def test_footnote_loss_reports_location(tmp_path):
    path = tmp_path / "脚注.docx"
    build_feature_docx(path, only=["footnote"])

    item = _item(path, "footnote")

    assert item.locations, "脚注没给出位置"
    assert any("段" in text for text in item.locations), item.locations


def test_textbox_loss_reports_location(tmp_path):
    path = tmp_path / "文本框.docx"
    build_feature_docx(path, only=["textbox"])

    item = _item(path, "textbox")

    assert item.locations, f"文本框没给出位置：{item.locations}"
    assert any("第" in text and "段" in text for text in item.locations)


def test_equation_loss_reports_location(tmp_path):
    path = tmp_path / "公式.docx"
    build_feature_docx(path, only=["equation"])

    assert _item(path, "equation").locations


def test_location_includes_text_snippet_when_available(tmp_path):
    """位置要带**该段文字片段**，光有段号对用户帮助有限。"""
    path = tmp_path / "文本框.docx"
    build_feature_docx(path, only=["textbox"])

    locations = _item(path, "textbox").locations

    assert any("『" in text and "』" in text for text in locations), locations


def test_snippet_is_truncated(tmp_path):
    """片段要截断，否则长段落会把报告撑爆。"""
    from docx import Document

    from beeoffice.plugins.fidelity_report import _paragraph_snippet

    doc = Document()
    paragraph = doc.add_paragraph("很长的正文内容" * 20)

    snippet = _paragraph_snippet(paragraph._p)

    assert len(snippet) <= SNIPPET_CHARS + 1, snippet
    assert snippet.endswith("…"), snippet


# ==================== 2. 上限与概括 ====================

def test_locations_are_capped_in_display():
    item = LossItem("k", "测试项", 20, "效果", HIGH,
                    locations=[f"第 {i} 段" for i in range(1, 21)])

    lines = item.location_lines()

    assert len(lines) == LOCATION_LIMIT + 1, lines
    assert any("…等 20 处" in line for line in lines), lines


def test_locations_within_limit_are_all_shown():
    item = LossItem("k", "测试项", 2, "效果", HIGH,
                    locations=["第 1 段", "第 2 段"])

    lines = item.location_lines()

    assert len(lines) == 2
    assert not any("等" in line for line in lines)


def test_no_locations_means_no_extra_lines():
    """定位不到的项（例如宏）**不编造位置**。"""
    item = LossItem("macro", "宏（VBA）", 1, "宏代码不会保留")

    assert item.location_lines() == []


def test_macro_loss_has_no_fabricated_location(tmp_path):
    """宏是包级部件，定位不到段落 —— 不许硬报一个位置。"""
    path = tmp_path / "宏.docx"
    from docx import Document

    doc = Document()
    doc.add_paragraph("正文")
    doc.save(str(path))
    # 手工塞一个 vbaProject.bin（只需要"文件名在包里"）
    import zipfile

    temp = tmp_path / "with_macro.docx"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(temp, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/vbaProject.bin", b"\x00")
    temp.replace(path)

    item = _item(path, "macro")

    assert item.count == 1
    assert item.locations == [], "宏没有段落位置，不该编造"


def test_deep_locate_stops_at_limit(tmp_path):
    """
    段落很多时不能在定位上花太久 —— `_locate` 到上限就停。

    （这条同时也是性能约束：报告是在**打开文件**时算的。）
    """
    from beeoffice.plugins.fidelity_report import _locate

    path = tmp_path / "多个文本框.docx"
    build_feature_docx(path, only=["textbox"])
    from lxml import etree
    import zipfile

    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    found = _locate(root, f"{W}txbxContent", limit=1)

    assert len(found) == 1


# ==================== 3. 两条用户可见路径都要有位置 ====================

def test_detail_lines_include_locations(tmp_path):
    path = tmp_path / "文本框.docx"
    build_feature_docx(path, only=["textbox"])

    lines = analyze_file(path).detail_lines()

    assert any("『" in line or "段" in line for line in lines), lines


def test_exported_report_text_includes_locations(tmp_path):
    """**导出的报告文件**也要带位置（R4 加的写文件路径）。"""
    from beeoffice.plugins.fidelity_report import write_report_file

    path = tmp_path / "文本框.docx"
    build_feature_docx(path, only=["textbox"])
    report = analyze_file(path)
    target = tmp_path / "报告.md"

    assert write_report_file(report, target) is True
    text = target.read_text(encoding="utf-8")

    assert "文本框" in text
    assert "段" in text, text


def test_report_without_locations_still_formats():
    """没有位置信息的历史报告仍要能正常渲染（向后兼容）。"""
    report = FidelityReport(path="a.docx", items=[
        LossItem("macro", "宏（VBA）", 1, "宏代码不会保留", HIGH),
    ])

    text = format_report(report)

    assert "宏（VBA）" in text
    assert "无法完整保留" in text
