# tests/fidelity/test_beesheet_number_formats_4_4.py
"""
蜜算数字格式 4.4 回归（2026-10-22）：万元 / 正负号。

- 万元：界面显示 `12.35万`（只改显示，**原始数值不变**），xlsx 写回带"万"的
  自定义格式；正负号：显示 `+1,234.50` / `-1,234.50`，xlsx 写回 `+...;-...`。
- 这两个格式是矩阵 B 的 S 级项：不改变数据，只补格式分支。
"""
import pandas as pd
from openpyxl import load_workbook

from beeoffice.plugins.beesheet import BeeSheetView
from beeoffice.plugins.beesheet_style_export import number_format_code


def test_number_format_codes():
    assert "万" in number_format_code("万元", 2)
    assert number_format_code("正负号", 2).count(";") == 1
    assert number_format_code("正负号", 2).startswith("+")
    assert number_format_code("正负号", 2).split(";")[1].startswith("-")


def test_view_displays_wan_and_sign(qtbot):
    view = BeeSheetView()
    qtbot.addWidget(view)
    view.model.set_full_dataframe(pd.DataFrame({"金额": [123456.0, -1234.5]}))
    view._cell_formats = {
        (0, 0): {"number_format": "万元", "decimals": 2},
        (1, 0): {"number_format": "正负号", "decimals": 2},
    }
    view.model._cell_formats_ref = view._cell_formats
    assert view.model.data(view.model.index(0, 0)) == "12.35万"
    assert view.model.data(view.model.index(1, 0)) == "-1,234.50"


def test_wan_and_sign_written_to_xlsx_without_changing_value(qtbot, tmp_path):
    view = BeeSheetView()
    qtbot.addWidget(view)
    view.model.set_full_dataframe(pd.DataFrame({"金额": [123456.0, -1234.5]}))
    view._cell_formats = {
        (0, 0): {"number_format": "万元", "decimals": 2},
        (1, 0): {"number_format": "正负号", "decimals": 2},
    }
    path = tmp_path / "formats.xlsx"
    assert view._write_xlsx(str(path)) is True

    wb = load_workbook(path)
    ws = wb.active
    wan_format = ws["A2"].number_format
    sign_format = ws["A3"].number_format
    assert "万" in wan_format, wan_format
    assert "+" in sign_format and ";" in sign_format, sign_format
    # 关键：格式只改显示，原始数值不能被除 10000
    assert ws["A2"].value == 123456
    assert ws["A3"].value == -1234.5
