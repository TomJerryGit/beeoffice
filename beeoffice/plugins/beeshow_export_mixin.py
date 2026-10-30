# beeoffice/plugins/beeshow_export_mixin.py
"""
蜜演导出 Mixin

职责：
  - _export_pptx         : 导出 PPTX（任务 3.1：原生元素映射 + 任务 5.1 动态数字预处理）
  - _export_pptx_legacy  : PPTX 保底方案（BeautifulSoup 解析）
  - _export_pdf          : 导出 PDF（P0-4：保持 QWebEngineView 引用 +
                           任务 5.3 视频封面替换 + 任务 5.1 动态数字预处理）
  - _export_html         : 导出 HTML（任务 5.3 复制视频 + 任务 5.1 注入动态数字 JS）
  - _export_images       : 导出 PNG 图片序列（P0-4）
  - _export_gif          : 导出 GIF 动图（依赖 bee_export）
  - _record_presentation : 录制占位

v5.1 增量汇总：
  - 任务 3.1：_export_pptx 改用 PPTXEngine 做原生元素映射；
             保留 _export_pptx_legacy 作为保底
  - 任务 5.1：_export_pdf / _export_pptx 前处理动态数字；
             _export_html 注入动态数字 JS
  - 任务 5.3：_export_html 复制视频文件；
             _export_pdf 把视频替换为封面图

v5.1.1 修复：
  - QTimer 从 PySide6.QtWidgets 改为 PySide6.QtCore
    （QTimer 属于 QtCore，不在 QtWidgets；原 import 会抛 ImportError）

修复记录：
  - 补 `import logging`：
      原 beeshow.py 的 _export_images 里用了 logging.error 但未导入 logging，
      实际运行时若截图失败会抛 NameError。

依赖：
  - logging / pathlib.Path
  - PySide6.QtWidgets（QFileDialog / QMessageBox / QInputDialog）
  - PySide6.QtCore（QTimer）
  - PySide6.QtWebEngineWidgets（QWebEngineView）
  - beeshow_constants.SLIDE_THEMES
  - bee_export.export_slides_to_gif（延迟 import）
  - beeshow_pptx_engine.PPTXEngine（延迟 import，任务 3.1）

注意：
  - 不导入其他 beeshow_* Mixin，避免循环依赖
  - P0-4 关键修复：_export_views 保持 QWebEngineView 引用，避免被 GC
"""

import logging
from pathlib import Path

from beeoffice.plugins.beeshow_html import render_slide_html

from PySide6.QtWidgets import (
    QFileDialog, QMessageBox, QInputDialog,
)
from PySide6.QtCore import QTimer
from PySide6.QtWebEngineWidgets import QWebEngineView

from beeoffice.plugins.beeshow_constants import SLIDE_THEMES


class ExportMixin:
    """蜜演导出 Mixin"""

    # ==================== 导出 PPTX（任务 3.1） ====================
    def _export_pptx(self):
        """
        导出为 PowerPoint（.pptx）。

        任务 3.1：改用 PPTXEngine 做原生元素映射。
        相比旧实现：
          - 文本保持原生可编辑、可搜索
          - 图片保持原生 Picture（可替换）
          - 表格保持原生 Table
          - 保留字号、颜色、加粗、对齐
          - 保留背景色 / 容器背景

        任务 5.1：导出前把 {{number:...}} 替换为最终值。

        降级：
          - 无 playwright 时，引擎自动降级为"仅文本"模式
            （不保真，但至少可编辑）
          - 引擎失败时，回退到旧的 BeautifulSoup 方案
        """
        if not self._slides_html:
            return

        path, _ = QFileDialog.getSaveFileName(
            self, "导出 PPTX", "演示.pptx", "PowerPoint (*.pptx)"
        )
        if not path:
            return

        theme = SLIDE_THEMES[self._current_theme]
        bg = self._custom_bg or theme.get("bg", "#FFF")

        # 背景色可能是渐变（linear-gradient(...)），提取纯色作为兜底
        if "gradient" in bg or "linear" in bg:
            bg_solid = "#FFFFFF"
        else:
            bg_solid = bg

        try:
            from beeoffice.plugins.beeshow_pptx_engine import PPTXEngine

            self._set_status("📄 正在导出 PPTX...", 0)

            # 任务 5.1：导出 PPTX 时把动态数字替换为最终值
            export_slides = []
            for s in self._slides_html:
                try:
                    s = self._render_dynamic_numbers_for_export(s)
                except AttributeError:
                    pass
                export_slides.append(s)

            engine = PPTXEngine(slide_width=10.0, slide_height=7.5)
            transitions = []
            for cfg in getattr(self, "_slide_animations", []):
                transitions.append((cfg or {}).get("transition", "none"))
            # 优化规划 §2.7：把文档所在目录作为相对路径图片的基准
            #（引擎会注入 <base href>）——否则幻灯片里的相对图片一律加载失败，
            # 导出的 PPTX 里图片整张消失。
            src_path = getattr(self, "_current_file_path", None)
            success = engine.export(
                slides_html=export_slides,
                output_path=path,
                theme_css=theme["css"],
                bg_color=bg_solid,
                base_dir=str(Path(src_path).parent) if src_path else None,
                transitions=transitions,
            )

            if success:
                QMessageBox.information(
                    self, "导出成功",
                    f"已保存: {path}\n\n"
                    f"✅ 文本可编辑、可搜索\n"
                    f"✅ 保留字号/颜色/对齐"
                )
                self._set_status("✅ PPTX 导出完成", 3000)
            else:
                QMessageBox.warning(
                    self, "导出失败",
                    "PPTX 引擎返回失败，尝试旧方案..."
                )
                self._export_pptx_legacy(path)

        except ImportError as e:
            logging.error(f"PPTX 引擎导入失败：{e}")
            QMessageBox.warning(
                self, "缺少依赖",
                f"PPTX 引擎不可用：\n{e}\n\n"
                f"请确认 python-pptx 已安装。\n"
                f"如需精确版式，请安装 playwright：\n"
                f"pip install playwright && playwright install chromium"
            )
            self._export_pptx_legacy(path)

        except Exception as e:
            logging.error(f"PPTX 导出失败：{e}")
            QMessageBox.warning(
                self, "导出失败", f"出错：{e}\n\n尝试旧方案..."
            )
            self._export_pptx_legacy(path)

    def _export_pptx_legacy(self, path: str):
        """
        旧的 PPTX 导出方案（保底）。

        - 基于 BeautifulSoup 解析 HTML
        - 提取标题和正文到占位符
        - 不保真，但依赖最少
        """
        try:
            from pptx import Presentation
            from pptx.util import Inches
            from bs4 import BeautifulSoup

            prs = Presentation()
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(7.5)

            for slide_html in self._slides_html:
                soup = BeautifulSoup(slide_html, 'html.parser')
                title_tag = soup.find(['h1', 'h2', 'h3'])
                title_text = (
                    title_tag.get_text().strip() if title_tag else ""
                )
                slide = prs.slides.add_slide(
                    prs.slide_layouts[1] if title_text
                    else prs.slide_layouts[0]
                )
                if title_text and slide.shapes.title:
                    slide.shapes.title.text = title_text

                content = []
                for elem in soup.find_all(['p', 'li']):
                    t = elem.get_text().strip()
                    if t and t != title_text:
                        prefix = '• ' if elem.name == 'li' else ''
                        content.append(f"{prefix}{t}")

                if content and len(slide.placeholders) > 1:
                    slide.placeholders[1].text = "\n".join(content)

            prs.save(path)
            QMessageBox.information(
                self, "导出成功（保底方案）",
                f"已保存: {path}\n\n"
                f"⚠️ 使用了保底方案，版式可能不精确。"
            )
        except Exception as e:
            logging.error(f"保底方案也失败：{e}")
            QMessageBox.warning(self, "导出失败", str(e))

    # ==================== 导出 PDF（P0-4 + 任务 5.3 + 任务 5.1） ====================
    def _build_pdf_html(self) -> str:
        """
        拼装"导出 PDF 用"的整份 HTML（视频换封面图、动态数字取终值、逐页分页）。

        抽出来是为了让「导出 PDF」的**全局入口**与蜜演自己的导出用同一套版面，
        避免两处逻辑漂移。
        """
        theme = SLIDE_THEMES[self._current_theme]
        bg = self._custom_bg or theme.get("bg", "#FFF")

        # 任务 5.3：视频替换为封面图
        slides_for_export = list(self._slides_html)
        try:
            slides_for_export = self._replace_videos_with_posters(
                slides_for_export
            )
        except AttributeError:
            pass

        parts = []
        for i, s in enumerate(slides_for_export):
            # 任务 5.1：动态数字替换为最终值
            try:
                s = self._render_dynamic_numbers_for_export(s)
            except AttributeError:
                pass

            pb = (
                '<div style="page-break-after:always"></div>'
                if i < len(slides_for_export) - 1 else ''
            )
            parts.append(
                f"<div class='slide' style='width:100%;min-height:100vh;"
                f"display:flex;flex-direction:column;"
                f"justify-content:center;align-items:center;"
                f"padding:60px;box-sizing:border-box'>{s}</div>{pb}"
            )
        slides_html = "".join(parts)

        # 任务 2.1：拼入版式 CSS
        # 优化规划 §3.3：骨架统一走 beeshow_html.render_slide_html()
        return render_slide_html(
            slides_html,
            theme_css=theme['css'],
            layout_css=self._get_layout_css(),
            bg=bg,
            extra_css="@page{size:A4 landscape;margin:0}body{margin:0}",
            wrap_slide=False,      # 每页已在上面单独包好 .slide 并带分页
        )

    def export_pdf(self, path: str) -> bool:
        """
        导出 PDF 到指定路径（供全局「导出 PDF」入口按模块分流时调用）。

        :return: **是否已开始导出**。QtWebEngine 的 `printToPdf` 是异步的，
            所以这里接上 `pdfPrintingFinished` 把**真实结果**回报给用户 ——
            原先的实现是"刚发起就弹『导出成功』"（审计报告指出：没等写完就报成功，
            失败时用户也会看到成功提示）。

        环境限制：本沙箱里 QtWebEngine 初始化即失败（`platform_channel.cc`），
        所以这条路径**无法在这里端到端验证**，只能保证结构与信号接线正确。
        """
        if not self._slides_html:
            return False

        full_html = self._build_pdf_html()
        temp = QWebEngineView()
        self._export_views.append(temp)

        def on_load(ok):
            if not ok:
                QMessageBox.warning(self, "导出失败", "幻灯片渲染失败")
                self._release_export_view(temp)
                return
            page = temp.page()
            page.pdfPrintingFinished.connect(
                lambda file_path, success: self._on_pdf_finished(
                    file_path, success, temp
                )
            )
            page.printToPdf(path)

        temp.page().loadFinished.connect(on_load)
        temp.setHtml(full_html)
        return True

    def _release_export_view(self, view):
        """清理临时导出视图的引用（避免被 GC 或残留）。"""
        try:
            if view in self._export_views:
                self._export_views.remove(view)
            view.deleteLater()
        except Exception:  # noqa: BLE001
            pass

    def _on_pdf_finished(self, file_path, success, view):
        """`pdfPrintingFinished` 回调：**按真实结果**提示用户。"""
        try:
            if success:
                QMessageBox.information(
                    self, "导出成功", f"已保存: {file_path}"
                )
            else:
                QMessageBox.warning(
                    self, "导出失败", f"PDF 写入失败：{file_path}"
                )
        finally:
            self._release_export_view(view)

    def _export_pdf(self):
        """
        导出 PDF（P0-4：保持 QWebEngineView 引用，避免被 GC）。

        任务 5.3：视频替换为封面图。
        任务 5.1：动态数字替换为最终值。

        流程：
          1. 拼接所有幻灯片为一份 HTML
          2. 创建临时 QWebEngineView（加入 self._export_views）
          3. loadFinished 后调 printToPdf
          4. 完成后清理引用

        关键：
          - 先 connect 再 setHtml，避免时序问题
          - 在 loadFinished 回调里做清理，避免提前 GC
        """
        if not self._slides_html:
            return

        path, _ = QFileDialog.getSaveFileName(
            self, "导出 PDF", "演示.pdf", "PDF (*.pdf)"
        )
        if not path:
            return

        # P0 修复：复用 `export_pdf` 的唯一实现（含"按真实结果提示"）
        self.export_pdf(path)

    # ==================== 导出 HTML（任务 5.3 + 任务 5.1） ====================
    def _export_html(self):
        """
        导出为单文件 HTML。

        任务 5.3：若含视频，自动复制到输出目录的 media/ 子目录。
        任务 5.1：注入动态数字 JS（HTML 里 JS 会执行，保留动画）。

        每张幻灯片作为 <div class='slide'>，通过 CSS 全屏显示。
        直接写文件，无需 QWebEngineView。
        """
        if not self._slides_html:
            return

        path, _ = QFileDialog.getSaveFileName(
            self, "导出 HTML", "演示.html", "HTML (*.html)"
        )
        if not path:
            return

        theme = SLIDE_THEMES[self._current_theme]
        bg = self._custom_bg or theme.get("bg", "#FFF")

        # 任务 5.3：复制视频文件
        slides_html = list(self._slides_html)
        try:
            output_dir = Path(path).parent
            slides_html = self._copy_videos_for_html_export(
                output_dir, slides_html
            )
        except AttributeError:
            pass

        # 构造 HTML 片段
        parts = []
        for s in slides_html:
            parts.append(
                f"<div class='slide' style='width:100%;min-height:100vh;"
                f"display:flex;flex-direction:column;"
                f"justify-content:center;align-items:center;"
                f"padding:60px;box-sizing:border-box'>{s}</div>"
            )
        slides_html_joined = "\n".join(parts)

        # 任务 5.1：动态数字 JS
        dynamic_js = ""
        try:
            dynamic_js = self._get_dynamic_number_js()
        except AttributeError:
            pass

        # 任务 2.1：拼入版式 CSS
        # 优化规划 §3.3：骨架统一走 beeshow_html.render_slide_html()
        full_html = render_slide_html(
            slides_html_joined,
            theme_css=theme['css'],
            layout_css=self._get_layout_css(),
            bg=bg,
            extra_css="body{margin:0}",
            body_extra=dynamic_js,
            wrap_slide=False,      # 每页已在上面单独包好 .slide
        )

        with open(path, 'w', encoding='utf-8') as f:
            f.write(full_html)

        QMessageBox.information(self, "导出成功", f"已保存: {path}")

    # ==================== 导出图片序列（P0-4） ====================
    def _export_images(self):
        """
        导出 PNG 图片序列（P0-4：保持 QWebEngineView 引用，避免被 GC）。

        为每张幻灯片创建一个独立的 QWebEngineView：
          - 加入 self._export_views 列表保持引用
          - loadFinished 后对**视图**取图（见下方 v5.0.5 说明）
          - 完成后从列表移除并 deleteLater

        v5.0.5 修复（优化规划 §1.5）：
          原实现对 QWebEnginePage 取图，而该能力只存在于
          QWidget / QWebEngineView 上，因此**必然**失败；异常被吞掉后
          仍在 finally 里提示「导出成功」——功能完全不存在却报成功。
          现在对视图取图，并让提示基于真实写入结果（检查 pixmap 是否为空
          与 save() 的返回值，有失败页时给出页码清单）。

        历史修复：
          - 原文件用了 logging.error 但未 import logging；
            现在顶部补上 import logging。
        """
        if not self._slides_html:
            return

        output_dir = QFileDialog.getExistingDirectory(
            self, "选择输出文件夹"
        )
        if not output_dir:
            return

        theme = SLIDE_THEMES[self._current_theme]
        bg = self._custom_bg or theme.get("bg", "#FFF")
        total = len(self._slides_html)
        failures = []

        # 3 秒后提示"导出中"，避免同步阻塞主线程。
        #
        # 优化规划 §1.5/§2.x：必须在**进入循环之前**建好这个定时器，
        # 这样最后一张截图完成时才能把它停掉。原先它建在循环之后，
        # 而假/快速渲染下 capture() 在循环里就已经跑完，于是：
        #   ① 停表语句拿不到定时器（还没建），成为空操作；
        #   ② 导出早已结束，它仍在 3 秒后弹「正在生成图片，请稍候…」——
        #      提示与服务端状态相反，实测还会以模态框把测试卡死。
        # 用绑定 self 的单次 QTimer，而不是 QTimer.singleShot：后者不绑定
        # 接收者生命周期，窗口/对象销毁后仍可能触发。
        self._export_hint_timer = QTimer(self)
        self._export_hint_timer.setSingleShot(True)
        self._export_hint_timer.timeout.connect(
            lambda: QMessageBox.information(
                self, "导出中", "正在生成图片，请稍候..."
            )
        )
        self._export_hint_timer.start(3000)

        for i, s in enumerate(self._slides_html):
            # 任务 2.1：拼入版式 CSS
            # 优化规划 §3.3：骨架统一走 beeshow_html.render_slide_html()
            full_html = render_slide_html(
                s,
                theme_css=theme['css'],
                layout_css=self._get_layout_css(),
                bg=bg,
            )

            # P0-4：保持引用，避免每次循环覆盖导致 GC
            temp = QWebEngineView()
            self._export_views.append(temp)
            temp.resize(1024, 768)

            def capture(view, idx=i, out=output_dir):
                """
                loadFinished 回调：截图并保存。

                优化规划 §1.5：原实现对 page.grab() 取值，而 page 是
                QWebEnginePage —— grab() 只存在于 QWidget / QWebEngineView
                上，因此必定抛 AttributeError，被 except 吞掉后仍在 finally
                里提示「导出成功」：功能完全不存在却报成功。
                现在改为对视图 grab()，并让提示基于真实写入结果。
                """
                target = Path(out) / f"slide_{idx+1:03d}.png"
                saved = False
                try:
                    pixmap = view.grab()
                    saved = (not pixmap.isNull()) and pixmap.save(str(target))
                    if not saved:
                        logging.error(f"截图保存失败 slide {idx+1}: {target}")
                except Exception as e:
                    logging.error(f"截图失败 slide {idx+1}: {e}")
                finally:
                    if not saved:
                        failures.append(idx + 1)
                    # 清理引用
                    if view in self._export_views:
                        self._export_views.remove(view)
                    view.deleteLater()
                    # 最后一张完成时提示（基于真实结果）
                    if idx == total - 1:
                        # 导出已结束，撤掉"导出中"提示定时器
                        hint = getattr(self, "_export_hint_timer", None)
                        if hint is not None:
                            try:
                                hint.stop()
                            except Exception:
                                pass
                        if failures:
                            QMessageBox.warning(
                                self, "导出未完成",
                                f"有 {len(failures)} 页截图失败"
                                f"（第 {', '.join(map(str, failures))} 页）。\n\n"
                                f"输出目录: {output_dir}\n"
                                "详见日志。",
                            )
                        else:
                            QMessageBox.information(
                                self, "导出成功",
                                f"已保存 {total} 张 PNG 到:\n{output_dir}",
                            )

            temp.page().loadFinished.connect(
                lambda ok, t=temp, idx=i: capture(t, idx)
            )
            temp.setHtml(full_html)

    # ==================== 导出 GIF ====================
    def _export_gif(self):
        """
        导出为 GIF 动图。

        依赖 bee_export.export_slides_to_gif（延迟 import）。
        用户可选每帧停留时长（默认 2000ms，范围 500~10000ms）。
        """
        if not self._slides_html:
            QMessageBox.information(self, "提示", "请先编写幻灯片内容")
            return

        path, _ = QFileDialog.getSaveFileName(
            self, "导出 GIF", "演示.gif", "GIF 动图 (*.gif)"
        )
        if not path:
            return

        theme = SLIDE_THEMES[self._current_theme]

        from beeoffice.plugins.bee_export import export_slides_to_gif

        duration, ok = QInputDialog.getInt(
            self, "GIF 设置", "每页停留时间（毫秒）:",
            2000, 500, 10000, 500
        )

        if ok and export_slides_to_gif(
            self._slides_html, path, theme['css'], duration
        ):
            QMessageBox.information(
                self, "导出成功", f"GIF 已保存至：\n{path}"
            )
        else:
            QMessageBox.warning(self, "导出失败", "GIF 导出失败")

    # ==================== 录制（占位） ====================
    def _record_presentation(self):
        """
        录制功能占位。

        真正的录制需要：
          - 屏幕捕获库（如 mss / pyautogui）
          - 音视频编码库（如 ffmpeg / moviepy）
          - 与全屏播放窗口的切换时机同步
        当前先以提示代替。
        """
        QMessageBox.information(
            self,
            "录制",
            "录制功能需要在全屏播放时自动记录幻灯片切换时间戳。\n"
            "按 F5 开始播放即可自动记录。"
        )