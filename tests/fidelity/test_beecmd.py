# tests/fidelity/test_beecmd.py
"""
`beecmd` v0（无界面 CLI）的回归。

## 它解决什么问题（战略文档 §5.2）
BeeOffice 的能力原先**全部绑在 GUI 上**，无法被脚本/服务调用 ——
这是"从编辑器变成流水线"最大的缺口。`beecmd` v0 要证明：

> **外部用户不开界面，用一条命令完成一件真实工作。**

## 本文件锁住的行为
* 表格互转（xlsx/csv/json/md/html）与**往返后零差异**（`verify` 逐格对比）
* `verify` 能**精确报差异**（格子地址 + 源值/产物值），并以**退出码 2** 区分"不一致"
* `batch` 失败不中断整批，且**失败也写审计流水**
* **BOM 容忍**（真实踩过的坑：PowerShell/记事本写的 JSON 带 BOM →
  第一次跑 batch 直接 `JSONDecodeError: Unexpected UTF-8 BOM`）
* **纯 Python、零 Qt**：表格路径不得 import PySide6（用子进程验证，避免被测试会话污染）
* 演示文稿 `md → .beshow`（带格式标记）
"""
import json
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


# ==================== 夹具 ====================

@pytest.fixture
def sheet_file(tmp_path):
    """造一个真实 xlsx（中文表名 + 数字/浮点），供往返测试。"""
    import pandas as pd

    path = tmp_path / "sales.xlsx"
    df = pd.DataFrame({
        "产品": ["A", "B", "C"],
        "数量": [10, 20, 30],
        "单价": [1.5, 2.0, 3.25],
    })
    with pd.ExcelWriter(path) as writer:
        df.to_excel(writer, sheet_name="销售", index=False)
    return path


def _run(args):
    from beeoffice import beecmd

    return beecmd.main(args)


# ==================== 1. 转换与往返 ====================

@pytest.mark.parametrize("fmt", ["csv", "json", "md", "html"])
def test_convert_then_verify_no_loss(sheet_file, tmp_path, fmt):
    """xlsx → 各格式 → verify 必须零差异（表格类格式逐格对比）。"""
    out = tmp_path / f"out.{fmt}"
    assert _run(["convert", str(sheet_file), "--to", fmt, "-o", str(out)]) == 0
    assert out.is_file()

    report = tmp_path / "report.json"
    code = _run(["verify", str(sheet_file), str(out), "--report", str(report)])
    data = json.loads(report.read_text(encoding="utf-8"))
    assert code == (0 if data["ok"] else 2)
    assert data["ok"] is True, data
    assert data["mismatches"] == 0, data


def test_roundtrip_csv_to_xlsx_has_zero_mismatch(tmp_path):
    """csv → xlsx → verify：往返零差异。"""
    csv = tmp_path / "in.csv"
    csv.write_text("名称,数量\n苹果,3\n梨,5\n", encoding="utf-8-sig")
    xlsx = tmp_path / "out.xlsx"

    assert _run(["convert", str(csv), "--to", "xlsx", "-o", str(xlsx)]) == 0
    report = tmp_path / "r.json"
    _run(["verify", str(csv), str(xlsx), "--report", str(report)])
    data = json.loads(report.read_text(encoding="utf-8"))
    assert data["ok"] is True, data


def test_info_reports_sheets(sheet_file, capsys):
    assert _run(["info", str(sheet_file)]) == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["sheets"][0]["name"] == "销售"
    assert payload["sheets"][0]["rows"] == 3
    assert payload["sheets"][0]["cols"] == 3


# ==================== 2. verify 的差异检测 ====================

def test_verify_detects_mismatch_and_exits_2(sheet_file, tmp_path):
    """**核心**：产物被改坏 → 必须报出具体格子，并以退出码 2 结束。"""
    bad = tmp_path / "bad.csv"
    bad.write_text("产品,数量,单价\nA,10,9.99\nB,20,2.0\nC,30,3.25\n",
                   encoding="utf-8-sig")

    report = tmp_path / "r.json"
    code = _run(["verify", str(sheet_file), str(bad), "--report", str(report)])
    data = json.loads(report.read_text(encoding="utf-8"))

    assert code == 2, "不一致必须以退出码 2 区分"
    assert data["ok"] is False
    assert data["mismatches"] == 1
    sample = data["samples"][0]
    assert sample["cell"] == "R1C3", sample
    assert sample["src"] == "1.5" and sample["artifact"] == "9.99", sample


def test_verify_matches_by_name_or_position(tmp_path):
    """跨格式时表名不同（xlsx 的"销售" vs csv 的文件名）→ 按位置配，不能误报缺表。"""
    xlsx = tmp_path / "a.xlsx"
    import pandas as pd

    with pd.ExcelWriter(xlsx) as writer:
        pd.DataFrame({"x": [1, 2]}).to_excel(writer, sheet_name="销售",
                                            index=False)
    csv = tmp_path / "a.csv"
    csv.write_text("x\n1\n2\n", encoding="utf-8-sig")

    report = tmp_path / "r.json"
    code = _run(["verify", str(xlsx), str(csv), "--report", str(report)])
    data = json.loads(report.read_text(encoding="utf-8"))
    assert code == 0, data
    assert data["ok"] is True


# ==================== 3. BOM 回归（真实踩过的坑）====================

def test_batch_tolerates_bom_in_jobs_json(sheet_file, tmp_path):
    """
    PowerShell/记事本写的 JSON **带 UTF-8 BOM** ——
    第一次跑 batch 就因此 `JSONDecodeError: Unexpected UTF-8 BOM`。
    现在用 `utf-8-sig` 读，必须容忍。
    """
    jobs = tmp_path / "jobs.json"
    payload = {"jobs": [
        {"op": "convert", "src": str(sheet_file), "to": "csv",
         "out": str(tmp_path / "a.csv")},
    ]}
    jobs.write_bytes(b"\xef\xbb\xbf" + json.dumps(payload).encode("utf-8"))

    audit = tmp_path / "audit.jsonl"
    assert _run(["batch", str(jobs), "--audit", str(audit)]) == 0
    assert audit.is_file()
    assert len(audit.read_text(encoding="utf-8").splitlines()) == 1


def test_convert_tolerates_bom_in_markdown(tmp_path):
    """带 BOM 的 Markdown 表格也要能转换。"""
    md = tmp_path / "t.md"
    md.write_bytes(b"\xef\xbb\xbf| a | b |\n| --- | --- |\n| 1 | 2 |\n")
    out = tmp_path / "t.csv"
    assert _run(["convert", str(md), "--to", "csv", "-o", str(out)]) == 0
    text = out.read_text(encoding="utf-8-sig")
    assert "a" in text and "1" in text


# ==================== 4. batch：失败隔离 + 审计 ====================

def test_batch_isolates_failure_and_audits(sheet_file, tmp_path):
    jobs = tmp_path / "jobs.json"
    payload = {"jobs": [
        {"op": "convert", "src": str(sheet_file), "to": "csv",
         "out": str(tmp_path / "ok.csv")},
        {"op": "convert", "src": str(tmp_path / "missing.xlsx"), "to": "csv"},
    ]}
    jobs.write_text(json.dumps(payload), encoding="utf-8")
    audit = tmp_path / "audit.jsonl"

    code = _run(["batch", str(jobs), "--audit", str(audit)])

    assert code == 1, "有失败时整批应为 1"
    lines = [json.loads(ln) for ln in
             audit.read_text(encoding="utf-8").splitlines()]
    assert len(lines) == 2, "失败的那条也必须写审计"
    assert lines[0]["ok"] is True
    assert lines[1]["ok"] is False
    assert (tmp_path / "ok.csv").is_file(), "第一条不该被第二条拖累"


# ==================== 5. 演示文稿 md → .beshow ====================

def test_convert_markdown_to_beshow(tmp_path):
    from beeoffice.plugins.beeshow_doc import is_deck_file, load_deck

    md = tmp_path / "deck.md"
    md.write_text("# 第一页\n\n- 要点\n\n---\n\n# 第二页\n", encoding="utf-8")
    out = tmp_path / "deck.beshow"

    assert _run(["convert", str(md), "--to", "beshow", "-o", str(out)]) == 0
    assert is_deck_file(out)
    assert "第一页" in (load_deck(out) or "")


# ==================== 6. 零 Qt 证明 ====================

def test_sheet_paths_do_not_import_qt(sheet_file, tmp_path):
    """
    **架构约束**：表格路径必须能在没有 Qt 的环境跑（服务器/容器）。
    在子进程里验证，避免被本测试会话已经 import 的 PySide6 污染。

    ⚠️ 沙箱边界：子进程**不能用管道捕获输出**（`subprocess` 的 `stdout=PIPE`
    在本沙箱里拿不到数据、返回 None）。所以让子进程把结果**写文件**，父进程读文件。
    """
    result_file = tmp_path / "qt_probe.txt"
    script = (
        "import sys\n"
        "from beeoffice import beecmd\n"
        f"code = beecmd.main(['convert', {str(sheet_file)!r}, '--to', 'md',"
        f" '-o', {str(tmp_path / 'noqt.md')!r}])\n"
        "loaded = [m for m in sys.modules if m.startswith(('PySide6', 'shiboken'))]\n"
        f"open({str(result_file)!r}, 'w', encoding='utf-8').write("
        "f'code={code};qt={\",\".join(loaded)}')\n"
    )
    completed = subprocess.run(
        [sys.executable, "-c", script], cwd=str(ROOT),
        env={**__import__("os").environ, "PYTHONPATH": str(ROOT)},
        timeout=180,
    )
    assert completed.returncode == 0, "子进程转换失败"
    payload = result_file.read_text(encoding="utf-8")
    assert payload.startswith("code=0;"), payload
    assert payload.endswith("qt="), f"表格路径不该加载 Qt：{payload}"


# ==================== 7. CLI 接口 ====================

def test_version_flag():
    with pytest.raises(SystemExit) as exc:
        _run(["--version"])
    assert exc.value.code == 0


def test_unknown_format_is_actionable(sheet_file, capsys, tmp_path):
    """不支持的格式 / 不支持的源类型，都要给出可操作提示。"""
    code = _run(["convert", str(sheet_file), "--to", "pptx",
                 "-o", str(tmp_path / "x.pptx")])
    assert code == 1
    err = capsys.readouterr().err
    assert "不支持" in err, err


# ==================== 8. v0.1：Markdown → docx / pdf ====================

@pytest.fixture
def markdown_doc(tmp_path):
    path = tmp_path / "报告.md"
    path.write_text(
        "# 季度报告\n\n这是一段**测试正文**，包含中文与 English words。\n\n"
        "- 要点一\n- 要点二\n",
        encoding="utf-8",
    )
    return path


def test_convert_markdown_to_docx_has_real_text(markdown_doc, tmp_path):
    """
    **v0.1 核心**：md → docx 必须写出**有内容**的 Word。

    我踩过的坑：用 `rich_editor.set_markdown()` 喂内容时，`_write_docx`
    在 markdown 模式下取 `get_rendered_content()` → 写出**空文档**，
    而 `_write_docx` 仍返回 True（假成功 ✗）。正确入口是 `set_content()`。
    """
    from docx import Document

    out = tmp_path / "报告.docx"
    assert _run(["convert", str(markdown_doc), "--to", "docx",
                 "-o", str(out)]) == 0
    assert out.is_file()

    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert len(doc.paragraphs) >= 2, f"段落太少：{len(doc.paragraphs)}"
    assert "季度报告" in text, text
    assert "测试正文" in text, text


def test_verify_markdown_vs_docx_is_clean(markdown_doc, tmp_path):
    """docx 有文本后，词元级 verify 应判为一致（退出码 0）。"""
    out = tmp_path / "报告.docx"
    _run(["convert", str(markdown_doc), "--to", "docx", "-o", str(out)])
    report = tmp_path / "r.json"
    code = _run(["verify", str(markdown_doc), str(out), "--report", str(report)])
    data = json.loads(report.read_text(encoding="utf-8"))
    assert code == 0, data
    assert data["ok"] is True
    assert data["mismatches"] == 0


def test_convert_markdown_to_pdf_produces_pages(markdown_doc, tmp_path):
    """md → pdf 至少产出一个有页面的 PDF（内容验证见下一条）。"""
    pytest.importorskip("fitz")
    import fitz

    out = tmp_path / "报告.pdf"
    assert _run(["convert", str(markdown_doc), "--to", "pdf",
                 "-o", str(out)]) == 0
    assert out.is_file() and out.stat().st_size > 0
    with fitz.open(str(out)) as doc:
        assert doc.page_count >= 1


def test_verify_pdf_is_inconclusive_when_text_unavailable(markdown_doc, tmp_path):
    """
    **诚实性**：PDF 提取不到文本时，verify 必须返回**退出码 3（不可判定）**，
    既不能假装通过、也不能假装失败。

    本机会因为 offscreen 无字体而提取不到文本；装了字体的机器上则应通过。
    两种都接受，但**必须自洽**。
    """
    pytest.importorskip("fitz")
    out = tmp_path / "报告.pdf"
    _run(["convert", str(markdown_doc), "--to", "pdf", "-o", str(out)])

    report = tmp_path / "r.json"
    code = _run(["verify", str(markdown_doc), str(out), "--report", str(report)])
    data = json.loads(report.read_text(encoding="utf-8"))

    assert code in (0, 3), data
    if code == 3:
        assert data.get("inconclusive") is True
        assert data.get("note"), "不可判定必须说明原因"
    else:
        assert data["ok"] is True


def test_table_source_to_pdf_is_honestly_rejected(sheet_file, capsys, tmp_path):
    """表格 → PDF 还没做：必须明确说不支持，而不是产出空白文件。"""
    code = _run(["convert", str(sheet_file), "--to", "pdf",
                 "-o", str(tmp_path / "x.pdf")])
    assert code == 1
    err = capsys.readouterr().err
    assert "只支持 Markdown" in err or "尚未支持" in err, err
    assert not (tmp_path / "x.pdf").exists(), "不支持时不该产出文件"
