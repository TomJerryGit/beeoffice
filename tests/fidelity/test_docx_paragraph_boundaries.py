# tests/fidelity/test_docx_paragraph_boundaries.py
"""
保真度阶段 A：DOCX **段落边界**必须保留。

缺陷（实测）
------------
`import_from_docx()` 每个段落输出**一行**、行间不留空行。而 Markdown 的语义是
**连续非空行属于同一个段落**，于是：

    以下表格含表头、合并单元格与空单元格。
    ![](f_media/image_001.png)
    [BeeOffice 项目主页](https://example.com)

在 Markdown 里是**一段**，重新导出 DOCX 后这三样东西就粘在了同一段里。
实测：顶层段落 **6 → 4**，`paragraphs` 也从 14 变 13。

修法：段落之间补空行（列表项之间除外 —— 空行会把一个列表拆成多个）。
"""
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document                                  # noqa: E402

from beeoffice.plugins.bee_export import (                 # noqa: E402
    export_to_docx,
    import_from_docx,
    read_docx_extras,
)


def _roundtrip(src: Path, tmp_path: Path) -> Path:
    out = tmp_path / "out.docx"
    extras = read_docx_extras(src)
    assert export_to_docx(
        import_from_docx(str(src)), str(out),
        page_settings=extras, base_dir=str(tmp_path),
        header_text=extras.get("header_text"),
        footer_text=extras.get("footer_text"),
    ) is True
    return out


def _texts(path: Path):
    return [p.text for p in Document(str(path)).paragraphs]


# ==================== 1. Markdown 里必须有段落分隔 ====================

def test_imported_markdown_separates_paragraphs(tmp_path):
    """相邻段落之间必须有空行，否则 Markdown 会把它们并成一段。"""
    from docx import Document as D

    src = tmp_path / "src.docx"
    doc = D()
    doc.add_paragraph("第一段")
    doc.add_paragraph("第二段")
    doc.save(str(src))

    md = import_from_docx(str(src))

    assert "第一段\n\n第二段" in md, (
        f"两段之间没有空行，Markdown 会把它们渲染成同一段: {md!r}"
    )


def test_markdown_from_fidelity_fixture_keeps_blocks_apart(tmp_path):
    """图片段与链接段必须各自独立，不能被并进上一段。"""
    from tools.fidelity_roundtrip import build_docx_fixture

    src = tmp_path / "f.docx"
    build_docx_fixture(src)

    md = import_from_docx(str(src))
    blocks = [b.strip() for b in md.split("\n\n") if b.strip()]

    assert any(b.startswith("![") for b in blocks), (
        f"图片没有独立成块: {blocks}"
    )
    assert any(b.startswith("[") for b in blocks), (
        f"超链接没有独立成块: {blocks}"
    )
    assert not any(
        "以下表格" in b and "![](f_media" in b for b in blocks
    ), "正文段落与图片被并进了同一块"


# ==================== 2. 列表不能被空行拆开 ====================

def test_list_items_stay_together(tmp_path):
    """列表项之间**不应**插空行 —— 那会把一个列表拆成多个。"""
    from docx import Document as D

    src = tmp_path / "src.docx"
    doc = D()
    doc.add_paragraph("甲项", style="List Bullet")
    doc.add_paragraph("乙项", style="List Bullet")
    doc.save(str(src))

    md = import_from_docx(str(src))
    lines = [ln for ln in md.splitlines() if ln.strip()]

    assert lines[0].startswith("- ") and lines[1].startswith("- "), lines
    assert "\n\n" not in md, f"列表项之间被插了空行: {md!r}"


# ==================== 3. 端到端：段落结构保持一致 ====================

def test_paragraph_count_is_preserved(tmp_path):
    """**核心断言**：顶层段落数不因往返而减少。"""
    from tools.fidelity_roundtrip import build_docx_fixture

    src = tmp_path / "f.docx"
    build_docx_fixture(src)

    out = _roundtrip(src, tmp_path)

    before = _texts(src)
    after = _texts(out)
    assert len(after) == len(before), (
        f"段落数变了: {len(before)} → {len(after)}\n"
        f"往返前: {before}\n往返后: {after}"
    )


def test_paragraph_texts_match_in_order(tmp_path):
    """逐段文字也应一致（不能把两段并成一段）。"""
    from tools.fidelity_roundtrip import build_docx_fixture

    src = tmp_path / "f.docx"
    build_docx_fixture(src)

    out = _roundtrip(src, tmp_path)

    assert _texts(src) == _texts(out), (
        f"往返前: {_texts(src)}\n往返后: {_texts(out)}"
    )


def test_image_and_link_are_in_their_own_paragraphs(tmp_path):
    """图片与超链接各自独立成段（原先被并进上一段）。"""
    from tools.fidelity_roundtrip import build_docx_fixture

    src = tmp_path / "f.docx"
    build_docx_fixture(src)
    out = _roundtrip(src, tmp_path)

    texts = _texts(out)
    assert any("项目主页" in t for t in texts), (
        f"链接没有独立成段: {texts}"
    )
    assert not any("空单元格" in t and "项目主页" in t for t in texts), (
        f"链接被并进了正文段: {texts}"
    )
