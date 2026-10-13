# tests/fidelity/test_docx_extras.py
"""
保真度阶段 A：DOCX 的**页面设置 + 页眉页脚**往返。

问题（实测基线）
----------------
`import_from_docx()` 只返回一个 **Markdown 字符串**，而 Markdown 没有
"页面设置/页眉页脚"的概念。于是往返后：

    page_margins  1417,1701,1417,1701 → 1440,1800,1440,1800（变默认 1 英寸）
    header_parts  1 → 0（`word/header1.xml` 整个不存在）
    footer_parts  1 → 0

修法：这些信息**绕开 Markdown** 传递 ——
导入时 `read_docx_extras()` 另读一份 → `widget.apply_docx_extras()` 存到文档上
→ 导出时 `export_to_docx(..., header_text=, footer_text=)` 写回。

本文件覆盖：读取、应用、写回、非标准纸张、以及**产品级**端到端。
"""
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from beeoffice.plugins.bee_export import (   # noqa: E402
    apply_header_footer,
    export_to_docx,
    import_from_docx,
    read_docx_extras,
)


def _make_docx(path: Path, *, top=25.0, bottom=25.0, left=30.0, right=30.0,
               header="页眉文字", footer="页脚文字",
               width_mm=None, height_mm=None):
    from docx import Document
    from docx.shared import Mm

    doc = Document()
    doc.add_paragraph("正文一段")
    section = doc.sections[0]
    section.top_margin = Mm(top)
    section.bottom_margin = Mm(bottom)
    section.left_margin = Mm(left)
    section.right_margin = Mm(right)
    if width_mm and height_mm:
        section.page_width = Mm(width_mm)
        section.page_height = Mm(height_mm)
    if header:
        section.header.paragraphs[0].text = header
    if footer:
        section.footer.paragraphs[0].text = footer
    doc.save(str(path))
    return path


# ==================== 1. 读取 ====================

def test_reads_margins_in_mm(tmp_path):
    """页边距必须以 mm 读出（项目的统一单位约定）。"""
    src = _make_docx(tmp_path / "a.docx", top=25, bottom=25, left=30, right=30)

    extras = read_docx_extras(src)

    assert extras["margin_top"] == pytest.approx(25.0, abs=0.2)
    assert extras["margin_left"] == pytest.approx(30.0, abs=0.2)
    assert extras["orientation"] == "portrait"


def test_reads_header_and_footer_text(tmp_path):
    src = _make_docx(tmp_path / "a.docx")

    extras = read_docx_extras(src)

    assert extras["header_text"] == "页眉文字"
    assert extras["footer_text"] == "页脚文字"


def test_missing_header_footer_yields_empty_strings(tmp_path):
    src = _make_docx(tmp_path / "a.docx", header="", footer="")

    extras = read_docx_extras(src)

    assert extras.get("header_text", "") == ""
    assert extras.get("footer_text", "") == ""


def test_detects_landscape_orientation(tmp_path):
    src = _make_docx(tmp_path / "a.docx", width_mm=297, height_mm=210)

    extras = read_docx_extras(src)

    assert extras["orientation"] == "landscape", extras


def test_keeps_exact_non_standard_page_size(tmp_path):
    """
    非标准纸张必须**精确**保留尺寸。

    否则会被当成 A4 —— 那是一次静默改变（`PAPER_SIZES_MM` 只认 5 种命名纸张）。
    """
    src = _make_docx(tmp_path / "a.docx", width_mm=200, height_mm=280)

    extras = read_docx_extras(src)

    assert extras["page_width_mm"] == pytest.approx(200, abs=0.5)
    assert extras["page_height_mm"] == pytest.approx(280, abs=0.5)


def test_unreadable_file_returns_empty(tmp_path):
    """损坏文件不应抛异常（打开流程不能被它打断）。"""
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a docx")

    assert read_docx_extras(bad) == {}


# ==================== 2. 写回 ====================

def test_apply_header_footer_creates_parts(tmp_path):
    from docx import Document
    from tools.ooxml_probe import probe

    doc = Document()
    doc.add_paragraph("正文")
    out = tmp_path / "out.docx"

    assert apply_header_footer(doc, "H", "F") is True
    doc.save(str(out))

    features = probe(out)["features"]
    assert features["header_parts"] == 1
    assert features["footer_parts"] == 1


def test_apply_header_footer_with_nothing_is_noop(tmp_path):
    """都是空 → 不应凭空造出页眉页脚部件。"""
    from docx import Document
    from tools.ooxml_probe import probe

    doc = Document()
    doc.add_paragraph("正文")
    out = tmp_path / "out.docx"

    assert apply_header_footer(doc, "", None) is False
    doc.save(str(out))

    features = probe(out)["features"]
    assert features["header_parts"] == 0
    assert features["footer_parts"] == 0


# ==================== 3. 端到端：往返后不再丢 ====================

def test_roundtrip_preserves_header_footer_and_margins(tmp_path):
    """**核心断言**：页眉/页脚部件与页边距必须逐项一致。"""
    from tools.ooxml_probe import compare, probe

    src = _make_docx(tmp_path / "src.docx")
    before = probe(src)

    out = tmp_path / "out.docx"
    extras = read_docx_extras(src)
    assert export_to_docx(
        import_from_docx(str(src)), str(out),
        page_settings=extras, base_dir=str(tmp_path),
        header_text=extras.get("header_text"),
        footer_text=extras.get("footer_text"),
    ) is True

    rows = {r["feature"]: r for r in compare(before, probe(out))}
    for key in ("header_parts", "footer_parts", "page_margins"):
        assert rows[key]["status"] == "same", (
            f"{key} 没保住: {rows[key]}"
        )


def test_roundtrip_keeps_exact_page_size(tmp_path):
    from docx import Document

    src = _make_docx(tmp_path / "src.docx", width_mm=200, height_mm=280)
    out = tmp_path / "out.docx"
    extras = read_docx_extras(src)

    export_to_docx(
        import_from_docx(str(src)), str(out),
        page_settings=extras, base_dir=str(tmp_path),
        header_text=extras.get("header_text"),
    )

    section = Document(str(out)).sections[0]
    assert round(float(section.page_width.mm), 1) == pytest.approx(200, abs=0.5)
    assert round(float(section.page_height.mm), 1) == pytest.approx(280, abs=0.5)


# ==================== 4. 编辑器 API（产品接线的一环） ====================

def test_editor_apply_and_get_extras(beewrite_editor):
    extras = {
        "paper": "Letter",
        "orientation": "portrait",
        "margin_top": 12.5,
        "margin_left": 20.0,
        "page_width_mm": 215.9,
        "page_height_mm": 279.4,
        "header_text": "我的页眉",
        "footer_text": "我的页脚",
    }

    applied = beewrite_editor.apply_docx_extras(extras)

    assert applied >= 6, applied
    settings = beewrite_editor.get_page_settings()
    assert settings["margin_top"] == 12.5
    assert settings["paper"] == "Letter"
    assert settings["page_width_mm"] == 215.9
    assert beewrite_editor.get_docx_extras() == {
        "header_text": "我的页眉", "footer_text": "我的页脚",
        "header_xml": "", "footer_xml": "",
        "table_merges": [], "sections": [], "bookmarks": [],
        "paragraph_styles": [], "style_definitions": {},
    }


def test_editor_apply_extras_ignores_empty_values(beewrite_editor):
    """解析失败的项（None/空串）不应覆盖编辑器当前值。"""
    before = beewrite_editor.get_page_settings()

    beewrite_editor.apply_docx_extras({
        "paper": "", "margin_top": None, "margin_left": 33.0,
    })

    after = beewrite_editor.get_page_settings()
    assert after["paper"] == before["paper"], "空值不应覆盖"
    assert after["margin_top"] == before["margin_top"], "None 不应覆盖"
    assert after["margin_left"] == 33.0, "有值的项应当生效"


def test_editor_get_extras_defaults_to_empty(beewrite_editor):
    """没导入过 DOCX 时，导出不应写出页眉页脚/分节/表格合并/书签。"""
    assert beewrite_editor.get_docx_extras() == {
        "header_text": "", "footer_text": "",
        "header_xml": "", "footer_xml": "",
        "table_merges": [], "sections": [], "bookmarks": [],
        "paragraph_styles": [], "style_definitions": {},
    }


# ==================== 5. 产品级：导出真的会写页眉页脚 ====================

def test_write_docx_writes_header_footer_from_extras(beewrite_editor, tmp_path):
    """
    端到端（产品路径）：编辑器收到 extras 后，`_write_docx()` 导出的文件里
    必须真有页眉页脚部件。

    这条防的是"模块写好了但没接上"——本项目踩过同类坑。
    """
    from tools.ooxml_probe import probe

    beewrite_editor.set_content("正文内容")
    beewrite_editor.apply_docx_extras({
        "header_text": "导出页眉",
        "footer_text": "导出页脚",
        "margin_top": 20.0,
    })

    out = tmp_path / "wired.docx"
    assert beewrite_editor._write_docx(str(out)) is True

    features = probe(out)["features"]
    assert features["header_parts"] == 1, "页眉部件没写出来"
    assert features["footer_parts"] == 1, "页脚部件没写出来"


def test_app_open_docx_file_applies_extras():
    """
    `app._open_docx_file` 必须读并应用 extras（否则整条链路断在入口）。

    用 **AST** 断言真实调用节点，不用文本搜索（本会话被自己的注释骗过多次）。
    """
    import ast
    import inspect
    import textwrap

    from beeoffice.app import BeeApp

    src = textwrap.dedent(inspect.getsource(BeeApp._open_docx_file))
    tree = ast.parse(src)

    imported = {
        alias.name
        for node in ast.walk(tree)
        if isinstance(node, ast.ImportFrom)
        for alias in node.names
    }
    called = {
        node.func.attr
        for node in ast.walk(tree)
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute)
    }

    assert "read_docx_extras" in imported, "没有导入 read_docx_extras"
    assert "apply_docx_extras" in called, "没有调用 widget.apply_docx_extras"
