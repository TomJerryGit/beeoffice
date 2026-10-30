"""Writer 题注：静态文字 → Word SEQ 域。"""
from docx import Document

from beeoffice.plugins.bee_export import export_to_docx


def test_caption_exports_seq_field(tmp_path):
    md = "图 1：系统架构\n\n表 2：数据字典\n"
    out = tmp_path / "captions.docx"
    assert export_to_docx(md, str(out)) is True
    doc = Document(str(out))
    xml = doc.element.body.xml
    assert "SEQ 图" in xml
    assert "SEQ 表" in xml
    assert "系统架构" in xml
    assert "数据字典" in xml


def test_figure_toc_field(tmp_path):
    md = "图 1：系统架构\n"
    out = tmp_path / "figure_toc.docx"
    assert export_to_docx(md, str(out), figure_toc=True) is True
    doc = Document(str(out))
    xml = doc.element.body.xml
    assert 'TOC \\h \\z \\c "图"' in xml
    assert "SEQ 图" in xml


def test_normal_paragraph_not_caption(tmp_path):
    md = "这是普通段落，不是题注。\n"
    out = tmp_path / "normal.docx"
    assert export_to_docx(md, str(out)) is True
    doc = Document(str(out))
    assert "SEQ " not in doc.element.body.xml
