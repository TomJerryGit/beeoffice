# tests/fidelity/test_atomic_saves.py
"""
P0 止血回归（审计 §4.3 G2）：主文档保存必须**原子**。

背景：`python-docx` / `openpyxl` / `python-pptx` 的 `save(path)` 直接以写模式
打开目标文件；写到一半崩溃/断电/磁盘满会把用户原文件留成半份 zip。
本文件验证 `utils.atomic_replace` 的语义，以及 DOCX/XLSX 两条主保存路径
在失败时会**保留原文件、不残留临时文件**。
"""
from contextlib import contextmanager
from pathlib import Path

import pytest

from beeoffice.utils import atomic_replace


# ==================== 原子替换原语 ====================

def test_atomic_replace_success_and_no_temp_left(tmp_path):
    target = tmp_path / "book.xlsx"
    target.write_bytes(b"OLD")

    with atomic_replace(target) as tmp:
        assert Path(tmp).parent == tmp_path, "临时文件必须与目标同目录"
        assert Path(tmp).suffix == ".xlsx", "必须保留扩展名（库按扩展名选引擎）"
        Path(tmp).write_bytes(b"NEW")

    assert target.read_bytes() == b"NEW"
    assert not list(tmp_path.glob(".*.tmp*")), "成功后不应残留临时文件"


def test_atomic_replace_failure_keeps_original(tmp_path):
    target = tmp_path / "book.xlsx"
    target.write_bytes(b"ORIGINAL")

    with pytest.raises(RuntimeError):
        with atomic_replace(target) as tmp:
            Path(tmp).write_bytes(b"PARTIAL")
            raise RuntimeError("disk full")

    assert target.read_bytes() == b"ORIGINAL", "失败时原文件必须保持不变"
    assert not list(tmp_path.glob(".*.tmp*")), "失败后不应残留临时文件"


# ==================== DOCX 主保存 ====================

def test_export_to_docx_failure_is_atomic(tmp_path, monkeypatch):
    """库在 save 时写到一半失败 → 原 docx 不被破坏。"""
    from docx.document import Document as DocxDocument

    from beeoffice.plugins import bee_export

    target = tmp_path / "out.docx"
    target.write_bytes(b"ORIGINAL")

    def _boom(self, path):
        # 模拟 python-docx 已经写了半份包再失败
        Path(path).write_bytes(b"PARTIAL")
        raise OSError("disk full")

    monkeypatch.setattr(DocxDocument, "save", _boom)

    assert bee_export.export_to_docx("# 标题\n\n正文", str(target)) is False
    assert target.read_bytes() == b"ORIGINAL"
    assert not list(tmp_path.glob(".*.tmp*"))


# ==================== XLSX 主保存 ====================

def test_write_xlsx_failure_is_atomic(tmp_path, monkeypatch):
    """写 xlsx 或贴样式时失败 → 原工作簿不被破坏。"""
    import pandas as pd

    from beeoffice.plugins.beesheet_data_mixin import DataMixin

    class _Fake(DataMixin):
        def __init__(self):
            self._active_sheet_name = "Sheet1"
            self._workbook = None

        def _get_export_df(self):
            return pd.DataFrame({"a": [1, 2]})

        def _apply_view_styles_to(self, *args, **kwargs):
            # 数据已经写进临时文件后再失败（最危险的时刻）
            raise RuntimeError("style boom")

    target = tmp_path / "book.xlsx"
    target.write_bytes(b"ORIGINAL")
    fake = _Fake()

    assert DataMixin._write_xlsx(fake, str(target)) is False
    assert target.read_bytes() == b"ORIGINAL"
    assert not list(tmp_path.glob(".*.tmp*"))


def test_write_xlsx_success_replaces_target(tmp_path):
    """正常保存路径仍然产出可读的 xlsx（原子替换没有破坏内容）。"""
    import pandas as pd
    from openpyxl import load_workbook

    from beeoffice.plugins.beesheet_data_mixin import DataMixin

    class _Fake(DataMixin):
        def __init__(self):
            self._active_sheet_name = "Sheet1"
            self._workbook = None

        def _get_export_df(self):
            return pd.DataFrame({"a": [1, 2], "b": ["x", "y"]})

        def _apply_view_styles_to(self, *args, **kwargs):
            return None

    target = tmp_path / "book.xlsx"
    fake = _Fake()

    assert DataMixin._write_xlsx(fake, str(target)) is True
    wb = load_workbook(str(target))
    ws = wb.active
    assert ws["A1"].value == "a"
    assert ws["A2"].value == 1
    wb.close()
    assert not list(tmp_path.glob(".*.tmp*"))


# ==================== Markdown 主保存 ====================

def test_write_md_failure_keeps_original(tmp_path, monkeypatch):
    """Markdown 保存失败 → 原文件不被截断（原先 open('w') 会截断）。"""
    from PySide6.QtWidgets import QMessageBox

    from beeoffice.plugins import beewrite_export_mixin as mod

    target = tmp_path / "note.md"
    target.write_text("旧内容", encoding="utf-8")

    class _Fake:
        _current_file_path = None
        _edit_mode = "markdown"

        def content(self):
            return "# 新内容"

        def set_current_file_path(self, path):
            self._current_file_path = path

    fake = _Fake()

    @contextmanager
    def _boom(path):
        # 模拟临时文件写入失败
        raise PermissionError("locked")
        yield  # pragma: no cover

    monkeypatch.setattr(mod, "atomic_replace", _boom)
    monkeypatch.setattr(
        QMessageBox, "warning", staticmethod(lambda *a, **k: QMessageBox.Ok)
    )

    assert mod.ExportMixin._write_md(fake, str(target)) is False
    assert target.read_text(encoding="utf-8") == "旧内容"
