"""Writer DOCX 浮动图片环绕导出。"""
from bs4 import BeautifulSoup
from docx import Document
from PIL import Image

from beeoffice.plugins.bee_export import _write_inline_image


def _tag(attrs):
    return BeautifulSoup(f'<img src="pic.png" {attrs}>', "html.parser").img


def test_square_wrap_exports_anchor(tmp_path):
    Image.new("RGB", (40, 20), "blue").save(tmp_path / "pic.png")
    doc = Document()
    para = doc.add_paragraph()
    _write_inline_image(_tag('data-wrap="square"'), para, str(tmp_path))
    xml = para._p.xml
    assert "wp:anchor" in xml
    assert "wp:wrapSquare" in xml


def test_behind_wrap_exports_wrap_none(tmp_path):
    Image.new("RGB", (40, 20), "blue").save(tmp_path / "pic.png")
    doc = Document()
    para = doc.add_paragraph()
    _write_inline_image(_tag('data-wrap="behind"'), para, str(tmp_path))
    xml = para._p.xml
    assert "wp:anchor" in xml
    assert "wp:wrapNone" in xml


def test_no_wrap_keeps_inline(tmp_path):
    Image.new("RGB", (40, 20), "blue").save(tmp_path / "pic.png")
    doc = Document()
    para = doc.add_paragraph()
    _write_inline_image(_tag(""), para, str(tmp_path))
    xml = para._p.xml
    assert "wp:inline" in xml
    assert "wp:anchor" not in xml
