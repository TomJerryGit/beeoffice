# tests/fidelity/test_p1_write_gaps.py
"""
第八批：条件格式**读回** + 蜜写三处"按下去没用/东西会变样"的修复。

## 为什么做这三处（审计报告 §2.13，均经复核）
1. **默认模式下 Ctrl+B/I/U 完全没反应**：默认是 markdown 模式，而三个方法
   第一句就是 `if self._edit_mode != "rich_text": return` —— 按下去不改内容、
   也不给提示，而快捷键帮助里**明确写着** "Ctrl+B 加粗"。
2. **「§ 分节符」不是分节、还污染导出**：只插入一行可见装饰文字
   `════ 分节符（新节开始） ════`，那行字**原样进入 DOCX/PDF**，
   而 Word 里根本没有"节"（tooltip 却承诺"用于改变页面设置"）。
   「分页符」同病（也插了 `──── 分页符 ────`）—— 一并修。
3. **「导出 Word」与「另存为 Word」保真度不同**：另存为传全套 extras
   （页面设置/页眉页脚/合并/书签/相对图片），而工具栏/菜单的"导出 Word"
   只传 content+path → 同一文档两个入口结果不一样。

## 另外两处一并修（同批）
4. **拼写检查「全部替换」会抛异常**：`self.items.index(item)` 在
   `_replace_item()` 重建列表之后使用 → 拿错行或直接 `ValueError` 中断。
   （老测试用**常量文本桩**且替换词等长，所以永远测不出来 —— 本文件的测试
   故意用**会变的文本**和**不等长**的替换词。）
5. **条件格式读回**（第四/五批只写了不读）：openpyxl **读回**的是通用
   `Rule`（不是 `DataBarRule`），信息在 `rule.dataBar` / `rule.colorScale`
   的 `cfvo` 里 —— 这条是我实测才发现的（第一版按类名判断，永远读不回来）。
"""
import sys
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def _load(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    return workbook, workbook[workbook.sheetnames[0]]


# ==================== 1. 条件格式读回 ====================

def test_databar_roundtrip(tmp_path):
    from beeoffice.plugins.beesheet_style_export import (
        apply_view_styles, read_view_styles,
    )

    path = tmp_path / "cf.xlsx"
    pd.DataFrame({"名称": ["甲", "乙", "丙"], "金额": [3.25, 100.0, 1234.5]}).to_excel(
        path, index=False
    )
    apply_view_styles(
        path, "Sheet1",
        conditional={1: {"type": "data_bars", "min": 3.25, "max": 1234.5}},
        row_count=3,
    )

    styles = read_view_styles(path)

    assert styles.get("conditional") == {
        1: {"type": "data_bars", "min": 3.25, "max": 1234.5}
    }


def test_colorscale_with_auto_endpoints_is_computed_from_column(tmp_path):
    """
    色阶的端点在 Excel 里常是"自动最小/最大"（没有数值）→ 读回时按该列数据推算。

    注意：这里**用产品自己的写入器**造色阶 —— 我第一版手工 `ColorScaleRule(...)`
    没给颜色，openpyxl 保存时直接把规则丢了（"组数 0"），白白怀疑了一轮产品代码。
    """
    from beeoffice.plugins.beesheet_style_export import (
        apply_view_styles, read_view_styles,
    )

    path = tmp_path / "scale.xlsx"
    pd.DataFrame({"金额": [10.0, 30.0, 20.0]}).to_excel(path, index=False)
    apply_view_styles(
        path, "Sheet1",
        conditional={0: {"type": "color_scale", "min": 10.0, "max": 30.0}},
        row_count=3,
    )

    styles = read_view_styles(path).get("conditional") or {}

    assert styles.get(0, {}).get("type") == "color_scale"
    assert styles[0]["min"] == 10.0
    assert styles[0]["max"] == 30.0


def test_unsupported_rule_is_not_fabricated(tmp_path):
    """视图渲染不了的规则（公式/图标集）**不读回**，也不编造。"""
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import FormulaRule

    from beeoffice.plugins.beesheet_style_export import read_view_styles

    path = tmp_path / "formula.xlsx"
    pd.DataFrame({"金额": [10.0, 30.0]}).to_excel(path, index=False)
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.conditional_formatting.add(
        "A2:A3", FormulaRule(formula=["A2>20"], fill=None)
    )
    workbook.save(path)

    styles = read_view_styles(path)

    assert not (styles.get("conditional") or {}), "把不支持的规则当成了数据条"


def test_conditional_read_back_reaches_view(qtbot, tmp_path):
    from beeoffice.plugins.beesheet import BeeSheetView

    target = tmp_path / "视图.xlsx"
    writer = BeeSheetView()
    qtbot.addWidget(writer)
    try:
        writer._set_status = lambda *a, **k: None
        writer.model.set_full_dataframe(
            pd.DataFrame({"名称": ["甲", "乙"], "金额": [3.0, 999.0]})
        )
        writer._conditional_formats = {
            1: {"type": "data_bars", "min": 3.0, "max": 999.0}
        }
        writer.set_current_file_path(target)
        assert writer.save() is True
    finally:
        writer.close()

    reader = BeeSheetView()
    qtbot.addWidget(reader)
    try:
        reader._set_status = lambda *a, **k: None
        reader._import_file_from_path(str(target))
        assert reader._conditional_formats.get(1, {}).get("type") == "data_bars"
        assert reader.model._conditional_formats_ref, "模型没拿到引用（渲染不会生效）"
    finally:
        reader.close()


# ==================== 2. markdown 模式的 B/I/U ====================

@pytest.fixture
def write_editor(qtbot):
    from beeoffice.plugins.beewrite import BeeWriteEditor

    editor = BeeWriteEditor()
    qtbot.addWidget(editor)
    editor._set_status = lambda *a, **k: None
    return editor


def _select(editor, start, end):
    from PySide6.QtGui import QTextCursor

    cursor = editor.markdown_editor.textCursor()
    cursor.setPosition(start)
    cursor.setPosition(end, QTextCursor.KeepAnchor)
    editor.markdown_editor.setTextCursor(cursor)


def test_markdown_bold_wraps_and_toggles(write_editor):

    editor = write_editor
    editor._switch_mode("markdown")
    editor.markdown_editor.setPlainText("加粗这段文字")
    _select(editor, 0, 4)

    editor._toggle_bold()
    assert editor.markdown_editor.toPlainText() == "**加粗这段**文字"

    # 再按一次应"取消"（选择仍在标记内部）
    editor._toggle_bold()
    assert editor.markdown_editor.toPlainText() == "加粗这段文字"


def test_markdown_italic_and_underline(write_editor):
    editor = write_editor
    editor._switch_mode("markdown")

    editor.markdown_editor.setPlainText("斜体")
    _select(editor, 0, 2)
    editor._toggle_italic()
    assert editor.markdown_editor.toPlainText() == "*斜体*"

    editor.markdown_editor.setPlainText("下划线")
    _select(editor, 0, 3)
    editor._toggle_underline()
    assert editor.markdown_editor.toPlainText() == "<u>下划线</u>"


def test_markdown_without_selection_inserts_markers_and_hints(write_editor):
    """没选中时插入空标记并**给提示**（原先按钮毫无反应、也没提示）。"""
    editor = write_editor
    statuses = []
    editor._set_status = lambda msg, t=0: statuses.append(msg)
    editor._switch_mode("markdown")
    editor.markdown_editor.setPlainText("")
    editor.markdown_editor.textCursor()

    editor._toggle_bold()

    # 插入的是**一对**空标记，光标停在中间（这样可以直接接着输入）
    assert editor.markdown_editor.toPlainText() == "****"
    assert editor.markdown_editor.textCursor().position() == 2, "光标应在标记中间"
    assert statuses, "没有给用户任何提示"


def test_rich_mode_still_uses_char_format(write_editor):
    """富文本模式的原有行为不能被改坏。"""
    editor = write_editor
    editor._switch_mode("rich_text")
    editor.rich_editor.setPlainText("富文本")
    cursor = editor.rich_editor.textCursor()
    cursor.select(cursor.SelectionType.Document)
    editor.rich_editor.setTextCursor(cursor)

    editor._toggle_bold()

    assert editor.rich_editor.textCursor().charFormat().fontWeight() > 50
    # 纯文本不应被插入任何 markdown 标记
    assert "**" not in editor.rich_editor.toPlainText()


# ==================== 3. 分页符 / 分节符不再污染文档 ====================

def test_breaks_leave_no_visible_text(write_editor):
    from beeoffice.plugins.beewrite_rich_editor import SECTION_MARK_PROPERTY

    editor = write_editor
    editor._switch_mode("rich_text")
    editor.rich_editor.setPlainText("第一段")

    editor.rich_editor.insert_page_break()
    editor.rich_editor.insert_section_break()

    text = editor.rich_editor.toPlainText()
    assert "分页符" not in text, "分页符的装饰文字还在"
    assert "分节符" not in text, "分节符的装饰文字还在"

    block = editor.rich_editor.document().firstBlock()
    marks = 0
    while block.isValid():
        if block.blockFormat().property(SECTION_MARK_PROPERTY):
            marks += 1
        block = block.next()
    assert marks == 1, "分节符的块属性标记没打上"


def test_section_break_becomes_real_docx_section(write_editor, tmp_path):
    """**核心验收**：分节符在 Word 里必须是真的"节"（可单独设页面）。"""
    from docx import Document

    editor = write_editor
    editor._switch_mode("rich_text")
    editor.rich_editor.setPlainText("第一段")
    editor.rich_editor.insert_section_break()

    target = tmp_path / "分节.docx"
    assert editor._write_docx(str(target)) is True

    document = Document(str(target))
    assert len(document.sections) == 2, f"节数不对：{len(document.sections)}"
    assert "分节符" not in "\n".join(p.text for p in document.paragraphs)


# ==================== 4. 拼写检查「全部替换」不再抛异常 ====================

class _Item:
    def __init__(self, word, start, end):
        self.word, self.start, self.end = word, start, end


def test_replace_all_survives_list_rebuild():
    """
    **回归**：`_replace_item()` 会重建列表，而旧实现随后还用
    `self.items.index(item)` 反查下标 → 拿错行或 `ValueError` 中断。

    这里故意用**会变的文本**与**不等长**的替换词（老测试用常量桩+等长词，
    所以永远测不到）。
    """
    from beeoffice.plugins.beewrite_spellcheck_dialog import (
        SpellCheckDialog,
    )

    class Harness:
        """只提供 replace_all 需要的最小环境，其余用真实逻辑拼装。"""

        def __init__(self):
            self.text = "苹果 香蕉 苹果"
            self.items = [_Item("苹果", 0, 2), _Item("香蕉", 3, 5)]
            self.applied = []

        # ---- 真实 replace_all 需要的东西 ----
        def _selected_replacement(self, row):
            return {"苹果": "apple", "香蕉": "banana"}[self.items[row].word]

        def _replace_item(self, item, replacement):
            # 模拟真实实现：替换（长度会变）→ 重取文本 → refresh() 重建列表
            self.apply_replace(item.start, item.end, replacement)
            self.text = self.text_provider()
            self.refresh()

        def apply_replace(self, start, end, replacement):
            self.applied.append((start, end, replacement))
            self.text = self.text[:start] + replacement + self.text[end:]

        def text_provider(self):
            return self.text

        def refresh(self):
            # 重建列表：**按当前文本重新扫描**（偏移是新的），旧对象作废
            rebuilt = []
            for word in ("苹果", "香蕉"):
                start = self.text.find(word)
                if start >= 0:
                    rebuilt.append(_Item(word, start, start + len(word)))
            self.items = rebuilt

    harness = Harness()
    # 绑定**真实的** replace_all，确保测的是产品代码而不是复刻品

    SpellCheckDialog.replace_all(harness)

    assert harness.applied, "一处都没替换"
    # 关键：替换发生在**正确的位置**。
    # 第一处用原始偏移 (0,2)；而第一处替换后文本长度变了（"苹果"→"apple"），
    # 第二处用的是**刷新后的新偏移 (6,8)** —— 这正是修复的要点：
    # 旧实现会在列表重建后还用旧对象/旧下标，替换到错位置或直接抛异常。
    # （原文 `苹果 香蕉 苹果` → 前两处被替换，第三处"苹果"保留 ——
    #  对话框的 items 是"每词一条"，所以"全部替换"目前实际是"每词替换一处"，
    #  这一点作为**发现**记录，不是本批要修的崩溃问题。）
    assert harness.applied[0][:2] == (0, 2)
    assert harness.applied[1][:2] == (6, 8), harness.applied
    assert harness.text == "apple banana 苹果", harness.text


# ==================== 5. 导出 Word 与另存为 Word 同一条路 ====================

def test_app_export_docx_prefers_rich_writer(qtbot, app, monkeypatch,
                                             tmp_path):
    from PySide6.QtWidgets import QFileDialog, QMessageBox

    from beeoffice.plugins.beewrite import BeeWriteEditor

    editor = BeeWriteEditor()
    qtbot.addWidget(editor)
    app.tab_manager.addTab(editor, "蜜写")
    app.tab_manager.setCurrentWidget(editor)

    target = tmp_path / "统一.docx"
    calls = []
    monkeypatch.setattr(editor, "_write_docx",
                        lambda path: calls.append(path) or True)
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName",
        staticmethod(lambda *a, **k: (str(target), "")),
    )
    monkeypatch.setattr(QMessageBox, "warning",
                        staticmethod(lambda *a, **k: None))

    app._export_docx()

    assert calls == [str(target)], (
        "导出 Word 没有走富文本路径（会丢页面设置/页眉页脚）"
    )
    editor.close()
