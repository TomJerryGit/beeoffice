# beeoffice/plugins/beewrite_export_mixin.py
"""
蜜写导出 / 保存 / 模板 Mixin

本次更新（P1-W5）：
  - 新增图片导出：_write_block_runs 检测 QTextImageFormat 并插入图片
  - 新增 _insert_picture_to_para 方法：px → Inches 换算
  - 新增表格列宽导出：_qtexttable_to_docx 应用 QTextTableFormat 列宽

P1-W4 已有：
  - docx 导出补齐段落对齐 / 行距 / 缩进 / 段前后
  - 字符级格式（粗体 / 斜体 / 字号 / 颜色）保留
  - 表格单元格对齐 / 表头加粗 / 背景色

v5.0.2 图片修复：
  - _write_md / _write_docx 保存成功后调用 self.set_current_file_path(path)
    使后续插入/粘贴的图片存到文档同目录的 media/ 而非 ~/media/
  - _insert_picture_to_para 遇到相对路径时，基于 _current_file_path 解析
  - _write_docx 保存前若富文本含相对路径图片，能正确定位

v5.0.2 增强：
  - _insert_picture_to_para 多基准目录兜底解析：
      1. 从 self._current_file_path 的父目录
      2. 从 self.get_current_file_dir() 方法（若存在）
      3. 从 self.rich_editor._current_file_dir
      4. 用户主目录
  - 图片找不到时记录所有尝试过的路径到日志
  - _rich_to_docx 在开始时打印文档图片统计（便于调试）

v5.0.4 优化：
  - _rich_to_docx 图片段落特殊处理：
      · 图片所在段落居中对齐（WD_ALIGN_PARAGRAPH.CENTER）
      · 段前段后各留 12pt 间距（space_before / space_after）
    目的：使图片与相邻表格有明显间隔，避免"贴着表格"的观感
  - 图片段落不再叠加 rich editor 的 block format，
    避免被覆盖为左对齐

v5.0.5 修复（本次）：
  - _write_docx / _write_md 增加 PermissionError 友好提示：
      当目标文件被 WPS / Word / 其他程序占用时，
      Windows 会拒绝写入并抛 PermissionError，
      此时弹出提示框，引导用户：
        · 关闭占用程序后重试
        · 另存为新文件
    避免用户看到底层 errno 13 错误而困惑
  - 保存失败时不再更新 _current_file_path / 标签标题：
      原实现在 _rich_to_docx 之后立即更新路径，
      若 _rich_to_docx 抛异常，路径已被更新为失败的文件，
      导致后续 Ctrl+S 行为异常（把状态污染）。
      修复后：只有文件真正写入成功才更新内部状态

已有功能保留：
  - save() 根据 _doc_kind 分派：word → .docx，markdown → .md
  - 保存成功后同步标签标题 / 窗口标题
  - 保存成功后通知主窗口"本次是自己保存的"，避免 fileChanged 误触发
  - _rich_to_docx 应用页面设置（纸张 / 方向 / 页边距）
  - _rich_to_docx 识别分页符（PageBreak_AlwaysBefore）
  - 首次保存记录路径，后续 Ctrl+S 直接写回同一文件
  - 「导出 Word」/「导出 PDF」作为独立的导出入口

依赖：
  - PySide6（QFileDialog / QMessageBox / QInputDialog / QTextCursor /
             QTextFormat / QTextBlockFormat / QTextImageFormat / QFont / Qt）
  - python-docx（延迟 import）
  - beeoffice.plugins.bee_export（延迟 import：export_to_pdf / export_to_docx）
  - beeoffice.plugins.beewrite_page_extras
        （build_watermark_html / build_page_number_css）
  - beeoffice.resources.template_manager.TemplateManager（延迟 import）
"""
import os
import sys
import shutil
import logging
import subprocess
from pathlib import Path

from PySide6.QtWidgets import (
    QFileDialog,
    QMessageBox,
    QInputDialog,
)
from PySide6.QtCore import Qt
from PySide6.QtGui import (
    QFont,
    QTextCursor,
    QTextFormat,
)

from beeoffice.plugins.beewrite_page_extras import (
    build_watermark_html,
    build_page_number_css,
)
from beeoffice.plugins.image_utils import (
    resolve_image_path,
    iter_image_formats,
    iter_block_images,
)
from beeoffice.utils import atomic_replace
from beeoffice.i18n import tr


# ==================== 常量 ====================
# 像素 → pt 换算（96 DPI 下，1px = 0.75pt）
_PX_TO_PT = 0.75

# QTextBlockFormat.LineHeightTypes 的 C++ 枚举值
_LINE_HEIGHT_SINGLE = 0
_LINE_HEIGHT_PROPORTIONAL = 1
_LINE_HEIGHT_FIXED = 2
_LINE_HEIGHT_MINIMUM = 3
_LINE_HEIGHT_LINE_DISTANCE = 4

# 像素 → 英寸（96 DPI）
_PX_TO_INCH = 1.0 / 96.0


# ==================== ExportMixin ====================
class ExportMixin:
    """蜜写导出 / 保存 / 模板 Mixin"""

    # ==================================================================
    # 统一保存入口（根据文档类型分派）
    # ==================================================================
    def save(self) -> bool:
        """
        保存当前文档：
          - _doc_kind == "word"     → 保存为 .docx
          - _doc_kind == "markdown" → 保存为 .md
        """
        self._save_canceled = False
        doc_kind = getattr(self, "_doc_kind", "markdown")
        if doc_kind == "word":
            return self._save_as_docx()
        return self._save_as_md()

    def save_as(self):
        """另存为（忽略当前路径，重新选择）"""
        self._save_canceled = False
        self._current_file_path = None
        self.save()

    # ==================================================================
    # 保存后同步标签标题 / 通知主窗口
    # ==================================================================
    def _update_tab_title(self, path: str):
        """保存成功后，把标签标题改为文件名，并同步窗口标题"""
        try:
            parent = self._find_main_window()
            if parent and hasattr(parent, "sync_tab_title"):
                parent.sync_tab_title(self, path)
        except Exception as e:
            logging.error(f"同步标签标题失败: {e}")

    def _notify_self_save(self, path: str):
        """通知主窗口：本次保存是自己触发的"""
        try:
            parent = self._find_main_window()
            if parent and hasattr(parent, "notify_self_save"):
                parent.notify_self_save(path)
        except Exception as e:
            logging.error(f"通知主窗口自己保存失败: {e}")

    # ==================================================================
    # 保存为 Markdown
    # ==================================================================
    def _save_as_md(self) -> bool:
        """保存为 .md（Markdown 源码）"""
        current = getattr(self, "_current_file_path", None)
        if current and current.lower().endswith(".md"):
            return self._write_md(current)

        default_name = "未命名.md"
        try:
            text = self.content().strip()
            if text:
                first_line = text.split('\n')[0]
                if first_line.startswith('#'):
                    t = first_line.lstrip('#').strip()[:50]
                    for c in '<>:"/\\|?*':
                        t = t.replace(c, '')
                    if t:
                        default_name = f"{t}.md"
        except Exception:
            pass

        path, _ = QFileDialog.getSaveFileName(
            self,
            "保存为 Markdown",
            default_name,
            "Markdown 文件 (*.md);;所有文件 (*)",
        )
        if not path:
            self._save_canceled = True
            return False

        if not path.lower().endswith(".md"):
            path = path + ".md"

        return self._write_md(path)

    def _write_md(self, path: str) -> bool:
        """
        实际写入 .md 文件（含批注 / 书签旁挂数据）

        v5.0.5 修复：
          - 增加 PermissionError 友好提示（文件被其他程序占用）
          - 保存失败时不更新 _current_file_path（避免状态污染）
        """
        try:
            # P0 止血：Markdown 也走"临时文件 + 原子替换" —— 写到一半失败不会
            # 把原文件截断。这里不用 atomic_write_text（它吞掉异常返回 bool），
            # 而是直接用 atomic_replace，让 PermissionError 保持原样抛给下面的
            # "文件被占用" 友好提示分支。
            with atomic_replace(path) as tmp_path:
                tmp_path.write_text(self.content(), encoding="utf-8")

            # 保存批注
            base = Path(path).with_suffix('')
            if (getattr(self, "_comment_manager", None)
                    and self._comment_manager.count() > 0):
                self._comment_manager.save_to_file(
                    str(base) + ".comments.json"
                )

            # 保存书签
            if (getattr(self, "_bookmark_manager", None)
                    and self._bookmark_manager._bookmarks):
                self._bookmark_manager.save_to_file(
                    str(base) + ".bookmarks.json"
                )

            # ---------- v5.0.2：同步文件路径到 rich_editor ----------
            # ★ v5.0.5：只有到这里才说明文件已写入成功，
            #          才更新内部状态（避免失败污染 _current_file_path）
            if hasattr(self, "set_current_file_path"):
                self.set_current_file_path(path)

            self._current_file_path = path
            self._update_tab_title(path)
            self._notify_self_save(path)

            self._set_status(
                f"✅ 已保存为 Markdown：{Path(path).name}", 3000
            )
            return True

        except PermissionError:
            # ---------- v5.0.5：文件被占用的友好提示 ----------
            logging.error(
                f"保存 Markdown 失败（文件被占用）: {path}"
            )
            QMessageBox.warning(
                self,
                "⚠️ 文件被占用",
                f"无法写入文件：\n{path}\n\n"
                f"该文件正在被其他程序打开。\n\n"
                f"请关闭占用该文件的程序（如 WPS、Word、记事本）后重试。",
            )
            return False

        except Exception as e:
            logging.error(f"保存 Markdown 失败: {e}")
            QMessageBox.warning(
                self, "保存失败", f"无法保存文件：\n{e}"
            )
            return False

    # ==================================================================
    # 读取批注 / 书签旁挂数据（优化规划 §2.6）
    # ==================================================================
    def load_md_sidecars(self, path: str) -> dict:
        """
        加载与 `.md` 同名的批注 / 书签旁挂数据。

        原缺陷（数据丢失）：`_write_md()` 会把批注写入
        `<base>.comments.json`、书签写入 `<base>.bookmarks.json`，
        但**打开文档时从不读回** —— `CommentManager.load_from_file()` 与
        `BookmarkManager.load_from_file()` 在整个代码库里没有任何调用点
        （已全仓 grep 确认）。结果是重开文档后批注书签全部消失，
        而数据其实还在磁盘上；用户会认为"保存没生效"或"软件把数据弄丢了"。

        :param path: 已打开的 Markdown 文件路径
        :return: {"comments": bool, "bookmarks": bool} 各自是否成功加载
        """
        loaded = {"comments": False, "bookmarks": False}
        if not path:
            return loaded

        base = str(Path(path).with_suffix(''))

        comment_mgr = getattr(self, "_comment_manager", None)
        if comment_mgr is not None:
            try:
                loaded["comments"] = bool(
                    comment_mgr.load_from_file(base + ".comments.json")
                )
            except Exception as e:
                logging.warning(f"载入批注失败: {e}")

        bookmark_mgr = getattr(self, "_bookmark_manager", None)
        if bookmark_mgr is not None:
            try:
                loaded["bookmarks"] = bool(
                    bookmark_mgr.load_from_file(base + ".bookmarks.json")
                )
            except Exception as e:
                logging.warning(f"载入书签失败: {e}")

        # 刷新批注面板（面板即使未显示也应拿到最新列表）
        refresh = getattr(self, "_refresh_comment_panel", None)
        if loaded["comments"] and callable(refresh):
            try:
                refresh()
            except Exception as e:
                logging.warning(f"刷新批注面板失败: {e}")

        if loaded["comments"] or loaded["bookmarks"]:
            logging.info(
                f"已载入旁挂数据: 批注={'是' if loaded['comments'] else '无'}, "
                f"书签={'是' if loaded['bookmarks'] else '无'}"
            )
        return loaded

    # ==================================================================
    # 保存为 DOCX
    # ==================================================================
    def _save_as_docx(self) -> bool:
        """保存为 .docx"""
        current = getattr(self, "_current_file_path", None)
        source = getattr(self, "_source_docx", None)

        # 优化规划 §1.3（数据完整性）：
        #   打开 .docx 后 _current_file_path 就是用户的原文件，因此 Ctrl+S 会走
        #   _rich_to_docx() 在原文件上"重新生成"一份文档。该重建保真度很低
        #   （超链接、列表编号、页眉页脚、脚注、域、东亚字体、单元格内多段
        #   全部丢失），且既无备份也无确认——属于不可逆的数据破坏。
        #   现在：当目标正好是本次打开的原文件时，先让用户明确选择。
        if (current and source
                and str(current) == str(source)
                and not getattr(self, "_docx_overwrite_ok", False)):
            choice = self._ask_docx_overwrite(current)
            if choice == "cancel":
                return False
            if choice == "save_as":
                current = None          # 落到下面的"另存为"分支
            else:                       # overwrite：先备份，失败则中止
                if not self._backup_before_overwrite(current):
                    return False
                self._docx_overwrite_ok = True

        if current and str(current).lower().endswith(".docx"):
            return self._write_docx(current)

        default_name = "未命名.docx"
        try:
            text = self.content().strip()
            if text:
                first_line = text.split('\n')[0].lstrip('#').strip()[:50]
                for c in '<>:"/\\|?*':
                    first_line = first_line.replace(c, '')
                if first_line:
                    default_name = f"{first_line}.docx"
        except Exception:
            pass

        path, _ = QFileDialog.getSaveFileName(
            self,
            "保存为 Word 文档",
            default_name,
            "Word (*.docx);;所有文件 (*)",
        )
        if not path:
            self._save_canceled = True
            return False

        if not path.lower().endswith(".docx"):
            path = path + ".docx"

        return self._write_docx(path)

    def _ask_docx_overwrite(self, path) -> str:
        """
        询问用户如何处理"保存到本次打开的原 .docx"。

        :return: "save_as" | "overwrite" | "cancel"
        """
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Warning)
        box.setWindowTitle("保存 Word 文档")
        box.setText(f"当前文档由 {Path(path).name} 打开。")
        box.setInformativeText(
            "直接覆盖会在这份原文件上重新生成文档，可能丢失超链接、列表编号、"
            "页眉页脚、脚注与中文字体设置。\n\n"
            "建议另存为新文件——原文件将保持完全不变。"
        )
        save_as_btn = box.addButton("另存为（推荐）", QMessageBox.AcceptRole)
        overwrite_btn = box.addButton(
            "覆盖原文件（先自动备份）", QMessageBox.DestructiveRole
        )
        box.addButton("取消", QMessageBox.RejectRole)
        box.setDefaultButton(save_as_btn)
        box.exec()

        clicked = box.clickedButton()
        if clicked is save_as_btn:
            return "save_as"
        if clicked is overwrite_btn:
            return "overwrite"
        return "cancel"

    def _backup_before_overwrite(self, path) -> bool:
        """
        覆盖原文件前先复制一份 <name>.docx.bak。

        备份失败时返回 False 并中止覆盖——宁可这次不保存，也不做不可逆的事。
        """
        try:
            src = Path(path)
            if not src.exists():
                return True
            backup = src.with_name(src.name + ".bak")
            shutil.copy2(str(src), str(backup))
            logging.info(f"覆盖前已备份原文件: {backup}")
            return True
        except Exception as e:
            logging.error(f"备份原文件失败，已取消覆盖: {e}")
            try:
                QMessageBox.warning(
                    self,
                    "无法备份",
                    f"无法备份原文件，为避免数据丢失已取消保存：\n{e}\n\n"
                    "请改用「另存为」。",
                )
            except Exception:
                pass
            return False

    def _write_docx(self, path: str) -> bool:
        """
        实际写入 .docx 文件（根据当前编辑模式选择转换方式）

        v5.0.5 修复：
          - 增加 PermissionError 友好提示（文件被 WPS/Word 占用）
          - 保存失败时不更新 _current_file_path（避免状态污染）
        """
        try:
            if self._edit_mode == "rich_text":
                self._rich_to_docx(path)
            else:
                from beeoffice.plugins.bee_export import export_to_docx
                content = self.get_rendered_content()
                # 传入页面设置（优化规划 §2.2）：原先 Markdown 路径不传，
                # 导出的纸张/页边距是 python-docx 模板默认值，与富文本路径不一致。
                # 同时传 base_dir，让相对路径的图片能被嵌入（§2.6）。
                src = getattr(self, "_current_file_path", None)
                base_dir = str(Path(src).parent) if src else None
                # 页眉页脚（保真度阶段 A）：这两个信息**不住在 Markdown 里**，
                # 由导入时读到的 extras 随文档保留（见 apply_docx_extras）。
                # 原先导出从不写页眉页脚，于是往返后部件直接消失。
                extras = (
                    self.get_docx_extras()
                    if hasattr(self, "get_docx_extras") else {}
                )
                if not export_to_docx(
                    content, path, self.get_page_settings(),
                    base_dir=base_dir,
                    header_text=extras.get("header_text"),
                    footer_text=extras.get("footer_text"),
                    header_xml=extras.get("header_xml", ""),
                    footer_xml=extras.get("footer_xml", ""),
                    header_footer_options=extras.get(
                        "header_footer_options"
                    ),
                    section_parts=extras.get("section_parts"),
                    table_merges=extras.get("table_merges"),
                    nested_tables=extras.get("nested_tables"),
                    sections=extras.get("sections"),
                    bookmarks=extras.get("bookmarks"),
                    paragraph_styles=extras.get("paragraph_styles"),
                    style_definitions=extras.get("style_definitions"),
                    page_number=getattr(self, "_page_number_config", None),
                    toc_field=getattr(self, "_docx_toc_field", False),
                    figure_toc=getattr(self, "_docx_figure_toc", False),
                ):
                    QMessageBox.warning(self, "失败", "无法生成 Word 文档")
                    return False

            # ---------- v5.0.2：同步文件路径到 rich_editor ----------
            # ★ v5.0.5：只有到这里才说明文件已写入成功，
            #          才更新内部状态（避免失败污染 _current_file_path）
            if hasattr(self, "set_current_file_path"):
                self.set_current_file_path(path)

            self._current_file_path = path
            self._update_tab_title(path)
            self._notify_self_save(path)

            self._set_status(
                f"✅ 已保存为 Word 文档：{Path(path).name}", 3000
            )
            return True

        except PermissionError:
            # ---------- v5.0.5：文件被占用的友好提示 ----------
            logging.error(
                f"保存 DOCX 失败（文件被占用）: {path}"
            )
            reply = QMessageBox.question(
                self,
                "⚠️ 文件被占用",
                f"无法写入文件：\n{path}\n\n"
                f"该文件正在被其他程序（如 WPS、Word）打开。\n\n"
                f"请选择操作：\n"
                f"• 「是」 → 关闭占用程序后重试\n"
                f"• 「否」 → 另存为新文件",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.Yes,
            )
            if reply == QMessageBox.No:
                # 引导用户另存为新文件
                try:
                    default_new = str(
                        Path(path).with_name(
                            Path(path).stem + "_副本" + Path(path).suffix
                        )
                    )
                except Exception:
                    default_new = "未命名_副本.docx"
                new_path, _ = QFileDialog.getSaveFileName(
                    self, "另存为", default_new,
                    "Word (*.docx);;所有文件 (*)",
                )
                if new_path:
                    if not new_path.lower().endswith(".docx"):
                        new_path += ".docx"
                    return self._write_docx(new_path)
            return False

        except Exception as e:
            import traceback
            traceback.print_exc()
            logging.error(f"保存 DOCX 失败: {e}")
            QMessageBox.warning(
                self, "保存失败", f"无法保存 Word 文档：\n{e}"
            )
            return False

    # ==================================================================
    # 导出 PDF
    # ==================================================================
    def export_pdf(self):
        """导出当前文档为 PDF"""
        path, _ = QFileDialog.getSaveFileName(
            self, "导出 PDF", "文档.pdf", "PDF (*.pdf)"
        )
        if not path:
            return

        page_settings = self.get_page_settings()

        try:
            from beeoffice.plugins.bee_export import export_to_pdf
            content = self.get_rendered_content()

            if page_settings:
                content = self._apply_page_settings_to_content(
                    content, page_settings
                )

            if export_to_pdf(content, path):
                QMessageBox.information(
                    self, "成功", f"已保存:\n{path}"
                )
            else:
                QMessageBox.warning(self, "失败", "需要 WeasyPrint")
        except Exception as e:
            logging.error(f"导出 PDF 失败: {e}")
            QMessageBox.warning(self, "导出失败", str(e))

    def export_html(self):
        """
        导出当前文档为独立 HTML 文件。

        优化规划 §2.6「补」：`bee_export.export_to_html` 早已实现且有测试覆盖，
        但蜜写侧一直没有调用方——菜单项走的是 `app._export_html` 的
        `hasattr(widget, "export_html")` 判断，永远不成立，于是点了没有任何反应。
        这里补上真正的实现，并把入口重新接上（菜单 / 命令面板 / 插件）。

        已知限制：不内联图片，图片仍以相对路径引用（与 Markdown 路径一致），
        分发时请连同 media/ 目录一起拷贝。
        """
        path, _ = QFileDialog.getSaveFileName(
            self,
            "导出 HTML",
            "文档.html",
            "HTML (*.html);;所有文件 (*)",
        )
        if not path:
            return

        if not path.lower().endswith((".html", ".htm")):
            path += ".html"

        try:
            from beeoffice.plugins.bee_export import export_to_html

            content = self.get_rendered_content()
            if export_to_html(content, path):
                self._set_status(
                    f"✅ 已导出 HTML：{Path(path).name}", 4000
                )
                QMessageBox.information(self, "成功", f"已保存:\n{path}")
            else:
                QMessageBox.warning(
                    self, "导出失败",
                    "无法生成 HTML。请确认已安装 markdown 依赖。",
                )
        except Exception as e:
            logging.error(f"导出 HTML 失败: {e}")
            QMessageBox.warning(self, "导出失败", str(e))

    def _apply_page_settings_to_content(self, content, page_settings):
        """把页面设置 + 水印 + 页码注入到 HTML 内容"""
        paper_sizes = {
            "A4": (210, 297),
            "A3": (297, 420),
            "A5": (148, 210),
            "Letter": (216, 279),
            "Legal": (216, 356),
        }
        w, h = paper_sizes.get(
            page_settings.get("paper", "A4"), (210, 297)
        )
        if page_settings.get("orientation") == "landscape":
            w, h = h, w

        # 页边距单位统一为 **mm**（优化规划 §2.2）。
        # 原实现对这四个值再乘 0.264583（px→mm 的换算），把 mm 当成了 px：
        # 默认 25.4mm 会被写成 6.7mm；而修复前的默认值 2.54 更是只剩下
        # 0.67mm。同一份 _page_settings 在 DOCX 路径里是直接用 Mm() 的
        # （即按 mm 解释），两条导出路径因此互相矛盾。
        mt = page_settings.get("margin_top", 25.4)
        mb = page_settings.get("margin_bottom", 25.4)
        ml = page_settings.get("margin_left", 31.7)
        mr = page_settings.get("margin_right", 31.7)

        page_css = f"""
        @page {{
            size: {w}mm {h}mm;
            margin: {mt:.1f}mm {mr:.1f}mm {mb:.1f}mm {ml:.1f}mm;
        }}
        """

        extras = self._build_export_extras()

        if "<head>" in content:
            return content.replace(
                "</head>",
                f"<style>{page_css}</style>{extras}</head>",
            )
        else:
            return f"""<!DOCTYPE html>
<html><head><meta charset="UTF-8">
<style>{page_css}</style>{extras}</head>
<body>{content}</body></html>"""

    def _build_export_extras(self):
        """构造水印 + 页码的 HTML / CSS 片段"""
        extras = ""
        if getattr(self, "_watermark_config", None):
            extras += build_watermark_html(self._watermark_config)
        if getattr(self, "_page_number_config", None):
            extras += build_page_number_css(self._page_number_config)
        return extras

    # ==================================================================
    # 导出 DOCX（独立入口，用户主动导出）
    # ==================================================================
    def spell_check(self):
        """
        拼写检查（**按需**，P1-B ★）。

        为什么不做"边打字边标红"：实测装上真英文词典后，对真实**中文技术文档**
        跑一遍有 **619 处误报**（`CentOS`/`mysql`/`yum`/`Xshell`/`VMware` 全被判错）。
        常驻波浪线会让用户满屏红字 —— 那是骚扰不是功能。
        改成用户主动触发 + 列表 + 「全部忽略 / 加入词典」，把控制权交给用户。
        """
        from beeoffice.plugins.beewrite_spellcheck_dialog import (
            SpellCheckDialog,
        )
        from beeoffice.plugins.spellcheck import SpellCheckerService

        service = getattr(self, "_spell_service", None)
        if service is None:
            service = SpellCheckerService()
            self._spell_service = service

        if not service.available:
            QMessageBox.information(
                self,
                "拼写检查不可用",
                f"{service.reason}\n\n"
                "说明：蜜蜂办公的拼写检查沿用**系统/本机**的英文词典，"
                "不会把任何数据传到网上。\n"
                "启用办法（任选其一）：\n"
                "· Windows 设置 → 时间和语言 → 语言 → 添加 English 语言包\n"
                "· 或用 pip 安装 pyspellchecker 后重启本程序\n\n"
                "（中文没有词形拼写检查，本功能只检查英文单词。）",
            )
            return

        # 取当前可见文本：富文本模式与 Markdown 模式分别取
        if getattr(self, "_edit_mode", "") == "rich_text":
            editor = self.rich_editor
        else:
            editor = getattr(self, "markdown_editor", None) or self.rich_editor
        text = editor.toPlainText()

        def apply_replace(start, end, replacement):
            from PySide6.QtGui import QTextCursor

            cursor = editor.textCursor()
            cursor.setPosition(max(0, start))
            cursor.setPosition(max(0, end), QTextCursor.MoveMode.KeepAnchor)
            cursor.insertText(replacement)

        dialog = SpellCheckDialog(
            self, service, text, apply_replace,
            on_locate=self._locate_in_editor,
            text_provider=lambda: editor.toPlainText(),
        )
        dialog.exec()

    def _locate_in_editor(self, start, end):
        """把编辑器光标选中到指定范围（双击列表项时用）。"""
        try:
            from PySide6.QtGui import QTextCursor

            mode = getattr(self, "_edit_mode", "")
            editor = self.rich_editor if mode == "rich_text" else (
                getattr(self, "markdown_editor", None) or self.rich_editor
            )
            cursor = editor.textCursor()
            cursor.setPosition(max(0, start))
            cursor.setPosition(max(0, end), QTextCursor.MoveMode.KeepAnchor)
            editor.setTextCursor(cursor)
            editor.setFocus()
        except Exception as e:  # noqa: BLE001 - 定位失败不影响检查
            logging.debug(f"[拼写检查] 定位失败: {e}")

    def _confirm_fidelity_losses(self, action: str = "导出") -> bool:
        """
        导出前提示"这份文件里哪些内容保不住"（P1-A）。

        背景：有些构造（脚注、批注、文本框、公式、嵌套表格、宏…）在
        导入→导出之间**必然**丢失或降级，而以往是**静默**发生的 ——
        用户要等导出后自己发现。竞品普遍也不提示。

        规则来自实测（见 `beeoffice/plugins/fidelity_report.py` 顶部说明），
        并且**经过"零误报"验收**：对我们已验证零丢失的真实文档不会弹窗。

        只对**真会丢内容**（high）或**有明确说明**（如旧格式/加密文档）的
        情况提示；纯"能力降级"（目录域、内容控件、修订痕迹——文字都在）
        不打断用户。

        :param action: 显示给用户的动作名（"导出"/"保存"）
        :return: True = 继续；False = 用户取消
        """
        report = getattr(self, "_fidelity_report", None)
        try:
            from beeoffice.plugins.fidelity_report import (
                format_report, should_prompt_before_export,
            )
        except Exception:  # noqa: BLE001 - 提示模块不可用不影响导出
            return True

        if not should_prompt_before_export(report):
            return True

        name = Path(report.path).name if report.path else "当前文档"
        detail = format_report(report, name)
        # 2026-10-21：区分"真会丢内容"与"仅说明性 note"。
        # 例：PPTX 导入报告 items=[]、summary="未发现无法保真的构造"，但 note 非空
        # （说明 PPTX 只支持导入）—— 此时不能再写"有内容无法完整保留"，自相矛盾。
        if getattr(report, "items", None):
            headline = tr("dialog.fidelity_loss_headline") + "："
        else:
            headline = tr("msg.hint") + "："
        reply = QMessageBox.warning(
            self,
            "保真度提示",
            f"{headline}\n\n{detail}\n\n"
            f"仍要继续{action}吗？",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.Yes,
        )
        if reply != QMessageBox.Yes:
            logging.info(f"[保真度提示] 用户取消{action}：{name}")
            return False
        logging.info(f"[保真度提示] 用户确认继续{action}：{name}")
        return True

    def export_docx(self):
        """导出当前文档为 DOCX（不改变 _current_file_path）"""
        if self._edit_mode == "rich_text":
            return self.export_rich_docx()

        if not self._confirm_fidelity_losses("导出"):
            return

        dn = "文档.docx"
        content_text = self.content().strip()
        fl = content_text.split('\n')[0] if content_text else ""
        if fl.startswith('#'):
            t = fl.lstrip('#').strip()[:50]
            for c in '<>:"/\\|?*':
                t = t.replace(c, '')
            if t:
                dn = f"{t}.docx"

        path, _ = QFileDialog.getSaveFileName(
            self, "导出 DOCX", dn,
            "Word (*.docx);;所有文件 (*)",
        )
        if not path:
            return

        try:
            from beeoffice.plugins.bee_export import export_to_docx
            if export_to_docx(self.get_rendered_content(), path):
                r = QMessageBox.information(
                    self, "成功",
                    f"已保存:\n{path}\n\n是否打开？",
                    QMessageBox.Yes | QMessageBox.No,
                )
                if r == QMessageBox.Yes:
                    self._open_file_externally(path)
            else:
                QMessageBox.warning(self, "失败", "无法生成")
        except PermissionError:
            logging.error(f"导出 DOCX 失败（文件被占用）: {path}")
            QMessageBox.warning(
                self,
                "⚠️ 文件被占用",
                f"无法写入文件：\n{path}\n\n"
                f"该文件正在被其他程序打开，请关闭后重试。",
            )
        except Exception as e:
            logging.error(f"导出 DOCX 失败: {e}")
            QMessageBox.warning(self, "导出失败", str(e))

    def export_rich_docx(self):
        """富文本模式导出 DOCX（保留表格结构 + 页面设置 + 分页符）"""
        if not self._confirm_fidelity_losses("导出"):
            return
        if self._edit_mode != "rich_text":
            QMessageBox.information(
                self, "提示", "请先切换到所见即所得模式"
            )
            return

        dn = "文档.docx"
        content = self.rich_editor.toPlainText().strip()
        if content:
            first_line = content.split('\n')[0]
            if first_line.startswith('#'):
                t = first_line.lstrip('#').strip()[:50]
                for c in '<>:"/\\|?*':
                    t = t.replace(c, '')
                if t:
                    dn = f"{t}.docx"

        path, _ = QFileDialog.getSaveFileName(
            self, "导出 DOCX（保留表格）", dn,
            "Word (*.docx);;所有文件 (*)",
        )
        if not path:
            return

        try:
            self._rich_to_docx(path)
            r = QMessageBox.information(
                self, "成功",
                f"已保存：\n{path}\n\n是否打开？",
                QMessageBox.Yes | QMessageBox.No,
            )
            if r == QMessageBox.Yes:
                self._open_file_externally(path)
        except PermissionError:
            logging.error(f"导出富文本 DOCX 失败（文件被占用）: {path}")
            QMessageBox.warning(
                self,
                "⚠️ 文件被占用",
                f"无法写入文件：\n{path}\n\n"
                f"该文件正在被其他程序打开，请关闭后重试。",
            )
        except Exception as e:
            import traceback
            traceback.print_exc()
            logging.error(f"导出富文本 DOCX 失败: {e}")
            QMessageBox.warning(self, "导出失败", str(e))

    # ==================================================================
    # 段落格式辅助：QTextBlockFormat → python-docx
    # ==================================================================
    def _apply_block_format_to_para(self, para, block_format):
        """
        从 QTextBlockFormat 提取段落格式并应用到 python-docx 段落。
        """
        try:
            from docx.shared import Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            return

        pf = para.paragraph_format

        # ---------- 对齐 ----------
        try:
            align = block_format.alignment()
            align_map = {
                Qt.AlignLeft: WD_ALIGN_PARAGRAPH.LEFT,
                Qt.AlignCenter: WD_ALIGN_PARAGRAPH.CENTER,
                Qt.AlignRight: WD_ALIGN_PARAGRAPH.RIGHT,
                Qt.AlignJustify: WD_ALIGN_PARAGRAPH.JUSTIFY,
            }
            if align in align_map:
                pf.alignment = align_map[align]
        except Exception as e:
            logging.debug(f"应用对齐失败: {e}")

        # ---------- 行距 ----------
        try:
            line_height = block_format.lineHeight()
            line_type = block_format.lineHeightType()
            if line_height > 0:
                if line_type == _LINE_HEIGHT_PROPORTIONAL:
                    pf.line_spacing = float(line_height) / 100.0
                elif line_type == _LINE_HEIGHT_FIXED:
                    pf.line_spacing = Pt(float(line_height))
        except Exception as e:
            logging.debug(f"应用行距失败: {e}")

        # ---------- 段前 / 段后 ----------
        try:
            top_margin = block_format.topMargin()
            if top_margin > 0:
                pf.space_before = Pt(top_margin * _PX_TO_PT)
            bottom_margin = block_format.bottomMargin()
            if bottom_margin > 0:
                pf.space_after = Pt(bottom_margin * _PX_TO_PT)
        except Exception as e:
            logging.debug(f"应用段前后失败: {e}")

        # ---------- 左右缩进 ----------
        try:
            left_margin = block_format.leftMargin()
            if left_margin > 0:
                pf.left_indent = Pt(left_margin * _PX_TO_PT)
            right_margin = block_format.rightMargin()
            if right_margin > 0:
                pf.right_indent = Pt(right_margin * _PX_TO_PT)
        except Exception as e:
            logging.debug(f"应用左右缩进失败: {e}")

        # ---------- 首行 / 悬挂缩进 ----------
        try:
            indent = block_format.textIndent()
            if indent != 0:
                pf.first_line_indent = Pt(indent * _PX_TO_PT)
        except Exception as e:
            logging.debug(f"应用首行缩进失败: {e}")

    def _apply_char_format_to_run(self, run, char_format):
        """从 QTextCharFormat 提取字符格式并应用到 python-docx Run"""
        try:
            from docx.shared import Pt, RGBColor
        except ImportError:
            return

        try:
            if char_format.fontWeight() == QFont.Bold:
                run.bold = True
            if char_format.fontItalic():
                run.italic = True
            if char_format.fontUnderline():
                run.underline = True
            if char_format.fontStrikeOut():
                run.font.strike = True
            size = char_format.fontPointSize()
            if size > 0:
                run.font.size = Pt(float(size))
            families = char_format.fontFamilies()
            if families:
                family = families[0] if isinstance(families, list) else families
                if family:
                    run.font.name = str(family)
            color = char_format.foreground()
            if color.style() != Qt.NoBrush:
                rgb = color.color()
                run.font.color.rgb = RGBColor(
                    rgb.red(), rgb.green(), rgb.blue()
                )
        except Exception as e:
            logging.debug(f"应用字符格式失败: {e}")

    # ==================================================================
    # 图片导出辅助
    # ==================================================================
    def _resolve_image_path_for_export(self, name: str):
        """
        解析图片路径（多基准目录兜底）。

        优化规划 §3.3：**候选来源**留在本类（它知道自己的文档状态），
        **解析算法**统一走 `image_utils.resolve_image_path()` —— 原先这里与
        `beewrite_rich_editor._insert_image_by_path()` 各有一份形态相同的实现。

        按优先级：
          1. 绝对路径：直接使用（不存在即失败，不落到下面）
          2. self._current_file_path 的父目录（最可靠）
          3. self.get_current_file_dir()（若存在）
          4. self.rich_editor._current_file_dir
          5. 用户主目录（由 resolve_image_path 兜底）

        :param name: 图片路径（来自 QTextImageFormat.name()）
        :return: (resolved_path_or_None, tried_paths_list)
        """
        candidates = []

        # 候选 1：self._current_file_path 的父目录
        try:
            current_file = getattr(self, "_current_file_path", None)
            if current_file:
                candidates.append(str(Path(current_file).parent))
        except Exception:
            pass

        # 候选 2：self.get_current_file_dir()（若实现）
        try:
            if hasattr(self, "get_current_file_dir"):
                d = self.get_current_file_dir()
                if d:
                    candidates.append(str(d))
        except Exception:
            pass

        # 候选 3：self.rich_editor._current_file_dir
        try:
            if hasattr(self, "rich_editor"):
                d = getattr(self.rich_editor, "_current_file_dir", None)
                if d:
                    candidates.append(str(d))
        except Exception:
            pass

        # 候选 4（home）由 resolve_image_path 内部兜底
        return resolve_image_path(name, candidates)

    def _insert_picture_to_para(self, para, img_format):
        """
        把 QTextImageFormat 表示的图片插入到 python-docx 段落。

        QTextImageFormat.name() 返回图片的路径或资源名：
          - 绝对路径（粘贴图片 / 插入图片常见）
          - 相对路径（相对于当前文档）
          - 资源 URL（qrc:// 少见）

        width/height 单位是 px（96 DPI），若为 0 表示使用原图尺寸。

        v5.0.2：相对路径解析失败时，基于 _current_file_path 的目录再尝试。
        v5.0.2 增强：多基准目录兜底解析 + 记录所有尝试过的路径。
        """
        try:
            from docx.shared import Inches
        except ImportError:
            return

        try:
            name = img_format.name()
            if not name:
                return

            # 处理 qrc 资源路径（跳过）
            if name.startswith("qrc:") or name.startswith(":/"):
                logging.debug(f"跳过 qrc 资源图片: {name}")
                return

            # v5.0.2 增强：多基准目录兜底解析
            img_path, tried = self._resolve_image_path_for_export(name)

            if img_path is None:
                logging.warning(
                    f"图片文件不存在，跳过: {name}\n"
                    f"尝试过的路径:\n  "
                    + "\n  ".join(tried)
                )
                return

            # 尺寸换算（px → Inches）
            width = img_format.width()
            height = img_format.height()

            run = para.add_run()
            if width > 0 and height > 0:
                run.add_picture(
                    str(img_path),
                    width=Inches(width * _PX_TO_INCH),
                    height=Inches(height * _PX_TO_INCH),
                )
            elif width > 0:
                run.add_picture(
                    str(img_path),
                    width=Inches(width * _PX_TO_INCH),
                )
            else:
                # 使用原图尺寸
                run.add_picture(str(img_path))

            logging.debug(
                f"[_insert_picture_to_para] 已插入图片: {img_path}"
            )
        except Exception as e:
            logging.error(f"插入图片失败: {e}")

    # ==================================================================
    # 富文本 → DOCX（核心转换）
    # ==================================================================
    def _add_docx_section(self, doc) -> bool:
        """
        在 DOCX 里开启**新的节**（真正的分节符）。

        Word 里的"节"可以各自设置页面大小/方向/页眉页脚 —— 这正是
        「§ 分节符」按钮承诺的能力（原先只插入一行装饰文字，Word 里没有节）。
        python-docx 不可用时**如实返回 False**（调用方跳过即可，不影响其余导出）。
        """
        try:
            from docx.enum.section import WD_SECTION

            doc.add_section(WD_SECTION.NEW_PAGE)
            return True
        except ImportError:
            logging.warning("[蜜写] 没有 python-docx，分节符降级为分页")
            return False
        except Exception as e:  # noqa: BLE001
            logging.warning(f"[蜜写] 插入分节失败（已跳过）: {e}")
            return False

    def _rich_to_docx(self, path):
        """
        把 QTextDocument 内容写到 DOCX。

        处理：
          - 页面设置（纸张 / 方向 / 页边距）
          - 分页符（PageBreak_AlwaysBefore）
          - 段落：标题 / 正文 + 对齐 / 行距 / 缩进 / 段前后
          - 段落内图片
          - 表格：单元格对齐 / 表头加粗 / 背景色 / 列宽

        v5.0.4：图片段落特殊处理——居中对齐 + 段前段后 12pt 间距。
        """
        from docx import Document

        doc = Document()

        # ========== 应用页面设置（优化规划 §2.2） ==========
        # 与 Markdown 导出路径共用同一实现（bee_export.apply_page_settings），
        # 单位统一为 mm，避免两条路径的纸张/页边距行为不一致。
        from beeoffice.plugins.bee_export import apply_page_settings
        apply_page_settings(doc, self.get_page_settings())

        # ========== 统计文档图片 ==========
        try:
            doc_has_img = self._count_document_images()
            logging.info(
                f"[_rich_to_docx] 开始导出，文档含 {doc_has_img} 张图片, "
                f"目标文件: {path}"
            )
        except Exception:
            pass

        # ========== 遍历文档 block ==========
        document = self.rich_editor.document()
        block = document.begin()
        while block.isValid():
            cursor = QTextCursor(block)
            table = cursor.currentTable()

            # ---------- 表格 ----------
            if table:
                self._qtexttable_to_docx(doc, table)
                last_cell = table.cellAt(
                    table.rows() - 1, table.columns() - 1
                )
                last_cursor = last_cell.lastCursorPosition()
                last_cursor.movePosition(QTextCursor.EndOfBlock)
                block = last_cursor.block().next()
                continue

            # ---------- 分节符（真正的 Word "节"）----------
            # P1：以前"§ 分节符"只插入可见装饰文字 `════ 分节符（新节开始） ════`，
            # 那行字会**原样进入 DOCX/PDF**，而 Word 里根本没有"节"。
            # 现在编辑器在块上打一个**用户属性**（不显示任何文字），
            # 导出时据此生成真正的 Word 节 —— 用户之后才能"只改这一节的页面设置"。
            bf = block.blockFormat()
            try:
                from beeoffice.plugins.beewrite_rich_editor import (
                    SECTION_MARK_PROPERTY,
                )

                if bf.property(SECTION_MARK_PROPERTY):
                    self._add_docx_section(doc)
            except Exception as e:  # noqa: BLE001
                logging.debug(f"分节符处理失败（跳过）: {e}")

            # ---------- 分页符 ----------
            try:
                if bf.pageBreakPolicy() & QTextFormat.PageBreak_AlwaysBefore:
                    doc.add_page_break()
            except Exception as e:
                logging.debug(f"分页符判断失败（跳过）: {e}")

            # ---------- 普通段落 / 标题（含图片） ----------
            has_text = bool(block.text().strip())
            has_image = self._block_has_image(block)

            if has_text or has_image:
                heading_level = bf.headingLevel()
                if heading_level > 0 and has_text and not has_image:
                    para = doc.add_heading(
                        block.text(), level=min(heading_level, 9)
                    )
                else:
                    para = doc.add_paragraph()
                    # 遍历 runs（保留字内格式 + 图片）
                    self._write_block_runs(para, block)

                # ---------- v5.0.4：图片段落特殊处理 ----------
                if has_image:
                    # 图片段落：居中 + 段前段后间距
                    # 目的：使图片与相邻表格有明显间隔，避免"贴着表格"的观感
                    try:
                        from docx.enum.text import WD_ALIGN_PARAGRAPH
                        from docx.shared import Pt
                        para.paragraph_format.alignment = (
                            WD_ALIGN_PARAGRAPH.CENTER
                        )
                        para.paragraph_format.space_before = Pt(12)
                        para.paragraph_format.space_after = Pt(12)
                    except Exception as e:
                        logging.debug(
                            f"应用图片段落格式失败: {e}"
                        )
                else:
                    # 普通段落：应用原 block format
                    self._apply_block_format_to_para(para, bf)

            block = block.next()

        # 保存（若文件被占用，会在此处抛 PermissionError）
        # P0 止血：先写同目录临时文件再原子替换（审计 §4.3 G2）。
        # 目标文件被占用时 os.replace 同样会抛 PermissionError，由上层友好提示。
        with atomic_replace(path) as tmp_path:
            doc.save(str(tmp_path))

    def _count_document_images(self) -> int:
        """
        统计富文本文档中的图片数量。

        优化规划 §3.3：改用 `image_utils.iter_image_formats()`（原先这里自己
        遍历 block/fragment）。

        :return: 图片 fragment 总数
        """
        try:
            return sum(
                1 for _ in iter_image_formats(self.rich_editor.document())
            )
        except Exception:
            return 0

    def _block_has_image(self, block):
        """判断 block 内是否包含图片 fragment（优化规划 §3.3：走共享实现）"""
        try:
            for _ in iter_block_images(block):
                return True
        except Exception:
            pass
        return False

    @staticmethod
    def _fragment_is_link(char_format) -> bool:
        """该 fragment 是否是超链接（Qt 把链接存成 anchor）。"""
        try:
            return bool(char_format.isAnchor() and char_format.anchorHref())
        except Exception:  # noqa: BLE001
            return False

    def _add_hyperlink(self, para, url: str, text: str,
                       char_format=None) -> bool:
        """
        往段落里写一个**真正的 Word 超链接**（可点击、带 URL）。

        python-docx 没有高层 API，需要手工拼 `w:hyperlink` + 外部关系。
        失败时返回 False，调用方会退回普通文本 run（**不丢文字**）。
        """
        try:
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            part = para.part
            r_id = part.relate_to(str(url), RT.HYPERLINK, is_external=True)
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), r_id)

            run = OxmlElement("w:r")
            run_props = OxmlElement("w:rPr")
            # Word 里链接的惯例样式：蓝色 + 下划线
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "0563C1")
            run_props.append(color)
            underline = OxmlElement("w:u")
            underline.set(qn("w:val"), "single")
            run_props.append(underline)
            run.append(run_props)

            text_element = OxmlElement("w:t")
            text_element.set(qn("xml:space"), "preserve")
            text_element.text = str(text)
            run.append(text_element)

            hyperlink.append(run)
            para._p.append(hyperlink)
            return True
        except Exception as e:  # noqa: BLE001
            logging.debug(f"[蜜写] 写入超链接失败（退回普通文本）: {e}")
            return False

    def _write_block_runs(self, para, block):
        """
        将 block 内的字符片段（按 run）写入段落，保留粗体 / 斜体 /
        字号 / 颜色等字符格式，并处理图片 fragment。
        """
        it = block.begin()
        has_content = False
        while not it.atEnd():
            fragment = it.fragment()
            if fragment.isValid():
                fmt = fragment.charFormat()

                # ---------- 图片 fragment ----------
                if fmt.isImageFormat():
                    img_fmt = fmt.toImageFormat()
                    self._insert_picture_to_para(para, img_fmt)
                    has_content = True
                    it += 1
                    continue

                # ---------- 普通文本 fragment ----------
                frag_text = fragment.text()
                # 超链接：审计 §2.13 —— 富文本→DOCX 原先只处理图片与文本 run，
                # 链接的 URL **整条丢掉**（文字还在、点不动）。
                # 现在写成真正的 Word 超链接（外部关系 + w:hyperlink）。
                if frag_text and self._fragment_is_link(fmt):
                    if self._add_hyperlink(para, fmt.anchorHref(),
                                           frag_text, fmt):
                        has_content = True
                        it += 1
                        continue
                if frag_text:
                    run = para.add_run(frag_text)
                    self._apply_char_format_to_run(run, fmt)
                    has_content = True
            it += 1

        # 若没有 fragment（罕见），退回到纯文本
        if not has_content:
            para.add_run(block.text())

    # ==================================================================
    # 表格 → DOCX
    # ==================================================================
    def _qtexttable_to_docx(self, doc, qtable):
        """
        将 QTextTable 转换为 python-docx 表格。

        处理：
          - 表头加粗
          - 单元格背景色
          - 单元格段落对齐（来自 cell 的 blockFormat）
          - 单元格字符格式（来自 fragment 的 charFormat）
          - 列宽（来自 QTextTableFormat 的 columnWidthConstraints）
          - 单元格内图片（v5.0.2 增强）
        """
        try:
            # 只导入本段用到的名字（Pt / RGBColor 未使用；第十一批 F401 清理核实）
            from docx.shared import Cm
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            from PySide6.QtGui import QTextLength
        except ImportError:
            return

        rows = qtable.rows()
        cols = qtable.columns()

        doc_table = doc.add_table(rows=rows, cols=cols)
        doc_table.style = 'Table Grid'

        # ---------- 应用列宽 ----------
        try:
            table_fmt = qtable.format()
            constraints = table_fmt.columnWidthConstraints()
            if constraints and len(constraints) == cols:
                for ci, length in enumerate(constraints):
                    try:
                        if length.type() == QTextLength.FixedLength:
                            # 像素 → 厘米：1 px = 0.026458 cm（96 DPI）
                            width_px = length.rawValue()
                            width_cm = width_px * 2.54 / 96.0
                            if width_cm > 0:
                                doc_table.columns[ci].width = Cm(width_cm)
                    except Exception as e:
                        logging.debug(f"列 {ci} 宽度设置失败: {e}")
        except Exception as e:
            logging.debug(f"应用列宽失败: {e}")

        # ---------- 逐单元格处理 ----------
        for r in range(rows):
            for c in range(cols):
                cell = qtable.cellAt(r, c)
                cell_cursor = cell.firstCursorPosition()
                cell_block = cell_cursor.block()

                docx_cell = doc_table.cell(r, c)
                # ⚠️ 审计 §2.13：原先是 `docx_cell.text = cell_block.text()`
                # —— **只取首段**，单元格里的多段落全被压成一段。
                # 现在按 QTextDocument 里该单元格的**每个块**写一个段落，
                # 并复用 `_write_block_runs`（顺带获得超链接/图片支持）。
                try:
                    end_block = cell.lastCursorPosition().block()
                    end_position = end_block.position()
                    block_iter = cell_block
                    para_index = 0
                    while block_iter.isValid() \
                            and block_iter.position() <= end_position:
                        if para_index == 0:
                            docx_para = docx_cell.paragraphs[0]
                        else:
                            docx_para = docx_cell.add_paragraph()
                        self._write_block_runs(docx_para, block_iter)
                        para_index += 1
                        if block_iter == end_block:
                            break
                        block_iter = block_iter.next()
                    if para_index == 0:
                        docx_cell.text = ""      # 空单元格（保持一个空段落）
                except Exception as e:  # noqa: BLE001
                    logging.warning(f"[蜜写] 单元格多段落导出失败，退回纯文本: {e}")
                    docx_cell.text = cell_block.text()

                # ---------- 段落对齐（每个 cell 的段落） ----------
                try:
                    cell_bf = cell_block.blockFormat()
                    cell_align = cell_bf.alignment()
                    align_map = {
                        Qt.AlignLeft: WD_ALIGN_PARAGRAPH.LEFT,
                        Qt.AlignCenter: WD_ALIGN_PARAGRAPH.CENTER,
                        Qt.AlignRight: WD_ALIGN_PARAGRAPH.RIGHT,
                        Qt.AlignJustify: WD_ALIGN_PARAGRAPH.JUSTIFY,
                    }
                    if cell_align in align_map:
                        for paragraph in docx_cell.paragraphs:
                            paragraph.alignment = align_map[cell_align]
                except Exception as e:
                    logging.debug(f"单元格对齐失败: {e}")

                # ---------- 字符格式（粗体等） ----------
                try:
                    char_fmt = cell_block.charFormat()
                    if char_fmt.fontWeight() == QFont.Bold:
                        for paragraph in docx_cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
                except Exception:
                    pass

                # ---------- 表头加粗 ----------
                if r == 0:
                    for paragraph in docx_cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True

                # ---------- 单元格背景色 ----------
                char_fmt = cell_block.charFormat()
                bg = char_fmt.background()
                if bg.style() != Qt.NoBrush:
                    try:
                        shd = OxmlElement('w:shd')
                        shd.set(
                            qn('w:fill'),
                            bg.color().name().lstrip('#'),
                        )
                        docx_cell._tc.get_or_add_tcPr().append(shd)
                    except Exception as e:
                        logging.warning(f"设置单元格背景色失败: {e}")

                # ---------- v5.0.2 增强：单元格内图片 ----------
                try:
                    if self._block_has_image(cell_block):
                        target_para = (
                            docx_cell.paragraphs[0]
                            if docx_cell.paragraphs
                            else docx_cell.add_paragraph()
                        )
                        # 优化规划 §3.3：图片遍历统一走 iter_block_images()
                        for img_fmt in iter_block_images(cell_block):
                            self._insert_picture_to_para(
                                target_para, img_fmt
                            )
                except Exception as e:
                    logging.warning(f"处理单元格图片失败: {e}")

    def _open_file_externally(self, path):
        """用系统默认程序打开导出的文件"""
        try:
            if sys.platform == 'win32':
                os.startfile(path)
            elif sys.platform == 'darwin':
                subprocess.run(['open', path])
            else:
                subprocess.run(['xdg-open', path])
        except Exception as e:
            logging.error(f"打开文件失败 {path}: {e}")

    # ==================================================================
    # 保存为模板
    # ==================================================================
    def save_as_template(self):
        """把当前文档内容保存为模板"""
        n, ok = QInputDialog.getText(self, "保存为模板", "模板名称:")
        if not ok or not n.strip():
            return
        try:
            from beeoffice.resources.template_manager import TemplateManager
            tm = TemplateManager()
            if tm.save_as_template(n.strip(), self.content()):
                self._set_status("模板已保存", 3000)
            else:
                QMessageBox.warning(self, "失败", "保存模板失败")
        except Exception as e:
            logging.error(f"保存模板失败: {e}")
            QMessageBox.warning(self, "失败", f"保存模板失败：\n{e}")