# tests/fidelity/test_p1_freeze_chart_names_author.py
"""
第六批：**冻结窗格 / 图表 / 命名区域 / 批注作者** 四项落盘。

## 先测量，再动手（本项目的固定动作）
我先造了一个同时带这四种特性的真实工作簿，走蜜算的写路径测了一遍，结果：
* 经**工作簿路径**（`read_workbook` + `write_workbook`，只改 `.value`）
  **四项全部天然保留** —— 文档没说假话；
* 我又专门确认了第四批加的"写盘后后处理"（`load_workbook` + 改 + `save`）
  会不会把图表弄丢（openpyxl 有"读不回图表"的坏名声）：**实测图表部件、
  绘图部件、表→绘图关系、绘图→图表关系、图表标题全在** ✓ 没有回归。

所以真实缺口是另外三处，本批就修它们：
1. **读侧**：命名区域/冻结状态/批注作者从不读进视图 → 定位对话框看不到名称、
   状态栏不显示冻结、作者信息全丢；
2. **`df.to_excel` 新路径**（CSV 导入/新建/透视）没有任何源结构，视图里新设的
   命名区域/冻结/图表都没地方去；
3. **图表**：应用内是 **matplotlib 预览**（弹窗 + 存 PNG），
   **从来没有进过 xlsx**；以及 ⚠️ **我在第四批引入的回归**：
   `_apply_comments` 把作者**硬编码成 "BeeOffice"**，导入带作者的批注再保存
   会**覆盖原作者**。

## 另一处必须做的换算
应用的 `_named_ranges` 存的是 **A1 风格、行号 = 视图行 + 1**（见"从当前选区填充"），
而**文件第 1 行是表头**、数据从第 2 行开始 → 两个方向都要整体平移 1 行，
否则保存后的命名区域会**指错行**。
"""
import sys
import zipfile
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def _load(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    return workbook, workbook[workbook.sheetnames[0]]


def _chart_parts(path) -> int:
    with zipfile.ZipFile(path) as archive:
        return sum(1 for n in archive.namelist() if n.startswith("xl/charts/"))


@pytest.fixture
def sheet_view(qtbot):
    from PySide6.QtWidgets import QMessageBox

    from beeoffice.plugins.beesheet import BeeSheetView

    view = BeeSheetView()
    qtbot.addWidget(view)
    view.model.set_full_dataframe(
        pd.DataFrame({"月份": ["1月", "2月", "3月"], "销量": [10, 20, 15]})
    )
    view._set_status = lambda *a, **k: None
    QMessageBox.information = staticmethod(lambda *a, **k: None)
    QMessageBox.warning = staticmethod(lambda *a, **k: None)
    return view


# ==================== 1. 行号换算（命名区域的前提）====================

def test_a1_row_shift_roundtrip():
    from beeoffice.plugins.beesheet_style_export import _shift_a1_rows

    assert _shift_a1_rows("A1", +1) == "A2"
    assert _shift_a1_rows("B2:B4", +1) == "B3:B5"
    assert _shift_a1_rows("B3:B5", -1) == "B2:B4"
    assert _shift_a1_rows("AA10:AB12", +1) == "AA11:AB13"


# ==================== 2. 命名区域 ====================

def test_named_range_shifts_rows_and_reads_back(sheet_view, qtbot, tmp_path):
    target = tmp_path / "名称.xlsx"
    sheet_view._named_ranges = {"销量区域": "B2:B4"}
    sheet_view.set_current_file_path(target)

    assert sheet_view.save() is True

    workbook, _ = _load(target)
    # 视图 B2:B4（行号 = 视图行+1）→ 文件 B3:B5（表头占第 1 行）
    assert workbook.defined_names["销量区域"].attr_text.endswith("B3:B5")

    from beeoffice.plugins.beesheet import BeeSheetView

    reader = BeeSheetView()
    qtbot.addWidget(reader)
    reader._set_status = lambda *a, **k: None
    try:
        reader._import_file_from_path(str(target))
        assert reader._named_ranges.get("销量区域") == "B2:B4"
    finally:
        reader.close()


# ==================== 3. 冻结窗格 ====================

def test_freeze_panes_written_and_read_back(sheet_view, qtbot, tmp_path):
    target = tmp_path / "冻结.xlsx"
    sheet_view._freeze_panes = "A2"
    sheet_view.set_current_file_path(target)

    sheet_view.save()

    _, ws = _load(target)
    assert ws.freeze_panes == "A2"

    from beeoffice.plugins.beesheet import BeeSheetView

    reader = BeeSheetView()
    qtbot.addWidget(reader)
    reader._set_status = lambda *a, **k: None
    try:
        reader._import_file_from_path(str(target))
        assert reader._freeze_panes == "A2"
        # 状态栏/信息面板用的旧字段也要跟着有值
        assert reader._frozen_top == "A2"
    finally:
        reader.close()


def test_file_without_freeze_stays_without(sheet_view, tmp_path):
    target = tmp_path / "无冻结.xlsx"
    sheet_view._freeze_panes = None
    sheet_view.set_current_file_path(target)
    sheet_view.save()

    _, ws = _load(target)
    assert ws.freeze_panes is None


# ==================== 4. 批注作者 ====================

def test_comment_author_is_written_and_read_back(sheet_view, qtbot, tmp_path):
    target = tmp_path / "作者.xlsx"
    sheet_view._cell_comments = {(0, 0): "这是备注"}
    sheet_view._comment_authors = {(0, 0): "张三"}
    sheet_view.set_current_file_path(target)

    sheet_view.save()

    _, ws = _load(target)
    assert ws["A2"].comment.text == "这是备注"
    assert ws["A2"].comment.author == "张三"

    from beeoffice.plugins.beesheet import BeeSheetView

    reader = BeeSheetView()
    qtbot.addWidget(reader)
    reader._set_status = lambda *a, **k: None
    try:
        reader._import_file_from_path(str(target))
        assert reader._comment_authors.get((0, 0)) == "张三"
    finally:
        reader.close()


def test_existing_comment_author_is_not_overwritten(qtbot, tmp_path):
    """
    **本批最重要的回归**：导入一份作者是"李四"的批注 → 原样保存 →
    作者必须还是"李四"。

    第四批的 `_apply_comments` 把作者硬编码成 "BeeOffice"，
    会在这种情况下**覆盖原作者**；现在的规则是
    "内容没变就别碰这个批注"。
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    from beeoffice.plugins.beesheet import BeeSheetView

    source = tmp_path / "原作者.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["月份", "销量"])
    sheet.append(["1月", 10])
    sheet["A2"].comment = Comment("原有备注", "李四")
    workbook.save(source)

    view = BeeSheetView()
    qtbot.addWidget(view)
    try:
        view._set_status = lambda *a, **k: None
        view._import_file_from_path(str(source))
        assert view._cell_comments.get((0, 0)) == "原有备注"
        assert view._comment_authors.get((0, 0)) == "李四"

        view.set_current_file_path(source)
        assert view.save() is True

        _, ws = _load(source)
        assert ws["A2"].comment.author == "李四", "原作者被覆盖了！"
        assert ws["A2"].comment.text == "原有备注"
    finally:
        view.close()


def test_new_comment_uses_settings_user_name(qtbot, monkeypatch):
    """新批注的作者来自设置的用户名（没设就留空，绝不硬编码假名）。"""
    from beeoffice.plugins.beesheet import BeeSheetView
    from beeoffice import settings as settings_module

    class FakeSettings:
        user_name = "王五"

    monkeypatch.setattr(settings_module, "AppSettings", lambda *a, **k: FakeSettings())

    view = BeeSheetView()
    qtbot.addWidget(view)
    try:
        assert view._comment_author_name() == "王五"
    finally:
        view.close()


# ==================== 5. 图表 ====================

def test_chart_spec_is_written_as_native_excel_chart(sheet_view, tmp_path):
    target = tmp_path / "图表.xlsx"
    sheet_view._charts = [{
        "type": "bar", "rows": (0, 2), "cols": (0, 1), "title": "销量图",
    }]
    sheet_view.set_current_file_path(target)

    assert sheet_view.save() is True
    assert _chart_parts(target) == 1, "图表没有写成原生 Excel 图表"
    # 规格**保留**但标记为已写：重建路径（df.to_excel）下次仍要全部重写，
    # 而工作簿路径靠这个标记避免叠加
    assert sheet_view._charts[0].get("written") is True


def test_unsupported_chart_type_is_skipped(sheet_view, tmp_path):
    target = tmp_path / "不支持.xlsx"
    sheet_view._charts = [{"type": "sunburst", "rows": (0, 2), "cols": (0, 1)}]
    sheet_view.set_current_file_path(target)

    sheet_view.save()

    assert _chart_parts(target) == 0, "不支持的类型不该写坏文件"


def test_second_save_keeps_chart_on_rebuild_path(sheet_view, tmp_path):
    """
    **重建路径**（没有源工作簿 → `df.to_excel`）保存两次：图表必须**还在**。

    我第一版是"写成功就清空规格"，结果这条路径第二次保存把文件整个重建、
    图表**直接消失**（比重复更糟）—— 被这条测试当场抓到。
    """
    target = tmp_path / "重建不丢.xlsx"
    sheet_view._charts = [{"type": "line", "rows": (0, 2), "cols": (0, 1)}]
    sheet_view.set_current_file_path(target)

    sheet_view.save()
    assert _chart_parts(target) == 1

    sheet_view.save()
    assert _chart_parts(target) == 1, "重建路径第二次保存把图表弄丢了"


def test_second_save_does_not_duplicate_chart_on_workbook_path(qtbot, tmp_path):
    """
    **工作簿路径**（导入的 xlsx）保存两次：只能有**一张**图。

    这条路径每次从源文件重开，上次写的图表天然还在 →
    若不看 `written` 标记就会叠加成两张。
    """
    from beeoffice.plugins.beesheet import BeeSheetView

    source = tmp_path / "源.xlsx"
    pd.DataFrame({"月份": ["1月", "2月", "3月"], "销量": [10, 20, 15]}).to_excel(
        source, index=False
    )

    view = BeeSheetView()
    qtbot.addWidget(view)
    try:
        view._set_status = lambda *a, **k: None
        view._import_file_from_path(str(source))
        assert view._workbook is not None, "没走工作簿路径，测试前提不成立"

        view._charts = [{"type": "pie", "rows": (0, 2), "cols": (0, 1)}]
        view.set_current_file_path(source)
        assert view.save() is True
        assert _chart_parts(source) == 1

        assert view.save() is True
        assert _chart_parts(source) == 1, "工作簿路径把图表重复写入了"
    finally:
        view.close()


def test_generate_chart_records_spec(sheet_view, monkeypatch):
    """应用内生成图表时要记下规格（原先只弹预览、什么都没留）。"""
    from PySide6.QtCore import QItemSelection, QItemSelectionModel

    selection = QItemSelection(sheet_view.model.index(0, 0),
                               sheet_view.model.index(2, 1))
    sheet_view.table_view.selectionModel().select(
        selection, QItemSelectionModel.SelectionFlag.Select
    )
    sheet_view._set_status = lambda *a, **k: None

    sheet_view._charts = []
    sheet_view._generate_chart("bar")

    assert sheet_view._charts, "生成图表后没有记录规格"
    spec = sheet_view._charts[0]
    assert spec["type"] == "bar"
    assert spec["rows"] == (0, 2)
    assert spec["cols"] == (0, 1)


# ==================== 6. 数据安全 ====================

def test_all_four_keep_data(sheet_view, tmp_path):
    target = tmp_path / "数据安全.xlsx"
    sheet_view._named_ranges = {"销量区域": "B2:B4"}
    sheet_view._freeze_panes = "A2"
    sheet_view._cell_comments = {(0, 0): "备注"}
    sheet_view._comment_authors = {(0, 0): "张三"}
    sheet_view._charts = [{"type": "pie", "rows": (0, 2), "cols": (0, 1)}]
    sheet_view.set_current_file_path(target)

    sheet_view.save()

    _, ws = _load(target)
    assert ws["A2"].value == "1月"
    assert ws["B2"].value == 10
    assert ws["B4"].value == 15
