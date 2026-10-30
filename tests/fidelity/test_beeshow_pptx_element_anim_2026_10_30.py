"""蜜演 PPTX 逐元素动画（实验：单元素淡入）。"""
from pptx import Presentation

from beeoffice.plugins.beeshow_pptx_engine import (
    MeasuredElement,
    MeasuredSlide,
    PPTXGenerator,
    add_shape_animations,
    add_shape_animations_click,
    add_shape_fade,
    selector_matches,
)


def test_selector_matches():
    el = MeasuredElement(tag="h1", element_id="title", class_name="hero big")
    assert selector_matches(el, "h1")
    assert selector_matches(el, "#title")
    assert selector_matches(el, ".hero")
    assert selector_matches(el, "h1.hero")
    assert not selector_matches(el, ".missing")


def test_add_shape_fade_xml():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(100, 100, 300, 80)
    assert add_shape_fade(slide, shape, 500, 0) is True
    xml = slide._element.xml
    assert "p:timing" in xml
    assert 'filter="fade"' in xml
    assert f'spid="{shape.shape_id}"' in xml


def test_add_shape_animations_sequence():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    a = slide.shapes.add_textbox(100, 100, 300, 80)
    b = slide.shapes.add_textbox(100, 220, 300, 80)
    assert add_shape_animations(slide, [
        (a, "fade_in", 400, 0),
        (b, "fade_out", 300, 100),
    ]) is True
    xml = slide._element.xml
    assert xml.count("<p:animEffect ") == 2
    assert 'transition="in"' in xml
    assert 'transition="out"' in xml


def test_generate_element_fade(tmp_path):
    ms = MeasuredSlide(elements=[
        MeasuredElement(
            tag="h1", text="标题", x=100, y=100,
            width=400, height=60, font_size=28,
        )
    ])
    out = tmp_path / "anim.pptx"
    gen = PPTXGenerator()
    assert gen.generate(
        [ms], str(out),
        element_animations=[{
            "pptx_element": True,
            "entries": [{
                "selector": "h1", "effect": "fade_in",
                "duration": 500, "delay": 0,
            }],
        }],
    ) is True
    reopened = Presentation(str(out))
    assert "p:timing" in reopened.slides[0]._element.xml


def test_slide_zoom_effects():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    a = slide.shapes.add_textbox(10, 10, 200, 50)
    b = slide.shapes.add_textbox(10, 100, 200, 50)
    assert add_shape_animations(slide, [
        (a, "slide_in", 400, 0), (b, "zoom_in", 300, 100),
    ]) is True
    xml = slide._element.xml
    assert "slide(fromLeft)" in xml
    assert "zoom(in)" in xml


def test_click_each_timeline():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    a = slide.shapes.add_textbox(10, 10, 200, 50)
    b = slide.shapes.add_textbox(10, 100, 200, 50)
    assert add_shape_animations_click(slide, [
        (a, "fade_in", 400, 0), (b, "fade_out", 300, 0),
    ]) is True
    xml = slide._element.xml
    assert xml.count('delay="indefinite"') == 2
    assert xml.count("<p:animEffect ") == 2
