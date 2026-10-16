# tests/fidelity/test_beesheet_workbook.py
"""
保真度阶段 A 第一刀：蜜算的「多工作表 + 公式」读写。

背景
----
蜜算原先是 `pd.read_excel(...)`（只读首张表）+ `df.to_excel(...)`，
实测 **9 项特征丢 8 项**：第二张工作表直接消失、公式全丢。
本模块把"工作簿"作为显式对象承载全部工作表与每格公式。

重点测什么
----------
1. **公式坐标换算**（最容易错的地方）：openpyxl 坐标 1 起、且第 1 行是表头，
   DataFrame 位置 0 起 —— 差一格的错位会让公式写到错误的单元格上。
2. **往返**：读进来再写出去，工作表数与公式数不变。
3. 边界：空表、无公式、工作表名需要净化、NaN 不能变成 "nan" 字符串。
"""
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from beeoffice.plugins.beesheet_workbook import (   # noqa: E402
    SheetData,
    WorkbookData,
    read_workbook,
    sync_active_sheet,
    write_workbook,
)


def _make_xlsx(path: Path, sheets: dict = None) -> Path:
    """用 openpyxl 直接构造工作簿（含公式），作为测试输入。"""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    sheets = sheets or {
        "明细": [
            ["产品", "数量", "单价", "金额"],
            ["甲", 10, 12.5, "=B2*C2"],
            ["乙", 80, 3.2, "=B3*C3"],
        ],
        "汇总": [
            ["项目", "值"],
            ["合计", "=SUM(明细!D2:D3)"],
        ],
    }
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=value)
    wb.save(str(path))
    return path


# ==================== 1. 公式坐标换算 ====================

def test_formula_coordinates_are_df_based(tmp_path):
    """
    核心：工作表 (R,C) → df (R-2, C-1)。

    "明细" 表：
      工作表 D2 = "=B2*C2"  →  df (0, 3)
      工作表 D3 = "=B3*C3"  →  df (1, 3)
    """
    src = _make_xlsx(tmp_path / "in.xlsx")

    wb = read_workbook(src)

    detail = wb.get("明细")
    assert detail is not None
    assert detail.formulas == {(0, 3): "=B2*C2", (1, 3): "=B3*C3"}, (
        f"公式坐标换算错误: {detail.formulas}"
    )


def test_header_row_is_not_treated_as_data(tmp_path):
    """第 1 行是表头，不应产生 df 行号 -1 的公式。"""
    src = _make_xlsx(tmp_path / "in.xlsx", {
        "表": [["A", "B"], ["=A2", 2]],
    })

    formulas = read_workbook(src).get("表").formulas

    assert all(r >= 0 for r, _ in formulas), f"出现负行号: {formulas}"
    assert (0, 0) in formulas, formulas


# ==================== 2. 读：多工作表 + 公式 ====================

def test_reads_all_sheets(tmp_path):
    """修复「第二张工作表消失」。"""
    src = _make_xlsx(tmp_path / "in.xlsx")

    wb = read_workbook(src)

    assert wb.sheet_names == ["明细", "汇总"], wb.sheet_names
    assert wb.formula_count == 3
    assert len(wb.get("汇总").df) == 1


def test_first_sheet_is_active(tmp_path):
    src = _make_xlsx(tmp_path / "in.xlsx")
    assert read_workbook(src).active == "明细"


def test_sheet_without_formulas(tmp_path):
    src = _make_xlsx(tmp_path / "in.xlsx", {"纯数据": [["A"], [1], [2]]})
    wb = read_workbook(src)
    assert wb.get("纯数据").formulas == {}
    assert len(wb.get("纯数据").df) == 2


def test_empty_sheet_does_not_crash(tmp_path):
    src = _make_xlsx(tmp_path / "in.xlsx", {"空": [], "有": [["A"], [1]]})
    wb = read_workbook(src)
    assert "空" in wb.sheet_names
    assert wb.get("空").formulas == {}


# ==================== 3. 写：全部工作表 + 公式 ====================

def test_write_preserves_all_sheets_and_formulas(tmp_path):
    """往返：表数与公式数都不变。"""
    src = _make_xlsx(tmp_path / "in.xlsx")
    wb = read_workbook(src)

    out = tmp_path / "out.xlsx"
    write_workbook(wb, out)

    again = read_workbook(out)
    assert again.sheet_names == ["明细", "汇总"], again.sheet_names
    assert again.formula_count == 3, again.formula_count
    assert again.get("明细").formulas == wb.get("明细").formulas


def test_written_formulas_are_real_formulas(tmp_path):
    """写出去的必须是**公式字符串**，不是 pandas 算出来的值。"""
    src = _make_xlsx(tmp_path / "in.xlsx")
    write_workbook(read_workbook(src), tmp_path / "out.xlsx")

    from openpyxl import load_workbook

    ws = load_workbook(tmp_path / "out.xlsx")["明细"]
    assert ws["D2"].value == "=B2*C2", (
        f"D2 应是公式，实际 {ws['D2'].value!r}"
    )


def test_values_are_preserved(tmp_path):
    src = _make_xlsx(tmp_path / "in.xlsx")
    out = tmp_path / "out.xlsx"
    write_workbook(read_workbook(src), out)

    df = pd.read_excel(out, sheet_name="明细", engine="openpyxl")
    assert list(df.columns) == ["产品", "数量", "单价", "金额"]
    assert df.loc[0, "产品"] == "甲"
    assert df.loc[1, "数量"] == 80


def test_no_extra_default_sheet(tmp_path):
    """不应多出一张空白的 "Sheet"。"""
    src = _make_xlsx(tmp_path / "in.xlsx")
    out = tmp_path / "out.xlsx"
    write_workbook(read_workbook(src), out)

    from openpyxl import load_workbook

    assert load_workbook(out).sheetnames == ["明细", "汇总"]


def test_nan_is_written_as_empty_not_string(tmp_path):
    """空单元格必须留空，不能写成长度 4 的 "nan" 字符串。"""
    src = _make_xlsx(tmp_path / "in.xlsx", {
        "表": [["A", "B"], [1, None], [None, 2]],
    })
    out = tmp_path / "out.xlsx"
    write_workbook(read_workbook(src), out)

    import openpyxl

    ws = openpyxl.load_workbook(out)["表"]
    assert ws["B2"].value is None, f"应是 None，实际 {ws['B2'].value!r}"
    assert ws["A3"].value is None, f"应是 None，实际 {ws['A3'].value!r}"


def test_long_sheet_name_is_trimmed(tmp_path):
    """Excel 工作表名 ≤31 字符。"""
    long_name = "很长的表名" * 10
    wb = WorkbookData(sheets=[
        SheetData(name=long_name, df=pd.DataFrame({"A": [1]}))
    ])
    out = tmp_path / "out.xlsx"
    write_workbook(wb, out)

    import openpyxl

    names = openpyxl.load_workbook(out).sheetnames
    assert len(names[0]) <= 31, names


def test_illegal_characters_in_sheet_name(tmp_path):
    wb = WorkbookData(sheets=[
        SheetData(name="a/b:c", df=pd.DataFrame({"A": [1]}))
    ])
    out = tmp_path / "out.xlsx"
    write_workbook(wb, out)

    import openpyxl

    assert "/" not in openpyxl.load_workbook(out).sheetnames[0]


# ==================== 4. sync_active_sheet：编辑后写回 ====================

def test_sync_active_sheet_updates_values_keeps_formulas(tmp_path):
    """用户在蜜算里改了单元格内容，导出时必须保存改动，同时公式不能丢。"""
    src = _make_xlsx(tmp_path / "in.xlsx")
    wb = read_workbook(src)

    edited = wb.get("明细").df.copy()
    edited.loc[0, "产品"] = "甲（已改）"
    sync_active_sheet(wb, "明细", edited)

    out = tmp_path / "out.xlsx"
    write_workbook(wb, out)
    again = read_workbook(out)

    assert again.get("明细").df.loc[0, "产品"] == "甲（已改）"
    assert again.get("明细").formulas == {(0, 3): "=B2*C2", (1, 3): "=B3*C3"}, (
        "改内容不应丢公式"
    )


def test_sync_active_sheet_adds_new_sheet():
    wb = WorkbookData(sheets=[SheetData(name="a", df=pd.DataFrame({"A": [1]}))])
    sync_active_sheet(wb, "b", pd.DataFrame({"B": [2]}))
    assert wb.sheet_names == ["a", "b"]


# ==================== 5. 已知约束：首行被当作表头 ====================

def test_headerless_sheet_loses_its_first_row(tmp_path):
    """
    **已知约束（特性化）**：pandas 固定把工作表第 1 行当表头。

    因此"没有表头行"的表，首行会变成列名、并从数据里消失；
    若首行恰好是公式，该公式也会随之丢失（实测：4 → 3）。

    这不是本模块引入的问题，而是蜜算一直使用的 pandas 约定。这里把它**显式
    钉住**，避免将来"公式怎么会少一个"这种问题又变成难查的谜题。
    若要支持无表头的表，需要导入时让用户选择，或改用 openpyxl 逐格读。

    附带发现（同样钉住）：**公式单元格的值读出来是 NaN** —— 因为 openpyxl
    写出的 xlsx 不含"缓存计算值"，而 pandas 读的正是缓存值。真实 Excel 保存的
    文件里有缓存值，因此这条只影响"由程序生成的文件"。公式字符串本身由本模块
    从 `data_only=False` 那一遍取回，不受影响。
    """
    src = _make_xlsx(tmp_path / "in.xlsx", {
        "无表头": [["合计", "=SUM(A1:A1)"], [1, 2]],
    })

    wb = read_workbook(src)
    sheet = wb.get("无表头")

    # 第 1 行成了列名；公式格没有缓存值 → pandas 起了 "Unnamed: 1"
    assert list(sheet.df.columns) == ["合计", "Unnamed: 1"], (
        f"实际列名 {list(sheet.df.columns)}；若约定改了请更新本测试"
    )
    # 于是它不再是"数据里的公式"
    assert sheet.formulas == {}, (
        "首行的公式被当作表头；若此处变为非空，说明约定改了，请更新本测试"
    )


def test_formula_value_is_nan_without_cached_result(tmp_path):
    """公式格的**值**在无缓存时读为 NaN，但**公式字符串**必须取回。"""
    src = _make_xlsx(tmp_path / "in.xlsx", {
        "表": [["A", "B"], [1, 2], [3, "=A3*2"]],
    })

    sheet = read_workbook(src).get("表")

    assert pd.isna(sheet.df.loc[1, "B"]), (
        "openpyxl 写的文件没有缓存值，pandas 应读到 NaN"
    )
    assert sheet.formulas == {(1, 1): "=A3*2"}, "但公式字符串必须保留下来"


# ==================== 7. 在原结构上更新（保留非数据属性） ====================

def _make_rich_xlsx(path: Path) -> Path:
    """构造一个"什么都有"的工作簿：合并/条件格式/验证/冻结/数字格式/图表/公式。"""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["产品", "数量", "金额"])
    ws.append(["甲", 10, "=B2*2"])
    ws.append(["乙", 80, "=B3*2"])
    ws.append(["丙", 5, "=B4*2"])
    ws["C2"].number_format = "#,##0.00"
    ws.merge_cells("A6:C6")
    ws["A6"] = "合并标题"
    ws.freeze_panes = "A2"
    ws.conditional_formatting.add(
        "B2:B4", CellIsRule(operator="greaterThan", formula=["50"],
                            fill=PatternFill("solid", fgColor="FFC7CE")))
    dv = DataValidation(type="list", formula1='"甲,乙,丙"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("E2:E4")
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4),
                   titles_from_data=True)
    ws.add_chart(chart, "G2")
    # 表格**之外**的内容：写回时绝不能碰
    ws["A20"] = "表外备注：请勿删除"
    ws["B20"] = 12345
    wb.save(str(path))
    return path


def _features(path: Path) -> dict:
    from tools.ooxml_probe import probe

    return probe(path)["features"]


def test_in_place_update_preserves_all_non_data_features(tmp_path):
    """
    核心：只改值，合并/条件格式/验证/冻结/数字格式/图表**全部保留**。

    这正是"新建工作簿"写法做不到的事（那一版实测 5 项丢失）。
    """
    src = _make_rich_xlsx(tmp_path / "rich.xlsx")
    wb = read_workbook(src)
    df = wb.get("明细").df.copy()
    df.loc[0, "产品"] = "甲（已改）"
    sync_active_sheet(wb, "明细", df)

    out = tmp_path / "out.xlsx"
    mode = write_workbook(wb, out)

    assert mode == "updated", f"应走「在原结构上更新」，实际 {mode}"
    before, after = _features(src), _features(out)
    for key in ("merged_ranges", "conditional_formatting",
                "data_validations", "freeze_panes", "charts", "cell_xfs"):
        assert after[key] == before[key], (
            f"{key} 没保留: {before[key]} → {after[key]}"
        )


def test_in_place_update_keeps_formulas_and_applies_edit(tmp_path):
    src = _make_rich_xlsx(tmp_path / "rich.xlsx")
    wb = read_workbook(src)
    df = wb.get("明细").df.copy()
    df.loc[1, "数量"] = 99
    sync_active_sheet(wb, "明细", df)

    out = tmp_path / "out.xlsx"
    write_workbook(wb, out)

    from openpyxl import load_workbook

    ws = load_workbook(out)["明细"]
    assert ws["C2"].value == "=B2*2", "公式被缓存值取代了（多半是 data_only=True 加载）"
    assert ws["B3"].value == 99, "用户改动没有落盘"


def _make_clean_xlsx(path: Path) -> Path:
    """
    **干净**的工作表：只有表头 + 3 行数据，没有合并标题、没有表外备注。

    为什么需要它：pandas 按**已用区域**判断表格范围，`_make_rich_xlsx` 里
    第 6 行的合并标题与第 20 行的备注会让 DataFrame 变成 **19 行**（见
    `test_used_range_absorbs_content_below_table`）。所以"增删行"这类
    按行号断言的测试必须用干净的表。
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["产品", "数量", "金额"])
    ws.append(["甲", 10, "=B2*2"])
    ws.append(["乙", 80, "=B3*2"])
    ws.append(["丙", 5, "=B4*2"])
    wb.save(str(path))
    return path


def test_deleted_row_clears_stale_cells(tmp_path):
    """用户删掉一行后，原来那一行的值不应残留。"""
    src = _make_clean_xlsx(tmp_path / "clean.xlsx")
    wb = read_workbook(src)
    assert wb.get("明细").df.shape[0] == 3, "前置条件：干净的表应是 3 行"

    df = wb.get("明细").df.iloc[:-1]        # 删掉最后一行（丙 / 5）
    sync_active_sheet(wb, "明细", df)

    out = tmp_path / "out.xlsx"
    write_workbook(wb, out)

    from openpyxl import load_workbook

    ws = load_workbook(out)["明细"]
    assert ws["A4"].value is None, f"被删的行还留着: {ws['A4'].value!r}"
    assert ws["B4"].value is None
    # 公式属于结构，保留（只是它所依赖的行没了）
    assert ws["C4"].value == "=B4*2"


def test_openpyxl_cell_value_none_is_a_noop(tmp_path):
    """
    **钉住一个 openpyxl 陷阱**（本模块曾因此静默失效）：

        ws.cell(row=1, column=1, value=None)   # ← 什么都不做！

    `Worksheet.cell()` 仅在 `value is not None` 时才赋值，所以"清空单元格"
    必须写成 `ws.cell(...).value = None`。

    实测后果：用户删掉的一行**仍留在导出文件里**。这条测试直接验证该陷阱
    本身，避免将来有人"顺手改回"更简洁的写法。
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "原值"

    ws.cell(row=1, column=1, value=None)          # 空操作
    assert ws["A1"].value == "原值", "若此处为 None，说明 openpyxl 行为变了"

    ws.cell(row=1, column=1).value = None         # 正确写法
    assert ws["A1"].value is None


def test_cleared_cell_does_not_resurrect_old_value(tmp_path):
    """
    用户把某个单元格清空后，导出文件里不应"复活"旧值。

    直接对应上面那个陷阱：`value=None` 是空操作时，旧值会留在文件里。
    """
    src = _make_clean_xlsx(tmp_path / "clean.xlsx")
    wb = read_workbook(src)
    df = wb.get("明细").df.copy()
    df.loc[0, "产品"] = None            # 用户清空了"甲"

    sync_active_sheet(wb, "明细", df)
    out = tmp_path / "out.xlsx"
    write_workbook(wb, out)

    from openpyxl import load_workbook

    assert load_workbook(out)["明细"]["A2"].value is None, (
        "旧值又回来了 —— 多半是用了 ws.cell(..., value=None)"
    )


def test_added_row_is_written(tmp_path):
    src = _make_clean_xlsx(tmp_path / "clean.xlsx")
    wb = read_workbook(src)
    df = wb.get("明细").df.copy()
    df.loc[len(df)] = ["丁", 7, None]
    sync_active_sheet(wb, "明细", df)

    out = tmp_path / "out.xlsx"
    write_workbook(wb, out)

    again = read_workbook(out).get("明细").df
    assert again.loc[3, "产品"] == "丁", again["产品"].tolist()
    assert again.shape[0] == 4


def test_used_range_absorbs_content_below_table(tmp_path):
    """
    **特性化（重要）**：pandas 按**已用区域**判断表格范围 —— 工作表下方/右方的
    说明文字会被当作数据行吸收进 DataFrame。

    实测：`_make_rich_xlsx` 只写了 3 行数据，却因为第 6 行有合并标题、
    第 20 行有备注，读出来是 **19 行**（中间全是 NaN）。

    这不是本模块引入的，而是"用 pandas 读表"的固有行为。它的用户可见后果是：
    表格下方有备注的工作表，在蜜算里会显示一长串空行。
    这里把它钉住，免得将来当成新 bug 排查。
    """
    src = _make_rich_xlsx(tmp_path / "rich.xlsx")

    sheet = read_workbook(src).get("明细")

    assert sheet.df.shape[0] == 19, (
        f"实测应为 19 行（第 20 行的备注被吸收）；实际 {sheet.df.shape[0]} 行。"
        f"若行为改变，请更新此测试与 docs/fidelity.md"
    )
    assert sheet.df["产品"].iloc[-1] == "表外备注：请勿删除", (
        "表格下方的备注被当成了数据行的首列"
    )


def test_content_below_table_survives_rewrite(tmp_path):
    """
    表格下方的备注在"只改一格"的重写后必须仍在原位置。

    注意它之所以能保留，是因为它被当作**数据行**原样重写回去了
    （见上一条），而不是因为本模块能识别"表外内容"。
    """
    src = _make_rich_xlsx(tmp_path / "rich.xlsx")
    wb = read_workbook(src)
    df = wb.get("明细").df.copy()
    df.loc[0, "产品"] = "甲（改）"
    sync_active_sheet(wb, "明细", df)

    out = tmp_path / "out.xlsx"
    write_workbook(wb, out)

    from openpyxl import load_workbook

    ws = load_workbook(out)["明细"]
    assert ws["A20"].value == "表外备注：请勿删除", (
        f"表格下方的备注丢了: {ws['A20'].value!r}"
    )
    assert ws["A2"].value == "甲（改）"


def test_update_in_place_of_same_path(tmp_path):
    """覆盖保存（path == source）也要能工作。"""
    src = _make_rich_xlsx(tmp_path / "rich.xlsx")
    wb = read_workbook(src)
    df = wb.get("明细").df.copy()
    df.loc[0, "产品"] = "就地改"
    sync_active_sheet(wb, "明细", df)

    mode = write_workbook(wb, src)

    assert mode == "updated"
    assert read_workbook(src).get("明细").df.loc[0, "产品"] == "就地改"
    assert _features(src)["charts"] == 1, "就地保存后图表丢了"


def test_missing_source_falls_back_to_new_workbook(tmp_path):
    """源文件不存在时（或没有源）退回新建，不应报错。"""
    wb = WorkbookData(sheets=[
        SheetData(name="新表", df=pd.DataFrame({"A": [1, 2]}))
    ])
    out = tmp_path / "new.xlsx"

    assert write_workbook(wb, out) == "created"
    assert read_workbook(out).get("新表").df["A"].tolist() == [1, 2]


def test_preserve_structure_can_be_disabled(tmp_path):
    """显式关闭保结构时应走新建路径（便于对照与排查）。"""
    src = _make_rich_xlsx(tmp_path / "rich.xlsx")
    wb = read_workbook(src)
    out = tmp_path / "flat.xlsx"

    assert write_workbook(wb, out, preserve_structure=False) == "created"
    assert _features(out)["charts"] == 0


# ==================== 8. 与保真度探针联动 ====================

def test_workbook_roundtrip_keeps_probe_features(tmp_path):
    """
    用**探针**验证：sheets 与 formulas 两项特征在往返后不再归零。

    这是 docs/fidelity.md 里 `sheets 2→1`、`formulas 4→0` 两条的直接对照。
    """
    from tools.fidelity_roundtrip import build_xlsx_fixture
    from tools.ooxml_probe import compare, probe

    src = tmp_path / "fixture.xlsx"
    build_xlsx_fixture(src)
    before = probe(src)

    out = tmp_path / "roundtrip.xlsx"
    write_workbook(read_workbook(src), out)
    after = probe(out)

    rows = {r["feature"]: r for r in compare(before, after)}
    assert rows["sheets"]["status"] == "same", rows["sheets"]
    assert rows["sheets"]["after"] == before["features"]["sheets"]
    assert rows["formulas"]["after"] == before["features"]["formulas"], (
        f"公式仍丢失: {rows['formulas']}"
    )
