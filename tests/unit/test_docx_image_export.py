# -*- coding: utf-8 -*-
"""
DOCX 导出：产物级验收测试（P0）

背景：
    原 `_write_element` 没有 <img> 分支，且 `_write_inline` 在
    "文本为空"时直接 continue —— 于是 Markdown 的 ![alt](path)
    渲染出来的 `<p><img/></p>` 被整段跳过，导出的 DOCX 里一张图都没有，
    而导出函数仍返回 True。

本文件不测"函数有没有被调用"，而是**把导出的文件重新打开，数图片**。
这是路线图 WP-07「产物级验收」的最小样本。
"""
import pytest

pytest.importorskip("docx", reason="需要 python-docx")
pytest.importorskip("markdown", reason="需要 markdown")
pytest.importorskip("bs4", reason="需要 beautifulsoup4")
pytest.importorskip("PIL", reason="需要 Pillow 生成测试图片")


@pytest.fixture
def workdir(tmp_path):
    from PIL import Image
    Image.new("RGB", (200, 100), (255, 0, 0)).save(tmp_path / "pic.png")
    return tmp_path


def _export(md_text, out_path):
    from beeoffice.plugins.bee_export import export_to_docx
    return export_to_docx(md_text, str(out_path))


def test_image_survives_docx_export(workdir):
    """Markdown 里的图片必须出现在 DOCX 里（回归 P0-3）。"""
    from docx import Document

    out = workdir / "out.docx"
    assert _export("# 标题\n\n正文\n\n![图](pic.png)\n\n结尾\n", out)
    assert out.exists() and out.stat().st_size > 0

    doc = Document(str(out))
    assert len(doc.inline_shapes) == 1, "导出的 DOCX 里应恰好有 1 张图片"


def test_image_paragraph_is_centered(workdir):
    """纯图片段落应居中（与蜜写富文本导出的 v5.0.4 约定一致）。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    out = workdir / "center.docx"
    assert _export("![图](pic.png)\n", out)

    doc = Document(str(out))
    img_paras = [
        p for p in doc.paragraphs
        if p.alignment == WD_ALIGN_PARAGRAPH.CENTER
    ]
    assert img_paras, "应存在居中的图片段落"


def test_missing_image_is_not_silently_dropped(workdir):
    """图片缺失时要留占位文字，而不是静默消失。"""
    from docx import Document

    out = workdir / "missing.docx"
    assert _export("# T\n\n![x](not_exists.png)\n", out)

    doc = Document(str(out))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "[图片缺失" in full_text


def test_text_only_export_still_works(workdir):
    """对照组：纯文本导出不应被图片逻辑影响。"""
    from docx import Document

    out = workdir / "text.docx"
    assert _export("# 标题\n\n第一段\n\n第二段\n", out)

    doc = Document(str(out))
    assert len(doc.inline_shapes) == 0
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "标题" in text and "第一段" in text and "第二段" in text
