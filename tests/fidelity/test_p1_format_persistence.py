# tests/fidelity/test_p1_format_persistence.py
"""
第四批：蜜算**视图侧格式落盘**（加粗/斜体/下划线/对齐、数字格式、合并、批注、条件格式）。

## 修前（审计报告 §2.9）
这些格式全在视图内存里（`_cell_formats`/`_merged_cells`/`_cell_comments`/
`_conditional_formats`），而写盘只有"值 + 公式"两条路 —— 它们在导出链上**零引用**。
于是用户自己加的加粗、对齐、数字格式、合并、批注、条件格式**保存后全部消失**。

## 修法
`beesheet_style_export.apply_view_styles()`：**两条写路径共用**的"写盘后后处理"
（打开刚写好的 xlsx，把样式贴上去再存）。这样逻辑只有一份，
而且可以用 openpyxl **读回来做往返断言**（本文件就是这么验的）。

## 两个必须守住的底线（本文件单独有测试）
1. **写格式不能丢数据**：openpyxl 合并会清掉非左上角的值 →
   若合并区内还有数据，**跳过合并**并计数（数据优先于外观）；
2. **视图的合并语义要与文档一致**：`_merge_cells` 的文档写着"其余置空"，
   但实现从未真的清 —— 结果是合并既存不下来（守卫跳过）又留着看不见的值。
   现在真的清空，并**明确告知清掉了几个值**（不静默改用户数据）。
"""
import sys
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


# ==================== 1. 纯函数：坐标与数字格式 ====================

def test_number_format_roundtrip_is_reversible():
    """
    **往返属性测试**（正向 → 反向必须回到原值）。

    这个测试是被一个真 bug 逼出来的：`_parse_number_format` 只剥了逗号、
    没剥 `#` 占位符，于是 `¥#,##0.00` 解析失败 → **读回时货币格式静默丢失**。
    我原先只测了正向（`number_format_code`），所以没抓到 ——
    写/读这种成对的逻辑，必须测"互逆"。
    """
    from beeoffice.plugins.beesheet_style_export import (
        _parse_number_format, number_format_code,
    )

    cases = [
        ("数值", 2, False), ("数值", 2, True), ("数值", 0, False),
        ("货币", 2, True), ("货币", 2, False),
        ("百分比", 1, False), ("日期", 2, False), ("文本", 2, False),
    ]
    for kind, decimals, thousands in cases:
        code = number_format_code(kind, decimals, thousands)
        parsed = _parse_number_format(code)
        assert parsed is not None, f"{kind} 的格式 {code!r} 解析不回来"
        assert parsed[0] == kind, f"{code!r} 解析成 {parsed[0]}，应为 {kind}"
        assert parsed[1] == decimals, f"{code!r} 的小数位数解析错：{parsed[1]}"


def test_view_to_cell_skips_header_row():
    """视图 (0,0) 是第一条数据 → 工作表 (第 2 行, A 列)。"""
    from beeoffice.plugins.beesheet_style_export import view_to_cell

    assert view_to_cell(0, 0) == (2, 1)
    assert view_to_cell(3, 2) == (5, 3)


@pytest.mark.parametrize("kind,decimals,thousands,expected", [
    ("常规", 2, False, "General"),
    ("文本", 2, False, "@"),
    ("数值", 2, False, "0.00"),
    ("数值", 2, True, "#,##0.00"),
    ("数值", 0, False, "0"),
    ("货币", 2, True, "¥#,##0.00"),
    ("货币", 2, False, "¥0.00"),
    ("百分比", 1, False, "0.0%"),
    ("日期", 2, False, "yyyy-mm-dd"),
])
def test_number_format_codes(kind, decimals, thousands, expected):
    from beeoffice.plugins.beesheet_style_export import number_format_code

    assert number_format_code(kind, decimals, thousands) == expected


# ==================== 2. 端到端：保存后真的写得进去 ====================

@pytest.fixture
def sheet_view(qtbot):
    from PySide6.QtWidgets import QMessageBox

    from beeoffice.plugins.beesheet import BeeSheetView

    view = BeeSheetView()
    qtbot.addWidget(view)
    view.model.set_full_dataframe(
        pd.DataFrame({"名称": ["甲", "乙", "丙"], "金额": [1234.5, 200.0, 3.25]})
    )
    view._set_status = lambda *a, **k: None
    # 保存时不该弹对话框
    QMessageBox.information = staticmethod(lambda *a, **k: None)
    QMessageBox.warning = staticmethod(lambda *a, **k: None)
    return view


def _load(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    return workbook, workbook[workbook.sheetnames[0]]


def test_cell_formats_are_persisted(sheet_view, tmp_path):
    target = tmp_path / "格式.xlsx"
    sheet_view._cell_formats = {
        (0, 0): {"bold": True, "align": "center"},
        (1, 0): {"italic": True, "underline": True},
        (1, 1): {"number_format": "货币", "decimals": 2, "thousands": True},
        (2, 1): {"number_format": "百分比", "decimals": 1},
    }
    sheet_view.set_current_file_path(target)

    assert sheet_view.save() is True

    _, ws = _load(target)
    assert ws["A2"].font.bold is True
    assert ws["A2"].alignment.horizontal == "center"
    assert ws["A3"].font.italic is True
    assert ws["A3"].font.underline == "single"
    assert ws["B3"].number_format == "¥#,##0.00"
    assert ws["B4"].number_format == "0.0%"
    # 数据仍然在
    assert ws["B2"].value == 1234.5


def test_comments_are_persisted(sheet_view, tmp_path):
    target = tmp_path / "批注.xlsx"
    sheet_view._cell_comments = {(2, 0): "这是批注内容"}
    sheet_view.set_current_file_path(target)

    sheet_view.save()

    _, ws = _load(target)
    assert ws["A4"].comment is not None
    assert ws["A4"].comment.text == "这是批注内容"


def test_conditional_formats_are_persisted(sheet_view, tmp_path):
    target = tmp_path / "条件格式.xlsx"
    sheet_view._conditional_formats = {
        1: {"type": "data_bars", "min": 200.0, "max": 1234.5}
    }
    sheet_view.set_current_file_path(target)

    sheet_view.save()

    _, ws = _load(target)
    ranges = [str(r) for r in ws.conditional_formatting]
    assert ranges, "条件格式没有写进文件"
    assert "B2" in " ".join(ranges)


def test_real_merge_is_persisted_and_hidden_values_reported(sheet_view,
                                                            tmp_path):
    """用**视图自己的合并功能**（不是手改状态）→ 合并要能存下来。"""
    from PySide6.QtCore import QItemSelection, QItemSelectionModel

    target = tmp_path / "合并.xlsx"
    statuses = []
    sheet_view._set_status = lambda msg, t=0: statuses.append(msg)

    selection = QItemSelection(sheet_view.model.index(0, 0),
                               sheet_view.model.index(0, 1))
    sheet_view.table_view.selectionModel().select(
        selection, QItemSelectionModel.SelectionFlag.Select
    )
    sheet_view._merge_cells()

    # 被遮住的值要真的清掉，并且**明确告知**（不静默改数据）
    assert sheet_view.model.df.iloc[0].tolist()[:2] == ["甲", ""]
    assert any("清除" in s for s in statuses), f"没有告知清掉了值：{statuses}"

    sheet_view.set_current_file_path(target)
    sheet_view.save()

    _, ws = _load(target)
    assert [str(r) for r in ws.merged_cells.ranges] == ["A2:B2"]


# ==================== 3. 底线：写格式不能丢数据 ====================

def test_merge_with_hidden_data_is_skipped_not_destroyed(sheet_view, tmp_path):
    """
    **关键回归**：合并区内还有数据（导入的合并区/状态不同步）时，
    必须**跳过合并**保住数据 —— openpyxl 的 merge_cells 会清掉非左上角的值，
    静默吃掉用户数据是绝不能接受的。
    """
    target = tmp_path / "危险合并.xlsx"
    # 直接伪造"合并状态存在、但数据还在"（导入的文件常见这种组合）
    sheet_view._merged_cells = {(0, 0): (1, 2)}
    sheet_view.set_current_file_path(target)

    assert sheet_view.save() is True

    _, ws = _load(target)
    assert [str(r) for r in ws.merged_cells.ranges] == [], "合并被执行了"
    assert ws["B2"].value == 1234.5, "被遮住的数据被吃掉了！"


def test_apply_view_styles_is_fast_noop_when_nothing_to_write(tmp_path):
    from beeoffice.plugins.beesheet_style_export import apply_view_styles

    target = tmp_path / "空.xlsx"
    target.write_bytes(b"not-a-workbook")     # 故意是坏文件

    result = apply_view_styles(target)        # 全 None → 直接返回，不去动文件

    # 契约：没有任何要写的东西时，所有计数为 0（键会随功能增加）
    assert all(value == 0 for value in result.values()), result
    assert {"formats", "merges", "merges_skipped", "comments",
            "conditional", "table_style", "validations"} <= set(result)
    assert target.read_bytes() == b"not-a-workbook", "空格式却改了文件"


def test_apply_view_styles_survives_corrupt_file(tmp_path):
    """格式写入失败**不能**抛异常（数据已经安全写入，格式是附加项）。"""
    from beeoffice.plugins.beesheet_style_export import apply_view_styles

    target = tmp_path / "坏文件.xlsx"
    target.write_bytes(b"definitely-not-a-workbook")

    result = apply_view_styles(target, formats={(0, 0): {"bold": True}})

    assert result["formats"] == 0
    assert target.read_bytes() == b"definitely-not-a-workbook"


# ==================== 5. 读回：存了要能看到 ====================

def test_formats_merges_comments_are_read_back_on_import(qtbot, tmp_path):
    """
    **闭环的另一半**：保存后再打开，格式/合并/批注都要回来。

    只写不读 = "存了却看不到" —— 用户会以为保存坏了。
    """
    from PySide6.QtCore import QItemSelection, QItemSelectionModel

    from beeoffice.plugins.beesheet import BeeSheetView

    target = tmp_path / "往返.xlsx"

    writer = BeeSheetView()
    qtbot.addWidget(writer)
    writer._set_status = lambda *a, **k: None
    writer.model.set_full_dataframe(
        pd.DataFrame({"产品": ["A", "B", "C"], "金额": [1234.5, 200.0, 3.25]})
    )
    writer._cell_formats = {
        (0, 0): {"bold": True, "align": "center"},
        (1, 1): {"number_format": "货币", "decimals": 2, "thousands": True},
        (2, 1): {"number_format": "百分比", "decimals": 1},
    }
    writer._cell_comments = {(2, 0): "季度初"}
    selection = QItemSelection(writer.model.index(0, 0),
                               writer.model.index(0, 1))
    writer.table_view.selectionModel().select(
        selection, QItemSelectionModel.SelectionFlag.Select
    )
    writer._merge_cells()
    writer.set_current_file_path(target)
    assert writer.save() is True
    writer.close()

    reader = BeeSheetView()
    qtbot.addWidget(reader)
    reader._set_status = lambda *a, **k: None
    try:
        reader._import_file_from_path(str(target))

        assert reader._cell_formats.get((0, 0), {}).get("bold") is True
        assert reader._cell_formats.get((0, 0), {}).get("align") == "center"
        # ← 这一条就是那个"没剥 # 占位符"的 bug 的回归点
        assert reader._cell_formats.get((1, 1), {}).get("number_format") == "货币"
        assert reader._cell_formats.get((1, 1), {}).get("decimals") == 2
        assert reader._cell_formats.get((2, 1), {}).get("number_format") == "百分比"
        assert reader._cell_comments.get((2, 0)) == "季度初"
        assert reader._merged_cells.get((0, 0)) == (1, 2)
        # 合并不仅要存进 "状态"，界面也要真的合并起来
        assert reader.table_view.rowSpan(0, 0) == 1
        assert reader.table_view.columnSpan(0, 0) == 2
    finally:
        reader.close()


def test_read_back_does_not_fabricate_conditional_formats(tmp_path):
    """
    条件格式**不读回**（如实说明）：也不许凭空造出视图状态。
    """
    from openpyxl import Workbook
    from openpyxl.formatting.rule import DataBarRule

    from beeoffice.plugins.beesheet_style_export import read_view_styles

    path = tmp_path / "有条件格式.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["金额"])
    sheet.append([100])
    sheet.conditional_formatting.add(
        "A2:A2",
        DataBarRule(start_type="num", start_value=0, end_type="num",
                    end_value=100, color="638EC6"),
    )
    workbook.save(path)

    styles = read_view_styles(path)

    assert "conditional" not in styles, "读回结果里不该有编造的条件格式"
    assert styles["formats"] == {}


def test_workbook_write_path_also_gets_formats(qtbot, tmp_path):
    """
    "导入 xlsx → 改 → 保存"走的是 `write_workbook` 那条路，
    它同样必须带上视图格式（否则两条路行为不一致）。
    """
    from beeoffice.plugins.beesheet import BeeSheetView
    from beeoffice.plugins.beesheet_workbook import read_workbook

    source = tmp_path / "源工作簿.xlsx"
    pd.DataFrame({"金额": [100.0, 200.0]}).to_excel(source, index=False)

    view = BeeSheetView()
    qtbot.addWidget(view)
    try:
        view._set_status = lambda *a, **k: None
        view._import_file_from_path(str(source))
        assert view._workbook is not None, "没走工作簿路径，测试前提不成立"

        view._cell_formats = {(0, 0): {"number_format": "货币", "decimals": 2}}
        view.set_current_file_path(source)
        assert view.save() is True

        _, ws = _load(source)
        assert ws["A2"].number_format.startswith("¥"), (
            f"工作簿路径没带数字格式：{ws['A2'].number_format}"
        )
        assert read_workbook(source).sheets, "文件结构被写坏了"
    finally:
        view.close()
