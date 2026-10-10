# tests/regression/v5/test_v505_docx_overwrite_guard.py
"""
优化规划 §1.3 回归：打开 .docx 后保存不得静默覆盖原文件。

原缺陷
------
app._open_docx_file() 把 widget._current_file_path 设为用户的原 .docx，
于是 Ctrl+S → _save_as_docx() → _write_docx(原路径) → _rich_to_docx()
会在这份原文件上"重新生成"一份文档。该重建保真度很低（超链接、列表编号、
页眉页脚、脚注、域、东亚字体、单元格内除第一段外的内容全部丢失），而且
既没有备份也没有确认——属于不可逆的数据破坏。

修复后
------
_save_as_docx() 发现保存目标正是本次打开的 _source_docx 时，必须先让用户
选择：另存为（推荐）/ 覆盖并自动备份 / 取消；且备份失败要中止覆盖。
"""
import pytest
from docx import Document
from PySide6.QtWidgets import QMessageBox


@pytest.fixture(autouse=True)
def _no_modal_dialogs(monkeypatch):
    """
    headless 环境不要弹模态框——QMessageBox.exec() 会把测试挂死。
    这里把静态提示方法替换为记录器，测试需要时可断言其内容。
    """
    shown = []

    def _record(*args, **kwargs):
        shown.append(args)
        return QMessageBox.Ok

    monkeypatch.setattr(QMessageBox, "warning", staticmethod(_record))
    monkeypatch.setattr(QMessageBox, "information", staticmethod(_record))
    monkeypatch.setattr(QMessageBox, "critical", staticmethod(_record))
    return shown


@pytest.fixture
def opened_docx(tmp_path):
    """造一个真实的 .docx，返回 (路径, 原始字节)。"""
    path = tmp_path / "original.docx"
    doc = Document()
    doc.add_heading("原标题", level=1)
    doc.add_paragraph("原始正文：覆盖后这段内容应当仍然存在于 .bak 里")
    doc.save(str(path))
    return path, path.read_bytes()


def _mark_as_opened(editor, path):
    """模拟 app._open_docx_file() 留下的状态。"""
    editor._doc_kind = "word"
    editor._current_file_path = str(path)
    editor._source_docx = str(path)


def test_cancel_keeps_original_untouched(beewrite_editor, opened_docx, monkeypatch):
    """用户选择取消：不得写任何文件。"""
    path, before = opened_docx
    _mark_as_opened(beewrite_editor, path)
    monkeypatch.setattr(
        type(beewrite_editor), "_ask_docx_overwrite", lambda self, p: "cancel"
    )

    assert beewrite_editor._save_as_docx() is False
    assert path.read_bytes() == before, "取消后原文件必须逐字节不变"
    assert not path.with_name(path.name + ".bak").exists()


def test_save_as_keeps_original_untouched(
    beewrite_editor, opened_docx, monkeypatch, tmp_path
):
    """用户选择另存为：原文件不变，新文件生成。"""
    path, before = opened_docx
    new_path = tmp_path / "copy.docx"
    _mark_as_opened(beewrite_editor, path)
    monkeypatch.setattr(
        type(beewrite_editor), "_ask_docx_overwrite", lambda self, p: "save_as"
    )
    monkeypatch.setattr(
        "beeoffice.plugins.beewrite_export_mixin.QFileDialog.getSaveFileName",
        lambda *a, **k: (str(new_path), "Word (*.docx)"),
    )

    assert beewrite_editor._save_as_docx() is True
    assert new_path.exists() and new_path.stat().st_size > 0
    assert path.read_bytes() == before, "另存为不得改动原文件"


def test_overwrite_backs_up_first(beewrite_editor, opened_docx, monkeypatch):
    """用户选择覆盖：必须先落一份内容等于原文件的 .bak。"""
    path, before = opened_docx
    _mark_as_opened(beewrite_editor, path)
    monkeypatch.setattr(
        type(beewrite_editor), "_ask_docx_overwrite", lambda self, p: "overwrite"
    )

    assert beewrite_editor._save_as_docx() is True
    backup = path.with_name(path.name + ".bak")
    assert backup.exists(), "覆盖前必须生成 .bak"
    assert backup.read_bytes() == before, ".bak 必须保留原文件内容"


def test_failed_backup_aborts_overwrite(beewrite_editor, opened_docx, monkeypatch):
    """备份失败时必须中止覆盖——宁可这次不保存，也不做不可逆的事。"""
    path, before = opened_docx
    _mark_as_opened(beewrite_editor, path)
    monkeypatch.setattr(
        type(beewrite_editor), "_ask_docx_overwrite", lambda self, p: "overwrite"
    )
    monkeypatch.setattr(
        type(beewrite_editor), "_backup_before_overwrite", lambda self, p: False
    )

    assert beewrite_editor._save_as_docx() is False
    assert path.read_bytes() == before


def test_overwrite_prompts_only_once_per_session(
    beewrite_editor, opened_docx, monkeypatch
):
    """确认过覆盖后，同一文档后续保存不应反复打扰用户。"""
    path, before = opened_docx
    _mark_as_opened(beewrite_editor, path)
    calls = []
    monkeypatch.setattr(
        type(beewrite_editor),
        "_ask_docx_overwrite",
        lambda self, p: calls.append(p) or "overwrite",
    )

    assert beewrite_editor._save_as_docx() is True
    assert beewrite_editor._save_as_docx() is True
    assert len(calls) == 1, f"应只询问一次，实际 {len(calls)} 次"


def test_new_document_does_not_prompt(beewrite_editor, monkeypatch, tmp_path):
    """新建文档（无 _source_docx）不应触发覆盖确认。"""
    called = []
    monkeypatch.setattr(
        type(beewrite_editor),
        "_ask_docx_overwrite",
        lambda self, p: called.append(p) or "cancel",
    )
    new_path = tmp_path / "brand_new.docx"
    monkeypatch.setattr(
        "beeoffice.plugins.beewrite_export_mixin.QFileDialog.getSaveFileName",
        lambda *a, **k: (str(new_path), "Word (*.docx)"),
    )
    beewrite_editor._doc_kind = "word"
    beewrite_editor._current_file_path = None
    beewrite_editor._source_docx = None

    assert beewrite_editor._save_as_docx() is True
    assert called == [], "新建文档不应触发覆盖确认"
    assert new_path.exists()


def test_saving_after_save_as_targets_new_file(
    beewrite_editor, opened_docx, monkeypatch, tmp_path
):
    """另存为之后，再保存应写新文件，且不再询问。"""
    path, before = opened_docx
    new_path = tmp_path / "copy2.docx"
    _mark_as_opened(beewrite_editor, path)
    monkeypatch.setattr(
        type(beewrite_editor), "_ask_docx_overwrite", lambda self, p: "save_as"
    )
    monkeypatch.setattr(
        "beeoffice.plugins.beewrite_export_mixin.QFileDialog.getSaveFileName",
        lambda *a, **k: (str(new_path), "Word (*.docx)"),
    )

    assert beewrite_editor._save_as_docx() is True          # 另存为
    stamp = new_path.stat().st_mtime_ns

    # 第二次保存：目标已是新文件，不应再询问，也不应改动原文件
    assert beewrite_editor._save_as_docx() is True
    assert new_path.exists()
    assert new_path.stat().st_mtime_ns >= stamp
    assert path.read_bytes() == before
