# tests/fidelity/test_pptx_import.py
"""
P0-1：**打开 `.pptx`**（导入成蜜演文稿）。

## 为什么这是 P0（门槛）
审计与战略都标它是门槛：连别人发来的 PPT 都打不开，"文档流水线"的故事就不成立。
`BeeOffice-与WPS差距分析` 里写的是「`.pptx` 打不开」—— 本批补上**导入**能力。

## 诚实的边界（本文件同样锁住）
**保留**：标题、正文段落**层级**、演讲者备注、页序
**不保留**：版式/主题/字体、图片、表格、图表、动画、超链接目标
→ 有损失时**必须给损失清单**（`loss_summary`），不能假装搬得过去。
→ 保存时**绝不覆盖原始 `.pptx`**（界面侧 `_current_file_path = None`，
   走"另存为"）。
"""
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


@pytest.fixture
def demo_pptx(tmp_path):
    """造一份真实 pptx：标题 + 两级正文 + 备注 + 图片 + 表格。"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    slide1.shapes.title.text = "季度汇报"
    body = slide1.placeholders[1].text_frame
    body.text = "收入增长"
    sub1 = body.add_paragraph()
    sub1.text = "华东区 +20%"
    sub1.level = 1
    sub2 = body.add_paragraph()
    sub2.text = "华南区 +15%"
    sub2.level = 1
    slide1.notes_slide.notes_text_frame.text = "这里讲增长结构"

    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "下半年计划"
    image = tmp_path / "pic.png"
    from PIL import Image

    Image.new("RGB", (40, 30), (200, 80, 40)).save(image)
    slide2.shapes.add_picture(str(image), Inches(1), Inches(2))
    table = slide2.shapes.add_table(
        2, 2, Inches(1), Inches(4), Inches(4), Inches(1)
    ).table
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "预算"
    table.cell(1, 0).text = "招聘"
    table.cell(1, 1).text = "10万"

    path = tmp_path / "demo.pptx"
    prs.save(str(path))
    return path


# ==================== 1. 抽取：标题/层级/备注 ====================

def test_read_pptx_keeps_text_hierarchy_and_notes(demo_pptx):
    from beeoffice.plugins.beeshow_pptx_import import read_pptx

    imported = read_pptx(demo_pptx)

    assert imported.slide_count == 2
    first = imported.slides[0]
    assert first.title == "季度汇报"
    # 标题**不能**被重复算进正文（踩过的坑：`slide.shapes.title` 每次返回新代理，
    # 用 `is` 比较永远 False → 标题被抽两遍）
    assert all(text != "季度汇报" for _level, text in first.bullets), first.bullets
    # 层级要保留：0 级 + 1 级
    levels = {level for level, _text in first.bullets}
    assert 0 in levels and 1 in levels, first.bullets
    assert first.notes == "这里讲增长结构"


def test_read_pptx_reports_losses(demo_pptx):
    """图片/表格必须**记进损失清单**，不能静默丢弃。"""
    from beeoffice.plugins.beeshow_pptx_import import loss_summary, read_pptx

    imported = read_pptx(demo_pptx)
    assert imported.losses.get("图片") == 1, imported.losses
    assert imported.losses.get("表格") == 1, imported.losses
    assert imported.has_losses() is True

    summary = loss_summary(imported)
    assert "图片" in summary and "表格" in summary
    assert "不会" in summary or "不保留" in summary


def test_loss_summary_is_honest_when_nothing_is_lost(tmp_path):
    """只有文字时，损失说明要明确说"全部保留"（不制造焦虑）。"""
    from pptx import Presentation

    from beeoffice.plugins.beeshow_pptx_import import loss_summary, read_pptx

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "只有文字"
    slide.placeholders[1].text_frame.text = "一条要点"
    path = tmp_path / "text-only.pptx"
    prs.save(str(path))

    imported = read_pptx(path)
    assert imported.has_losses() is False
    assert "全部保留" in loss_summary(imported)


# ==================== 2. 渲染：蜜演文稿格式 ====================

def test_render_deck_uses_beshow_conventions(demo_pptx):
    from beeoffice.plugins.beeshow_pptx_import import import_pptx

    _imported, markdown = import_pptx(demo_pptx)

    assert "# 季度汇报" in markdown
    assert "- 收入增长" in markdown
    assert "  - 华东区 +20%" in markdown, markdown      # 层级缩进
    assert "> 备注：这里讲增长结构" in markdown
    assert "\n---\n" in markdown                       # 页分隔符


def test_import_pptx_to_file_writes_a_real_beshow(demo_pptx, tmp_path):
    """写出的文件必须是**合法的蜜演文稿**（带标记，能被 load_deck 读回）。"""
    from beeoffice.plugins.beeshow_doc import is_deck_file, load_deck
    from beeoffice.plugins.beeshow_pptx_import import import_pptx_to_file

    out = tmp_path / "demo.beshow"
    imported = import_pptx_to_file(demo_pptx, out)

    assert imported.slide_count == 2
    assert is_deck_file(out), "缺少蜜蜂演文稿标记"
    text = load_deck(out)
    assert text is not None and "季度汇报" in text


# ==================== 3. 错误处理：给可操作提示 ====================

def test_read_pptx_rejects_missing_file(tmp_path):
    from beeoffice.plugins.beeshow_pptx_import import read_pptx

    with pytest.raises(ValueError) as exc:
        read_pptx(tmp_path / "nope.pptx")
    assert "不存在" in str(exc.value)


def test_read_pptx_rejects_non_pptx(tmp_path):
    from beeoffice.plugins.beeshow_pptx_import import read_pptx

    path = tmp_path / "x.docx"
    path.write_text("not a pptx", encoding="utf-8")
    with pytest.raises(ValueError) as exc:
        read_pptx(path)
    assert ".pptx" in str(exc.value)


# ==================== 4. CLI 与界面接线 ====================

def test_cli_converts_pptx_to_beshow(demo_pptx, tmp_path, capsys):
    from beeoffice import beecmd

    out = tmp_path / "out.beshow"
    code = beecmd.main(["convert", str(demo_pptx), "--to", "beshow",
                        "-o", str(out)])
    assert code == 0
    assert out.is_file()
    captured = capsys.readouterr().out
    assert "已导入 2 页" in captured or "2 页" in captured


def test_show_extensions_include_pptx():
    """应用的"打开"分流必须认得 `.pptx`（否则用户拖进来仍会得到"不支持"）。"""
    import beeoffice.app as app_module

    assert ".pptx" in app_module.SHOW_EXTENSIONS


def test_beeshow_view_exposes_open_pptx():
    """蜜演视图要有 `open_pptx`（类级检查：本机 WebEngine 起不来，不能实例化）。"""
    from beeoffice.plugins.beeshow import BeeShowView

    assert callable(getattr(BeeShowView, "open_pptx", None))


def test_app_dispatches_pptx_to_open_pptx():
    """源码级确认：`.pptx` 分派到 `open_pptx`，而不是 `open_deck`。"""
    source = (ROOT / "beeoffice" / "app.py").read_text(encoding="utf-8")
    assert 'suffix.lower() == ".pptx"' in source
    assert "open_pptx(" in source


# ==================== 5. 纯 Python（不 import Qt）====================

def test_pptx_import_module_does_not_require_qt(demo_pptx, tmp_path):
    """导入器必须能在没有 Qt 的环境用（CLI/服务器）；子进程写文件取证（沙箱禁管道）。"""
    result_file = tmp_path / "probe.txt"
    script = (
        "import sys\n"
        "from beeoffice.plugins.beeshow_pptx_import import read_pptx\n"
        f"imp = read_pptx({str(demo_pptx)!r})\n"
        "loaded = [m for m in sys.modules if m.startswith(('PySide6', 'shiboken'))]\n"
        f"open({str(result_file)!r}, 'w', encoding='utf-8').write("
        "f'pages={imp.slide_count};qt={\",\".join(loaded)}')\n"
    )
    completed = subprocess.run(
        [sys.executable, "-c", script], cwd=str(ROOT),
        env={**__import__("os").environ, "PYTHONPATH": str(ROOT)},
        timeout=180,
    )
    assert completed.returncode == 0
    payload = result_file.read_text(encoding="utf-8")
    assert payload.startswith("pages=2;"), payload
    assert payload.endswith("qt="), f"导入器不该加载 Qt：{payload}"
