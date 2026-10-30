"""Writer 页码域 + 分节页脚。"""
from docx import Document

from beeoffice.plugins.bee_export import apply_header_footer, apply_page_numbers
from beeoffice.plugins.beewrite_page_extras import (
    HeaderFooterDialog,
    PageNumberDialog,
)


def test_page_number_dialog_all_sections(qapp):
    dlg = PageNumberDialog({"enabled": True, "all_sections": True})
    cfg = dlg.get_config()
    assert cfg["enabled"] is True
    assert cfg["all_sections"] is True


def test_page_numbers_all_sections(tmp_path):
    doc = Document()
    doc.add_paragraph("第一页")
    doc.add_section()
    doc.add_paragraph("第二页")
    config = {
        "enabled": True, "position": "底部居中", "format": "第 1 页 / 共 N 页",
        "start": 1, "num_format": "lowerRoman", "all_sections": True,
    }
    assert apply_page_numbers(doc, config) is True
    out = tmp_path / "pages.docx"
    doc.save(str(out))
    reopened = Document(str(out))
    assert len(reopened.sections) >= 2
    for section in reopened.sections:
        xml = section.footer._element.xml
        assert "PAGE" in xml
    pg = reopened.sections[0]._sectPr.xml
    assert "w:pgNumType" in pg and "lowerRoman" in pg


def test_header_footer_dialog_all_sections(qapp):
    dlg = HeaderFooterDialog()
    dlg._all_sections_check.setChecked(True)
    assert dlg.options()["all_sections"] is True


def test_section_parts_footer_text(tmp_path):
    doc = Document()
    doc.add_paragraph("A")
    doc.add_section()
    doc.add_paragraph("B")
    parts = [
        {"footer_text": "第一节页脚", "footer_linked": False,
         "header_linked": True, "header_text": "", "header_xml": "",
         "footer_xml": ""},
        {"footer_text": "第二节页脚", "footer_linked": False,
         "header_linked": True, "header_text": "", "header_xml": "",
         "footer_xml": ""},
    ]
    assert apply_header_footer(doc, section_parts=parts) is True
    out = tmp_path / "sections.docx"
    doc.save(str(out))
    reopened = Document(str(out))
    footers = [
        s.footer.paragraphs[0].text for s in reopened.sections
        if s.footer.paragraphs
    ]
    assert "第一节页脚" in footers
    assert "第二节页脚" in footers
