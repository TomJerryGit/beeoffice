# -*- coding: utf-8 -*-
"""
蜜写页眉 / 页脚：写进 Word 的节，而不是正文里的一行注释（P1）

背景：
    `_insert_header` 原本把 `<!-- header: 机密 -->` 直接插进正文，
    在编辑器里就是一行可见的注释文字；而且全仓没有任何解析它的代码，
    导出 Word 时也不会生效——既难看又没用。

修复后：
    · 文案只记在 `_header_text` / `_footer_text`；
    · 导出 DOCX 时写入 `section.header` / `section.footer`。
"""
import pytest

pytest.importorskip("PySide6.QtWidgets")
pytest.importorskip("docx", reason="需要 python-docx")


@pytest.fixture
def editor(qtbot):
    from beeoffice.plugins.beewrite import BeeWriteEditor

    ed = BeeWriteEditor()
    qtbot.addWidget(ed)
    ed.set_content("# 标题\n\n正文\n")
    try:
        ed._switch_mode("rich_text")
    except Exception:
        pass
    yield ed
    try:
        ed.close()
    except Exception:
        pass


def test_insert_header_does_not_touch_body(editor, monkeypatch):
    """设置页眉不得改动正文（原实现会插入一行 HTML 注释）。"""
    from PySide6.QtWidgets import QInputDialog

    monkeypatch.setattr(
        QInputDialog, "getText",
        staticmethod(lambda *a, **k: ("机密文件", True)),
    )
    before = editor.rich_editor.toPlainText()

    editor._insert_header()

    assert editor._header_text == "机密文件"
    assert editor.rich_editor.toPlainText() == before, "正文被插入了额外内容"
    assert "<!-- header" not in editor.rich_editor.toPlainText()


def test_apply_header_footer_writes_sections(editor):
    """页眉/页脚必须落到 section 上，而不是段落里。"""
    from docx import Document

    editor._header_text = "机密文件"
    editor._footer_text = "第 X 页"

    doc = Document()
    assert editor._apply_header_footer(doc) is True

    section = doc.sections[0]
    assert section.header.paragraphs[0].text == "机密文件"
    assert section.footer.paragraphs[0].text == "第 X 页"
    # 正文段落不应包含页眉文字
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "机密文件" not in body


def test_apply_header_footer_noop_when_empty(editor):
    """没设置时不应改动文档。"""
    from docx import Document

    editor._header_text = ""
    editor._footer_text = ""
    assert editor._apply_header_footer(Document()) is False
