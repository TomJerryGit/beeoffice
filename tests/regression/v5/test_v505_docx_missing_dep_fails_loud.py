# tests/regression/v5/test_v505_docx_missing_dep_fails_loud.py
"""
优化规划 §3.3 回归：DOCX 导出的"缺依赖降级"实现已删除。

原状态
------
`bee_export.export_to_docx()` 原先是：

    if HAS_BS4 and HAS_MARKDOWN:
        ... BeautifulSoup 中间层（真正的实现）
    else:
        doc = Document()
        _write_plain_to_docx(content, doc)      # 逐行降级的第二套实现

而 `markdown` 与 `beautifulsoup4` 在 `requirements.txt` 里是**必需**依赖，
所以那个 else 分支实际不可达；`_write_plain_to_docx` / `_write_inline_md`
（约 55 行）也**从未有测试覆盖**。留着它有两个害处：
  1. 让人以为 Markdown→DOCX 有第二套实现（消重时的心智负担）；
  2. 真缺依赖时会**悄悄产出一份低保真文档**，用户以为导出成功了。

现在：缺依赖时明确失败 + 记录明确的日志原因，不产出文件。
"""
import pytest

from beeoffice.plugins import bee_export


@pytest.fixture
def minimal_md():
    return "# 标题\n\n正文段落，含 **加粗**。\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"


# ==================== 正常路径不变 ====================

def test_export_still_works_when_deps_present(minimal_md, tmp_path):
    out = tmp_path / "ok.docx"
    assert bee_export.export_to_docx(minimal_md, str(out)) is True
    assert out.exists() and out.stat().st_size > 0

    from docx import Document
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "标题" in text and "正文段落" in text


# ==================== 缺依赖：明确失败，不产出文件 ====================

@pytest.mark.parametrize("missing", ["HAS_BS4", "HAS_MARKDOWN"])
def test_missing_dep_fails_loud_and_writes_nothing(
    monkeypatch, minimal_md, tmp_path, missing
):
    monkeypatch.setattr(bee_export, missing, False)

    out = tmp_path / f"missing_{missing}.docx"
    result = bee_export.export_to_docx(minimal_md, str(out))

    assert result is False, (
        f"{missing}=False 时应明确失败，而不是降级产出一份低保真文档"
    )
    assert not out.exists(), "失败时不应留下文件（原实现会写出一份粗糙文档）"


# ==================== 降级实现确实已删除（防止长回来） ====================

def test_plain_fallback_functions_are_gone():
    assert not hasattr(bee_export, "_write_plain_to_docx"), (
        "`_write_plain_to_docx` 又出现了 —— 不可达的降级实现被加回来了？"
    )
    assert not hasattr(bee_export, "_write_inline_md"), (
        "`_write_inline_md` 只服务于上面那个降级实现，不应单独存在"
    )


def test_only_one_markdown_to_docx_renderer_remains():
    """
    Markdown 路径应只剩一套渲染实现：`_write_html_to_docx` 及其 helpers。
    """
    assert hasattr(bee_export, "_write_html_to_docx")
    assert hasattr(bee_export, "_write_element")
    assert hasattr(bee_export, "_write_table")
    assert hasattr(bee_export, "_write_inline")
