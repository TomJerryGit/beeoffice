# tests/regression/v5/test_v505_pptx_engine_fixes.py
"""
优化规划 §2.7 回归：PPTX 导出的三处缺陷。

原缺陷
------
1. **每页重启一次 Chromium**
   `ElementMeasurer._measure_playwright()` 每次调用都
   `with sync_playwright() as p: browser = p.chromium.launch(...)` … `browser.close()`，
   而 `PPTXEngine.export()` 是**逐页**调用 `measure()` 的 —— 一页一次浏览器冷启动，
   几十页的演示文稿光启动就要几十秒。
   修复：会话级复用（`start_session()` / `close_session()`），导出结束在 finally 里释放。

2. **相对路径图片一律加载失败**
   `page.set_content(full_html)` 没有 base URL，幻灯片里的 `<img src="media/x.png">`
   相对路径无从解析 → 图片在导出结果里整张消失。
   修复：`export(base_dir=...)` → `_wrap_html` 注入 `<base href="file:///…/">`。

3. **文本框没有字体名**
   `_add_textbox()` 只设了字号/粗斜/颜色，没有设 `run.font.name` ——
   PowerPoint 只能用主题默认字体渲染，中文尤其容易被换字形；
   "字体嵌入"（FontEmbedder）也因此没有可对应的字体名。
   修复：设置 `a:latin`（`font.name`）与 `a:ea`（东亚字体）。
"""
import sys
import types

import pytest

from beeoffice.plugins.beeshow_pptx_engine import (
    DEFAULT_CJK_FONT,
    ElementMeasurer,
    MeasuredElement,
    MeasuredSlide,
    PPTXEngine,
    PPTXGenerator,
)


# ==================== 1. 浏览器会话复用 ====================

class _FakePage:
    def __init__(self):
        self.closed = False

    def set_content(self, html, **kwargs):
        self.html = html

    def evaluate(self, script):
        # 第一个 evaluate 是等图片；MEASURE_JS 返回元素列表
        if "document.images" in script:
            return None
        if "backgroundColor" in script:
            return "rgb(255, 255, 255)"
        return []

    def wait_for_timeout(self, ms):
        pass

    def close(self):
        self.closed = True


class _FakeBrowser:
    def __init__(self, counter):
        self._counter = counter
        self.pages = []

    def new_page(self, **kwargs):
        page = _FakePage()
        self.pages.append(page)
        return page

    def close(self):
        self._counter["closed"] += 1


class _FakeChromium:
    def __init__(self, counter):
        self._counter = counter

    def launch(self, **kwargs):
        self._counter["launch"] += 1
        return _FakeBrowser(self._counter)


class _FakePlaywright:
    def __init__(self, counter):
        self.chromium = _FakeChromium(counter)

    def stop(self):
        self._counter_stopped = True


@pytest.fixture
def fake_playwright(monkeypatch):
    """注入一个假的 playwright.sync_api，统计 launch 次数。"""
    counter = {"launch": 0, "closed": 0}

    def _sync_playwright():
        return types.SimpleNamespace(
            chromium=_FakeChromium(counter),
            start=lambda: _FakePlaywright(counter),
        )

    sync_api = types.ModuleType("playwright.sync_api")
    sync_api.sync_playwright = _sync_playwright
    playwright_mod = types.ModuleType("playwright")
    playwright_mod.sync_api = sync_api

    monkeypatch.setitem(sys.modules, "playwright", playwright_mod)
    monkeypatch.setitem(sys.modules, "playwright.sync_api", sync_api)
    return counter


def test_browser_is_launched_once_for_all_slides(fake_playwright, tmp_path):
    """
    核心回归：3 页幻灯片只应启动 **1 次** Chromium（原实现是 3 次）。
    """
    engine = PPTXEngine(slide_width=10.0, slide_height=7.5)
    engine.measurer._playwright_available = True

    out = tmp_path / "deck.pptx"
    ok = engine.export(
        slides_html=["<h1>一</h1>", "<h1>二</h1>", "<h1>三</h1>"],
        output_path=str(out),
        theme_css="",
        bg_color="#FFFFFF",
    )

    assert ok is True
    assert counter_launches(fake_playwright) == 1, (
        f"Chromium 被启动了 {counter_launches(fake_playwright)} 次"
        f"（每页一次）；应整份导出只启动 1 次"
    )


def counter_launches(counter):
    return counter["launch"]


def test_session_is_released_after_export(fake_playwright, tmp_path):
    """导出结束后必须释放浏览器（否则留下 Chromium 进程）。"""
    engine = PPTXEngine()
    engine.measurer._playwright_available = True

    engine.export(
        slides_html=["<h1>一</h1>"],
        output_path=str(tmp_path / "a.pptx"),
    )

    assert engine.measurer._browser is None, "会话没有释放"
    assert fake_playwright["closed"] == 1


def test_session_released_even_when_measure_raises(fake_playwright, tmp_path):
    """测量抛异常时也要释放会话（finally 保证）。"""
    engine = PPTXEngine()
    engine.measurer._playwright_available = True

    def _boom(*a, **k):
        raise RuntimeError("模拟测量崩溃")

    engine.measurer._measure_playwright = _boom

    with pytest.raises(RuntimeError):
        engine.export(
            slides_html=["<h1>一</h1>"],
            output_path=str(tmp_path / "b.pptx"),
        )

    assert engine.measurer._browser is None, "异常路径下会话泄漏了"


def test_close_session_is_idempotent():
    """close_session 可重复调用，不应抛异常。"""
    m = ElementMeasurer()
    m.close_session()
    m.close_session()
    assert m._browser is None


# ==================== 2. base_url ====================

def test_wrap_html_injects_base_href(tmp_path):
    engine = PPTXEngine()
    href = engine._make_base_href(str(tmp_path), str(tmp_path / "o.pptx"))

    html = engine._wrap_html("<img src='a.png'>", "", "#FFF", href)

    assert "<base href=" in html, "没有注入 <base href>，相对路径图片会加载失败"
    assert href.startswith("file:///"), f"base 应为 file:/// 形式: {href}"
    assert href.endswith("/"), "base href 应以 / 结尾"


def test_wrap_html_without_base_is_unchanged():
    engine = PPTXEngine()
    html = engine._wrap_html("<p>x</p>", "", "#FFF")
    assert "<base" not in html


def test_make_base_href_falls_back_to_output_dir(tmp_path):
    engine = PPTXEngine()
    out = tmp_path / "sub" / "deck.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)

    href = engine._make_base_href(None, str(out))

    assert href, "没有 base_dir 时应退回输出目录"
    assert tmp_path.name in href or "sub" in href


def test_make_base_href_prefers_given_dir(tmp_path):
    engine = PPTXEngine()
    docs = tmp_path / "docs"
    docs.mkdir()
    href = engine._make_base_href(str(docs), str(tmp_path / "o.pptx"))
    assert "docs" in href, "应优先使用调用方给的基准目录"


def test_export_passes_base_dir_into_html(fake_playwright, tmp_path):
    """export 必须把 base_dir 转成 <base href> 传进页面 HTML。"""
    docs = tmp_path / "my_docs"
    docs.mkdir()

    engine = PPTXEngine()
    engine.measurer._playwright_available = True
    captured = {}

    real_measure = engine.measurer.measure

    def _spy(html):
        captured["html"] = html
        return real_measure(html)

    engine.measurer.measure = _spy

    engine.export(
        slides_html=["<img src='pic.png'>"],
        output_path=str(tmp_path / "o.pptx"),
        base_dir=str(docs),
    )

    assert "my_docs" in captured.get("html", ""), (
        f"页面 HTML 里没有 base href：{captured.get('html','')[:200]}"
    )


# ==================== 3. 字体名 ====================

def test_textbox_sets_latin_and_east_asian_font(tmp_path):
    """文本框必须同时设置 a:latin 与 a:ea，否则中文会用主题默认字体。"""
    from pptx import Presentation
    from pptx.oxml.ns import qn

    gen = PPTXGenerator()
    slide_data = MeasuredSlide(
        elements=[
            MeasuredElement(tag="h1", text="中文标题", font_size=32),
        ],
        bg_color="#FFFFFF",
    )
    out = tmp_path / "font.pptx"
    assert gen.generate([slide_data], str(out)) is True

    prs = Presentation(str(out))
    runs = [
        run
        for shape in prs.slides[0].shapes
        if shape.has_text_frame
        for para in shape.text_frame.paragraphs
        for run in para.runs
    ]
    assert runs, "没有生成文本框"
    run = runs[0]

    assert run.font.name == DEFAULT_CJK_FONT, (
        f"run.font.name 未设置（当前 {run.font.name!r}）——"
        f"PowerPoint 会用主题默认字体渲染"
    )

    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    assert ea is not None, "缺少 a:ea（东亚字体），中文会回退到主题字体"
    assert ea.get("typeface") == DEFAULT_CJK_FONT


def test_element_font_family_overrides_default(tmp_path):
    """测量端给出字体族时应优先使用它。"""
    from pptx import Presentation

    gen = PPTXGenerator()
    slide_data = MeasuredSlide(
        elements=[
            MeasuredElement(
                tag="p", text="正文", font_size=18,
                font_family="Consolas",
            ),
        ],
        bg_color="#FFFFFF",
    )
    out = tmp_path / "font2.pptx"
    assert gen.generate([slide_data], str(out)) is True

    prs = Presentation(str(out))
    run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Consolas"


# ==================== 导出调用点接线 ====================

def test_export_mixin_passes_base_dir():
    import inspect

    from beeoffice.plugins.beeshow_export_mixin import ExportMixin

    src = inspect.getsource(ExportMixin._export_pptx)
    assert "base_dir=" in src, (
        "导出 PPTX 时没有传 base_dir，相对路径图片仍会丢失"
    )
