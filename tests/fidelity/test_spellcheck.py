# tests/fidelity/test_spellcheck.py
"""
拼写检查核心模块保真度测试（``beeoffice/plugins/spellcheck.py``）

本机（Windows 10 19041）后端实测
--------------------------------
* **windows-native**：COM 服务器能调用（``CreateSpellChecker('zh-Latn-CN-x-ext')`` 成功），
  但 ``SupportedLanguages`` 只有 ``zh-Latn-CN-x-ext``，``IsSupported('en-US')`` = FALSE，
  ``CreateSpellChecker('en-US')`` = ``E_INVALIDARG(0x80070057)``；而那个中文检查器把
  ``hello/world/this/is/test`` **全部**判为拼写错误。→ 不能用于英文，必须降级。
* **pyspellchecker / pyenchant**：未安装（``ModuleNotFoundError``）。
* 结论：本机 ``detect_backend().available is False``，``check_text()`` 恒返回 ``[]``。

因此本文件分两层
----------------
1. **逻辑层**（后端无关）：用桩后端把 ``check_text`` / 忽略 / 自定义词典 / 建议
   跑满 —— 避免"后端不可用 ⇒ 测试全绿但其实什么都没测"。桩后端让每条断言都能
   在实现被改坏时失败。
2. **真后端层**：只有真后端能验的（如 ``teh`` → 建议含 ``the``）一律
   ``pytest.skip``，绝不假装通过。
"""
import re
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from beeoffice.plugins import spellcheck as sc  # noqa: E402

CORPUS_DIR = ROOT / "tests" / "fidelity" / "corpus"


# ====================================================================
# 工具 / 桩
# ====================================================================

def words(text):
    """只要词，丢位置。"""
    return [word for word, _, _ in sc.tokenize_words(text)]


def has_chinese(text):
    return any("\u4e00" <= char <= "\u9fff" for char in text or "")


class DictBackend(sc.SpellBackend):
    """小词典桩后端：只有词典里的词算正确。"""

    name = "stub-dict"

    def __init__(self, known):
        self.known = {word.lower() for word in known}

    def check(self, word):
        return word.lower() in self.known

    def suggest(self, word):
        return ["the"] if word.lower() == "teh" else []


class FlagAllBackend(sc.SpellBackend):
    """把所有词都判错的桩后端：用来暴露"本来不该被切出来的东西"。"""

    name = "stub-flag-all"

    def check(self, word):
        return False

    def suggest(self, word):
        return []


class ManySuggestionsBackend(sc.SpellBackend):
    """一次给 10 条建议，用来验证上限。"""

    name = "stub-many-suggestions"

    def check(self, word):
        return False

    def suggest(self, word):
        return ["s%d" % index for index in range(10)]


def unavailable(reason):
    return sc.UnavailableBackend(reason)


# ====================================================================
# 1. 后端探测
# ====================================================================

def test_detect_backend_never_raises():
    backend = sc.detect_backend()
    assert isinstance(backend, sc.SpellBackend)
    assert isinstance(backend.available, bool)
    assert backend.name


def test_detect_backend_unavailable_has_chinese_reason():
    """不可用时必须给出**非空中文**原因（不允许静默失败）。"""
    backend = sc.detect_backend()
    if backend.available:
        pytest.skip("本机有可用后端（%s），降级原因分支无法验证" % backend.name)
    assert backend.available is False
    assert isinstance(backend.reason, str) and backend.reason.strip()
    assert has_chinese(backend.reason), backend.reason


def test_backend_priority_order_is_windows_then_pyspellchecker_then_enchant():
    """优先级顺序是接口的一部分：改顺序必须让这条测试失败。"""
    assert [name for name, _ in sc._BACKEND_PROBES] == [
        "windows-native", "pyspellchecker", "enchant",
    ]


def test_detect_backend_prefers_first_available():
    calls = []

    def first():
        calls.append("first")
        return unavailable("第一候选不可用：测试用原因")

    def second():
        calls.append("second")
        return DictBackend(["the"])

    def third():
        calls.append("third")
        return DictBackend(["the"])

    backend = sc.detect_backend([("first", first), ("second", second), ("third", third)])
    assert backend.name == "stub-dict"
    assert calls == ["first", "second"], "命中后不该继续探测后面的后端"


def test_detect_backend_collects_every_reason():
    backend = sc.detect_backend([
        ("甲", lambda: unavailable("甲后端不可用：原因一")),
        ("乙", lambda: unavailable("乙后端不可用：原因二")),
    ])
    assert backend.available is False
    assert "原因一" in backend.reason and "原因二" in backend.reason, backend.reason
    assert has_chinese(backend.reason)


def test_detect_backend_survives_probe_exception():
    def boom():
        raise RuntimeError("探测炸了")

    backend = sc.detect_backend([("bad", boom)])
    assert backend.available is False
    assert "RuntimeError" in backend.reason and has_chinese(backend.reason)

    # 探测异常不该阻断后续候选
    backend2 = sc.detect_backend([("bad", boom), ("good", lambda: DictBackend(["the"]))])
    assert backend2.name == "stub-dict"


def test_detect_backend_handles_probe_returning_none():
    backend = sc.detect_backend([("none", lambda: None)])
    assert backend.available is False
    assert has_chinese(backend.reason)


def test_unavailable_backend_never_flags_words():
    """不可用后端必须"什么都不报"，否则整篇文档会被标红。"""
    backend = unavailable("没有可用后端（测试）")
    assert backend.check("teh") is True
    assert backend.suggest("teh") == []


# ====================================================================
# 2. Windows 原生后端的语言选择策略（本机失效的真实原因就在这里）
# ====================================================================

def test_select_windows_language_picks_english_that_passes_sanity():
    tags = ["zh-Latn-CN-x-ext", "en-US"]

    def usable(tag):
        return tag == "en-US"

    assert sc._select_windows_language(tags, usable) == ("en-US", "")


def test_select_windows_language_rejects_english_tag_that_flags_correct_words():
    """
    **本机真实情况**：只有 ``zh-Latn-CN-x-ext``，且它把 hello/world/the 全判错。

    策略必须因此拒绝它并给出中文原因 —— 若把"语言存在"当成"可用"，这条会失败。
    """
    tag, reason = sc._select_windows_language(["zh-Latn-CN-x-ext"], lambda _tag: False)
    assert tag is None
    assert reason and has_chinese(reason)
    assert "没有英文" in reason


def test_select_windows_language_reports_english_pack_present_but_unusable():
    tag, reason = sc._select_windows_language(["en-US", "zh-Latn-CN-x-ext"], lambda _tag: False)
    assert tag is None
    assert "en-US" in reason and has_chinese(reason)
    assert "未通过自检" in reason


def test_select_windows_language_accepts_non_english_tag_that_really_checks_english():
    """判据是**功能自检**，不是标签名：非 en 标签若真能查英文也接受。"""
    assert sc._select_windows_language(["de-DE"], lambda _tag: True) == ("de-DE", "")


def test_select_windows_language_with_no_languages():
    tag, reason = sc._select_windows_language([], lambda _tag: True)
    assert tag is None
    assert has_chinese(reason)


def test_windows_native_probe_is_honest_about_this_machine():
    """
    真机探测：要么给出通过自检的原生后端，要么给出**中文原因**拒绝。

    本机上 ``_probe_windows_native()`` 拿到的 zh-Latn 检查器会把正确英文词判错，
    所以断言"拒绝"。若自检被删掉，这条会失败（会返回 available=True 但自检为假）。
    """
    backend = sc._probe_windows_native()
    assert isinstance(backend, sc.SpellBackend)
    if backend.available:
        assert backend.name == "windows-native"
        assert backend._checker.is_usable_for_english() is True
    else:
        assert has_chinese(backend.reason), backend.reason


def test_is_english_language_tag():
    assert sc._is_english_language_tag("en-US")
    assert sc._is_english_language_tag("en_GB")
    assert sc._is_english_language_tag("en")
    assert not sc._is_english_language_tag("zh-Latn-CN-x-ext")
    assert not sc._is_english_language_tag("eng")


# ====================================================================
# 3. 词法：每条"必须跳过"的规则，都同时验正例与反例
# ====================================================================
# 每条测试的写法：断言"该跳过的没出现"**且**"旁边的错词仍在" ——
# 只验前者会被"什么都不返回"的坏实现骗过去。

def test_tokenize_english_words_have_exact_positions():
    text = "The quick brown fox"
    assert sc.tokenize_words(text) == [
        ("The", 0, 3), ("quick", 4, 9), ("brown", 10, 15), ("fox", 16, 19),
    ]


def test_tokenize_skips_chinese_but_keeps_english_typo():
    assert words("中文测试内容") == []
    assert words("中文测试，标点。 qwick") == ["qwick"]
    assert words("中文与english混排") == []          # 紧贴中文 → 跳过
    assert words("中文与 english 混排") == ["english"]  # 有空格分隔 → 检查


def test_tokenize_skips_chinese_punctuation():
    assert words("中文，。！？；：（）【】「」") == []
    assert words("中文（括号） qwick") == ["qwick"]


def test_tokenize_skips_fullwidth_latin_letters():
    """全角字母（ＡＢＣ）不是英文单词；若正则改成 Unicode 感知就会被报出来。"""
    assert words("ＡＢＣ ａｂｃ ＯＫ") == []
    assert words("全角ＡＢＣ qwick") == ["qwick"]


def test_tokenize_skips_numbers_decimals_percent_and_thousands():
    assert words("1,234.5") == []
    assert words("12%") == []
    assert words("共 1,234.5 元，增长 12%") == []
    assert words("共 1,234.5 元 qwick") == ["qwick"]
    assert words("数量 1234 与 0.5 与 1,000") == []


def test_number_spans_are_masked_out_before_matching():
    """
    数字区间在匹配前被显式换成空格（``_mask_ignored`` 的直接验证）。

    注：拉丁词正则只匹配 ``[A-Za-z]``，所以数字**本来就**不会变成"词" ——
    这条钉住的是"数字区间被显式屏蔽 + 屏蔽不改变下标"这一实现事实。
    """
    text = "共 1,234.5 元"
    masked = sc._mask_ignored(text)
    assert len(masked) == len(text), "屏蔽必须保持长度，否则位置会错位"
    assert "1" not in masked and "234" not in masked and "." not in masked
    assert masked.startswith("共 ") and masked.endswith(" 元")
    assert sc._mask_ignored("12%") == "   "


def test_tokenize_skips_version_numbers():
    assert words("5.0.2") == []
    assert words("2021.10.20") == []
    assert words("V5.0.2") == []
    assert words("version 5.0.2 qwick") == ["version", "qwick"]


def test_tokenize_skips_letter_runs_glued_to_digits():
    """紧贴数字的字母片段不是英语单词（``Windows10`` / ``python3`` / ``ipv4``）。"""
    assert words("Windows10 python3 ipv4 grub2 ens33") == []
    assert words("Windows 10 与 python 3 与 grub 2") == ["Windows", "python", "grub"]


def test_tokenize_skips_urls_and_emails_but_keeps_typos_around_them():
    text = "访问 https://a.com 或 a@b.com 了解 qwick"
    assert words(text) == ["qwick"]
    assert "com" not in words(text) and "https" not in words(text)
    assert words("www.example.com") == []
    assert words("ftp://files.example.org/path") == []
    assert words("https://a.com teh") == ["teh"]


def test_tokenize_skips_markdown_markers():
    assert words("**bold**") == ["bold"]
    assert words("# 标题\n| 表格 | 列 |") == []
    assert words("**qwick** # 标题 | 单元格 |") == ["qwick"]


def test_tokenize_skips_inline_code_and_fenced_code_blocks():
    assert words("`teh qwick` 与普通 teh") == ["teh"]
    assert words("使用 `code_snippet` 与 print() 输出") == ["print"]
    assert words("```\nteh qwick\n```\nreal teh") == ["real", "teh"]
    assert words("~~~\nteh qwick\n~~~") == []
    assert words("```\nteh qwick") == []            # 未闭合围栏：一直到文末都跳过


def test_tokenize_skips_identifiers():
    assert words("sheet_name getValue iCounter MyClass") == []
    assert words("sheet_name qwick") == ["qwick"]
    assert words("this_is_a_long_snake_case_name") == []
    assert words("_italic_ 效果") == ["italic"]      # markdown 斜体不是 snake_case
    assert words("a_b") == []


def test_tokenize_skips_acronyms_and_camel_case_but_keeps_capitalised_words():
    assert words("README IPADDR GPL CPU") == []
    assert words("CentOS MySQL IPv4 WinSCP") == []
    assert words("Linux Windows README qwick") == ["Linux", "Windows", "qwick"]


def test_tokenize_skips_short_words():
    """固定规则：短于 3 个字母的词不检查（``a`` / ``ok`` / ``is`` 不报，``teh`` 要报）。"""
    assert sc.MIN_WORD_LENGTH == 3
    assert words("a ok is of an") == []
    assert words("a ok is teh") == ["teh"]


def test_tokenize_skips_contraction_parts():
    assert words("don't worry") == ["worry"]
    assert words("it's a test") == ["test"]
    assert words("'hello' there") == ["hello", "there"]
    assert words("don't qwick") == ["qwick"]


def test_tokenize_reports_misspelled_words_and_their_positions():
    text = "Teh qwick brown fox recieve"
    found = sc.tokenize_words(text)
    assert [word for word, _, _ in found] == ["Teh", "qwick", "brown", "fox", "recieve"]
    for word, start, end in found:
        assert text[start:end] == word
        assert end - start == len(word)


def test_tokenize_handles_empty_and_none():
    assert sc.tokenize_words("") == []
    assert sc.tokenize_words(None) == []


def test_tokenize_positions_stay_consistent_in_a_mixed_document():
    text = (
        "# 配置说明\n"
        "编辑 `ifcfg-ens33` 文件：BOOTPROTO=static\n"
        "```\nIPADDR=192.168.1.100\n```\n"
        "然后执行 systemctl restart network，如果报错请看 https://a.com 或写信 a@b.com。\n"
        "常见拼写错误：accomodate、recieve、seperate。\n"
    )
    found = sc.tokenize_words(text)
    assert all(isinstance(item, tuple) and len(item) == 3 for item in found)
    for word, start, end in found:
        assert text[start:end] == word
        assert re.fullmatch(r"[A-Za-z]+", word)
        assert len(word) >= sc.MIN_WORD_LENGTH
    assert words(text) == [
        "static", "systemctl", "restart", "network",
        "accomodate", "recieve", "seperate",
    ]


# ====================================================================
# 4. 检查服务（用桩后端把逻辑跑满）
# ====================================================================

def test_service_reports_misspelling_with_position_and_suggestions():
    service = sc.SpellCheckerService(backend=DictBackend(["the", "word", "is"]))
    found = service.check_text("the word is teh")
    assert len(found) == 1
    miss = found[0]
    assert (miss.word, miss.start, miss.end) == ("teh", 12, 15)
    assert "the" in miss.suggestions, "建议没接上（写好又被清空时这条会失败）"


def test_service_reports_every_occurrence_with_its_own_position():
    service = sc.SpellCheckerService(backend=DictBackend(["the", "and"]))
    found = service.check_text("teh and teh")
    assert [(miss.word, miss.start, miss.end) for miss in found] == [("teh", 0, 3), ("teh", 8, 11)]


def test_service_ignores_chinese_text_entirely():
    service = sc.SpellCheckerService(backend=FlagAllBackend())
    assert service.check_text("这是一段纯中文，包含标点：。，！？（）「」，全角ＡＢＣ。") == []


def test_service_can_disable_suggestions():
    service = sc.SpellCheckerService(backend=DictBackend(["the"]), suggest=False)
    found = service.check_text("teh")
    assert [miss.word for miss in found] == ["teh"]
    assert found[0].suggestions == []


def test_service_caps_suggestion_count():
    service = sc.SpellCheckerService(backend=ManySuggestionsBackend(), max_suggestions=2)
    found = service.check_text("teh")
    assert len(found) == 1
    assert len(found[0].suggestions) == 2


def test_service_ignore_is_case_insensitive_and_persists_for_the_session():
    service = sc.SpellCheckerService(backend=DictBackend(["the"]))
    assert [miss.word for miss in service.check_text("teh the")] == ["teh"]
    service.ignore("TEH")
    assert service.check_text("teh the") == []
    assert service.knows("teh") is True


def test_service_add_to_dictionary_is_case_insensitive():
    service = sc.SpellCheckerService(backend=DictBackend(["the"]))
    service.add_to_dictionary("Qwick")
    assert service.check_text("qwick the") == []
    assert service.knows("QWICK") is True


def test_service_dictionary_argument_is_applied():
    service = sc.SpellCheckerService(backend=DictBackend(["the"]), dictionary=["Qwick"])
    assert service.check_text("qwick the") == []


def test_service_check_word():
    service = sc.SpellCheckerService(backend=DictBackend(["the"]))
    assert service.check_word("the") is True
    assert service.check_word("teh") is False
    assert service.check_word("") is True
    service.ignore("teh")
    assert service.check_word("teh") is True


def test_service_unavailable_backend_reports_nothing():
    """降级行为：不可用时 ``available is False``、中文 ``reason``、``check_text`` 空。"""
    service = sc.SpellCheckerService(backend=unavailable("没有可用的拼写检查后端（测试用原因）"))
    assert service.available is False
    assert service.reason == "没有可用的拼写检查后端（测试用原因）"
    assert has_chinese(service.reason)
    assert service.check_text("teh qwick recieve 中文") == []
    assert service.check_word("teh") is True


def test_service_default_backend_matches_detect_backend():
    service = sc.SpellCheckerService()
    backend = sc.detect_backend()
    assert service.available is backend.available
    assert service.backend_name == backend.name
    if not service.available:
        assert service.check_text("teh qwick recieve") == []


def test_service_survives_backend_exceptions():
    class BrokenBackend(sc.SpellBackend):
        name = "stub-broken"

        def check(self, word):
            raise ValueError("后端炸了")

        def suggest(self, word):
            raise ValueError("后端炸了")

    service = sc.SpellCheckerService(backend=BrokenBackend())
    assert service.check_text("teh qwick") == []      # 出错时宁可不报
    assert service.check_word("teh") is True


# ====================================================================
# 5. 自定义词典持久化
# ====================================================================

def test_user_dictionary_roundtrip(tmp_path):
    path = tmp_path / "user_dict.txt"
    assert sc.save_user_dictionary(path, ["BeeOffice", "spellcheck", "BeeOffice", "  ", "TEH"]) is True
    assert path.exists()
    assert sc.load_user_dictionary(path) == {"beeoffice", "spellcheck", "teh"}

    raw = path.read_text(encoding="utf-8")
    lines = raw.splitlines()
    assert lines[0].startswith("#") and has_chinese(lines[0])
    body = [line for line in lines if line and not line.startswith("#")]
    assert body == sorted(body), "写出的词表应排序，便于版本管理与人工核对"


def test_load_user_dictionary_tolerates_bom_comments_and_blank_lines(tmp_path):
    path = tmp_path / "dict_bom.txt"
    path.write_text("# 注释\n\nbeeoffice\n  BeeTool  \n# 又一行注释\nbeeoffice\n", encoding="utf-8-sig")
    assert sc.load_user_dictionary(path) == {"beeoffice", "beetool"}


def test_load_user_dictionary_missing_file_returns_empty_set(tmp_path):
    assert sc.load_user_dictionary(tmp_path / "不存在.txt") == set()


def test_user_dictionary_file_takes_effect_in_service(tmp_path):
    path = tmp_path / "user_dict.txt"
    assert sc.save_user_dictionary(path, ["qwick"]) is True

    service = sc.SpellCheckerService(backend=DictBackend(["the"]))
    assert [miss.word for miss in service.check_text("qwick the")] == ["qwick"]

    assert service.load_dictionary(path) == 1
    assert service.check_text("qwick the") == [], "重新加载的词典没有生效"


def test_service_save_dictionary_writes_ignored_and_dictionary_words(tmp_path):
    service = sc.SpellCheckerService(backend=DictBackend(["the"]))
    service.ignore("teh")
    service.add_to_dictionary("qwick")
    path = tmp_path / "saved.txt"

    assert service.save_dictionary(path) is True
    assert sc.load_user_dictionary(path) == {"teh", "qwick"}
    reloaded = sc.SpellCheckerService(backend=DictBackend(["the"]), dictionary=sc.load_user_dictionary(path))
    assert reloaded.check_text("teh qwick the") == []


# ====================================================================
# 6. 建议（只有真后端可用才断言）
# ====================================================================

def test_real_backend_suggests_the_for_teh():
    backend = sc.detect_backend()
    if not backend.available:
        pytest.skip("本机没有可用的拼写后端（%s），无法验证真实建议" % backend.reason[:60])
    service = sc.SpellCheckerService(backend=backend)
    found = service.check_text("teh")
    assert len(found) == 1
    assert any(suggestion.lower() == "the" for suggestion in found[0].suggestions), found[0].suggestions


def test_stub_suggestion_path_is_exercised_without_a_real_backend():
    """与上一条等价的逻辑层版本：保证"建议"这条链路任何时候都被测到。"""
    service = sc.SpellCheckerService(backend=DictBackend(["the"]))
    found = service.check_text("teh")
    assert found and found[0].suggestions == ["the"]


# ====================================================================
# 7. 真实中文语料（tests/fidelity/corpus/*.docx）
# ====================================================================

@pytest.fixture(scope="module")
def corpus_texts():
    if not CORPUS_DIR.is_dir():
        pytest.skip("没有语料目录: %s" % CORPUS_DIR)
    files = sorted(CORPUS_DIR.glob("*.docx"))
    if not files:
        pytest.skip("语料目录里没有 docx")
    from docx import Document

    documents = []
    for path in files:
        document = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += "\n" + cell.text
        documents.append((path.name, text))
    return documents


def test_corpus_tokenization_invariants(corpus_texts):
    """
    对**真实中文文档**跑词法，逐条从**原文**独立复核不变量。

    这些断言独立于实现内部状态：只要哪条跳过规则失效（比如不再跳过中文紧邻、
    不再跳过驼峰/缩写、不再跳过数字紧邻），就有词违反不变量 → 立刻失败。
    """
    raw_runs = 0
    total = 0
    for name, text in corpus_texts:
        raw_runs += len([run for run in re.findall(r"[A-Za-z]+", text) if len(run) >= 3])
        for word, start, end in sc.tokenize_words(text):
            total += 1
            assert text[start:end] == word, (name, word)
            assert 0 <= start < end <= len(text), (name, word)
            assert len(word) >= sc.MIN_WORD_LENGTH, (name, word)
            assert re.fullmatch(r"[a-z]+", word) or re.fullmatch(r"[A-Z][a-z]+", word), (
                "%s: 报出非普通英文词形的词 %r" % (name, word))
            before = text[start - 1] if start > 0 else ""
            after = text[end] if end < len(text) else ""
            assert not before.isdigit() and not after.isdigit(), (name, word)
            assert "_" not in (before, after), (name, word)
            assert not (before and sc._CJK_WORD_RE.match(before)), (name, word)
            assert not (after and sc._CJK_WORD_RE.match(after)), (name, word)
            assert before not in "'\u2019" or start < 2 or not text[start - 2].isalpha(), (name, word)
    assert total > 0, "真实语料里一个可检查的拉丁词都没切出来 —— 词法切分坏了"
    assert total < raw_runs, "词法规则没起作用：可检查词数等于原始拉丁片段数"


def test_corpus_chinese_only_text_is_never_flagged(corpus_texts):
    """
    中文零误报：把真实文档里的拉丁字母全部抹掉，剩下的纯中文内容**一个都不能报**。

    用"把所有词都判错"的桩后端，所以任何被切出来的东西都会报 —— 这条只要失败
    就说明中文（或全角字母/数字）被当成英文单词了。
    """
    service = sc.SpellCheckerService(backend=FlagAllBackend())
    checked = 0
    for name, text in corpus_texts:
        chinese_only = re.sub(r"[A-Za-z]+", "", text)
        assert service.check_text(chinese_only) == [], name
        checked += len(chinese_only)
    assert checked > 5000, checked


def test_corpus_documents_have_no_latin_false_positives(corpus_texts):
    """
    对 7 份真实中文文档跑整体检查。

    * 后端**可用**时：断言报出的每一处都是纯拉丁词、位置自洽，并且后端不会把
      **每一个**可检查词都判错（那正是本机 zh-Latn 检查器的失效模式）。
      ⚠️ 本机后端不可用，所以"误报数很少"这一条**没有实测**，这里不编数字。
    * 后端**不可用**时（本机实测结果）：``available is False`` + ``check_text`` 恒为 ``[]``。
    """
    service = sc.SpellCheckerService()
    total_tokens = 0
    total_errors = 0
    for name, text in corpus_texts:
        tokens = sc.tokenize_words(text)
        total_tokens += len(tokens)
        for miss in service.check_text(text):
            assert text[miss.start:miss.end] == miss.word, (name, miss)
            assert re.fullmatch(r"[A-Za-z]+", miss.word), ("%s: 报出非拉丁词" % name, miss.word)
            assert isinstance(miss.suggestions, list)
            total_errors += 1

    assert total_tokens > 0, "语料里没有可检查的拉丁词 —— 词法切分坏了"

    if not service.available:
        assert service.available is False
        assert has_chinese(service.reason)
        assert total_errors == 0, "后端不可用时不该报任何错"
        return

    assert total_errors < total_tokens, (
        "后端把语料里每一个拉丁词都判成错误（检查器与语言不匹配）：%d/%d"
        % (total_errors, total_tokens))


def test_corpus_live_pipeline_reports_when_backend_flags_everything(corpus_texts):
    """
    反向确认：同一批语料在"把所有词都判错"的桩后端下**必须报出东西**。

    没有这条，上面那个 ``check_text(...) == []`` 可能是"管道整个是空的"造成的假绿。
    """
    service = sc.SpellCheckerService(backend=FlagAllBackend(), suggest=False)
    reported = 0
    for name, text in corpus_texts:
        found = service.check_text(text)
        assert len(found) == len(sc.tokenize_words(text)), name
        reported += len(found)
    assert reported > 100, reported
