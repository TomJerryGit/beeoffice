# tests/fidelity/test_r7_media_fallback.py
"""
R7：图片回退目录**全都不可写**时必须**快速失败**，而不是每张图各失败一次。

## 背景（为什么值得单独修）
`_media_cache_root()` 原来在"所有候选都不可写"时**兜底返回
`tempfile.gettempdir()`** —— 而它本身可能也不可写（受限账户 / 磁盘满 / 沙箱）。
于是调用方拿到一个**已知不可写**的目录，接着**每张图各失败一次**：
23 张图 = 23 次异常，日志一片 error，真正的原因反而看不清。

改成：候选全部不可写 → 返回 `None` → `_resolve_media_dir` 返回 `(None, "")`
→ `import_from_docx` **一次性放弃提取图片**，正文照常导入，
并给用户**一条明确说明**（而不是 23 行失败）。

## 怎么造"全都不可写"（确定性，不依赖权限）
把一个**文件**放在本该是目录的位置：`mkdir(parents=True, exist_ok=True)`
对"已存在的文件"会抛 `FileExistsError`（`exist_ok` 只忽略已存在的**目录**）。
"""
import sys
from pathlib import Path

import io

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from beeoffice.plugins import bee_export                            # noqa: E402
from beeoffice.plugins.bee_export import (                          # noqa: E402
    _media_cache_root, _resolve_media_dir, import_from_docx,
)


def _make_docx_with_image(path: Path, images: int = 3) -> Path:
    """造一份带若干图片的 docx（用 Pillow 生成**真** PNG）。"""
    from docx import Document
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 30, 30)).save(buffer, format="PNG")
    png = buffer.getvalue()

    doc = Document()
    doc.add_paragraph("正文第一段")
    for _ in range(images):
        doc.add_paragraph().add_run().add_picture(io.BytesIO(png))
    doc.add_paragraph("正文最后一段")
    doc.save(str(path))
    return path


@pytest.fixture
def nothing_writable(tmp_path, monkeypatch):
    """
    让**所有**图片目录候选都不可写：
      · 文档同级：放一个同名文件挡住 `<stem>_media`
      · 缓存候选（LOCALAPPDATA / XDG_CACHE_HOME / ~/.cache / tempdir）：
        全部指向一个"是文件不是目录"的路径
    """
    blocker = tmp_path / "blocker"
    blocker.write_text("我是文件，不是目录", encoding="utf-8")

    monkeypatch.setenv("LOCALAPPDATA", str(blocker))
    monkeypatch.setenv("XDG_CACHE_HOME", str(blocker))
    monkeypatch.setattr(Path, "home", staticmethod(lambda: blocker))
    import tempfile

    monkeypatch.setattr(tempfile, "gettempdir", lambda: str(blocker))
    return blocker


# ==================== 1. 单点：候选全不可写 → None ====================

def test_media_cache_root_returns_none_when_nothing_writable(nothing_writable):
    """
    **核心回归**：不能再用 `tempfile.gettempdir()` 兜底返回一个不可写目录。

    （那条兜底正是"23 张图各失败一次"的根源。）
    """
    assert _media_cache_root() is None


def test_resolve_media_dir_returns_none_when_nothing_writable(
        nothing_writable, tmp_path):
    docx = tmp_path / "样本.docx"
    docx.write_bytes(b"x")
    (tmp_path / "样本_media").write_text("挡住同级目录", encoding="utf-8")

    media_dir, prefix = _resolve_media_dir(docx)

    assert media_dir is None, f"应为 None，实际 {media_dir}"
    assert prefix == ""


def test_resolve_media_dir_still_prefers_adjacent_when_writable(tmp_path):
    """能写的时候行为不变（同级目录 + 相对链接）—— 不能为了修异常情况改坏常规路径。"""
    docx = tmp_path / "样本.docx"
    docx.write_bytes(b"x")

    media_dir, prefix = _resolve_media_dir(docx)

    assert media_dir == tmp_path / "样本_media"
    assert prefix == "样本_media/"


# ==================== 2. 端到端：正文要导入，图片一次性放弃 ====================

def test_import_keeps_text_and_gives_one_clear_notice(
        nothing_writable, tmp_path, caplog):
    """
    **验收**：无可用图片目录时
      · 正文照常导入（不能因为图片问题把文字也丢了）
      · **不再有图片链接**（`![](...)`）
      · 只给**一条**明确说明（而不是每张图一行失败）
    """
    docx = _make_docx_with_image(tmp_path / "带图.docx", images=3)
    (tmp_path / "带图_media").write_text("挡住同级", encoding="utf-8")

    with caplog.at_level("WARNING"):
        md = import_from_docx(str(docx))

    assert "正文第一段" in md, "正文丢了"
    assert "正文最后一段" in md, "正文丢了"
    assert "![" not in md, f"不该有图片链接：{md[:200]}"

    notices = [
        line for line in md.splitlines()
        if "图片未能提取" in line
    ]
    assert len(notices) == 1, f"说明应只有一条，实际 {len(notices)}：{notices}"

    # 日志里应当**只出现一次**"无可用图片目录"，而不是每张图一次
    fallback_warnings = [
        r for r in caplog.records if "无可用图片目录" in r.getMessage()
    ]
    assert len(fallback_warnings) == 1, (
        f"应只警告一次，实际 {len(fallback_warnings)} 次"
    )
    per_image_errors = [
        r for r in caplog.records if "图片" in r.getMessage()
        and "失败" in r.getMessage()
    ]
    assert not per_image_errors, (
        "不应出现『每张图各失败一次』的日志："
        f"{[r.getMessage() for r in per_image_errors]}"
    )


def test_import_with_writable_fallback_still_embeds_images(tmp_path, monkeypatch):
    """对照组：缓存目录**可写**时图片仍要正常提取（别把正常路径修坏了）。"""
    docx = _make_docx_with_image(tmp_path / "带图.docx", images=2)
    # 挡住同级目录，迫使走缓存回退
    (tmp_path / "带图_media").write_text("挡住同级", encoding="utf-8")
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "cache"))
    monkeypatch.setattr(Path, "home", staticmethod(lambda: tmp_path / "home"))
    import tempfile

    monkeypatch.setattr(tempfile, "gettempdir", lambda: str(tmp_path / "temp"))

    md = import_from_docx(str(docx))

    assert "![" in md, f"可写时应当提取出图片链接：{md[:200]}"
    assert "图片未能提取" not in md
    written = list((tmp_path / "cache").rglob("image_*"))
    assert written, "缓存目录里应有图片文件"


def test_document_without_images_is_unaffected(nothing_writable, tmp_path):
    """没有图片的文档：不该出现"图片未能提取"的说明。"""
    from docx import Document

    docx = tmp_path / "纯文字.docx"
    doc = Document()
    doc.add_paragraph("只有文字")
    doc.save(str(docx))

    md = import_from_docx(str(docx))

    assert "只有文字" in md
    assert "图片未能提取" not in md, "没有图片却提示图片问题"
