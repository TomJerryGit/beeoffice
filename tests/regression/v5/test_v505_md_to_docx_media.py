# tests/regression/v5/test_v505_md_to_docx_media.py
"""
优化规划 §2.6 回归：Markdown→DOCX 不得静默丢弃图片与超链接。

原缺陷
------
`bee_export.py` 的 HTML→DOCX 转换器**完全没有 `<img>` 分支**，而 `_write_inline`
又有"文本为空就跳过"的前置判断 —— `<img>` 的 `get_text()` 恒为空，于是：

    ![说明](img.png)   →  导出的 Word 里什么都没有（静默丢弃）

超链接则只被**染色加下划线**（`run.font.color` / `run.font.underline`），
`href` 直接丢失：导出的 Word 里链接既不可点击，也看不出指向哪里。

修复：
  - `_write_inline_image()` 用 `run.add_picture()` 真正嵌入图片，相对路径按
    Markdown 所在目录解析；
  - 找不到文件/不支持的格式 → 插入**可见占位文字**（而不是静默丢弃）；
  - `_add_hyperlink()` 拼 `w:hyperlink` + 外部关系，生成可点击链接。
"""
import pytest

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:  # pragma: no cover
    HAS_DOCX = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:  # pragma: no cover
    HAS_PIL = False

pytestmark = pytest.mark.skipif(
    not HAS_DOCX, reason="python-docx 未安装"
)

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _export(md_text, tmp_path, name="out.docx", base_dir=None):
    from beeoffice.plugins.bee_export import export_to_docx

    out = tmp_path / name
    ok = export_to_docx(
        md_text, str(out),
        base_dir=str(base_dir if base_dir is not None else tmp_path),
    )
    assert ok is True, "导出应成功"
    assert out.exists() and out.stat().st_size > 0
    return out


def _document_text(path):
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _count_drawings(path):
    doc = Document(str(path))
    return len(doc.element.body.findall(f".//{_W_NS}drawing"))


@pytest.fixture
def sample_png(tmp_path):
    if not HAS_PIL:
        pytest.skip("Pillow 未安装")
    path = tmp_path / "img.png"
    Image.new("RGB", (300, 200), color="green").save(str(path))
    return path


# ==================== 图片 ====================

def test_image_is_embedded_in_docx(tmp_path, sample_png):
    """相对路径的图片必须真正嵌入（原实现静默丢弃）。"""
    out = _export(f"# 标题\n\n![说明]({sample_png.name})\n", tmp_path)
    assert _count_drawings(out) >= 1, (
        "导出的 DOCX 里没有图片（<img> 被静默丢弃）"
    )


def test_image_in_own_paragraph(tmp_path, sample_png):
    """只含图片的段落也不能被跳过。"""
    out = _export(f"![只有图]({sample_png.name})", tmp_path)
    assert _count_drawings(out) >= 1


def test_missing_image_produces_visible_placeholder(tmp_path):
    """找不到图片时要留下可见痕迹，而不是当无事发生。"""
    out = _export("![缺失的图](not_here.png)", tmp_path)
    assert _count_drawings(out) == 0
    assert "图片未导出" in _document_text(out), (
        "图片丢失了却没有任何提示，用户不会知道导出不完整"
    )


def test_remote_image_is_not_downloaded_but_flagged(tmp_path):
    """http(s) 图片不联网下载，但必须给出可见占位。"""
    out = _export("![网络图](https://example.com/x.png)", tmp_path)
    assert _count_drawings(out) == 0
    assert "图片未导出" in _document_text(out)


def test_absolute_path_image(tmp_path, sample_png):
    """
    绝对路径同样可用。

    用 `as_posix()` 而不是 `str()`：Markdown/HTML 里的图片地址是 URL 风格、
    使用正斜杠；反斜杠写在 Markdown 里本身也不合法。
    """
    out = _export(f"![绝对]({sample_png.as_posix()})", tmp_path)
    assert _count_drawings(out) >= 1


def test_file_url_image(tmp_path, sample_png):
    """file:/// 形式的本地地址要能解析（去掉盘符前的斜杠）。"""
    url = sample_png.as_uri()          # file:///C:/...
    out = _export(f"![file url]({url})", tmp_path)
    assert _count_drawings(out) >= 1


def test_image_resolved_against_base_dir(tmp_path):
    """相对路径应相对 Markdown 所在目录解析，而不是当前工作目录。"""
    if not HAS_PIL:
        pytest.skip("Pillow 未安装")
    md_dir = tmp_path / "docs"
    md_dir.mkdir()
    img = md_dir / "pic.png"
    Image.new("RGB", (120, 80), color="blue").save(str(img))

    out = tmp_path / "out.docx"
    from beeoffice.plugins.bee_export import export_to_docx

    assert export_to_docx(
        "![相对](pic.png)", str(out), base_dir=str(md_dir)
    ) is True
    assert _count_drawings(out) >= 1, "base_dir 未生效，相对路径没解析到"


# ==================== 超链接 ====================

def test_hyperlink_is_real_and_clickable(tmp_path):
    """必须生成 w:hyperlink + 外部关系，而不是只染个颜色。"""
    url = "https://example.com/some/page"
    out = _export(f"[点这里]({url})", tmp_path)

    doc = Document(str(out))
    # 用元素查找而不是字符串包含：序列化后是 `<w:hyperlink>`（有前缀声明），
    # 永远不会出现 `{命名空间}hyperlink` 这种 Clark 记法，字符串断言会假失败。
    hyperlinks = doc.element.body.findall(f".//{_W_NS}hyperlink")
    assert hyperlinks, "没有生成 w:hyperlink（href 丢失）"

    # 外部关系必须指向原 URL
    targets = [
        rel.target_ref for rel in doc.part.rels.values()
        if getattr(rel, "is_external", False)
    ]
    assert url in targets, f"未建立指向 {url} 的外部关系，实际: {targets}"
    assert "点这里" in _document_text(out), "链接文字应保留"


def test_hyperlink_text_and_surrounding_text(tmp_path):
    """链接前后的普通文本也要保留。"""
    out = _export("开头 [链接](https://a.example.com) 结尾", tmp_path)
    text = _document_text(out)
    assert "开头" in text and "结尾" in text and "链接" in text


def test_link_without_href_does_not_crash(tmp_path):
    """没有 href 的 <a> 不应导致导出失败。"""
    out = _export("<a>只有文字</a>", tmp_path, name="nohref.docx")
    assert "只有文字" in _document_text(out)


# ==================== 不回归：常规内容 ====================

def test_regular_content_still_works(tmp_path):
    md = (
        "# 标题\n\n"
        "普通段落，含 **加粗** 与 *斜体*。\n\n"
        "- 项目一\n- 项目二\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n"
    )
    out = _export(md, tmp_path, name="plain.docx")
    text = _document_text(out)
    assert "标题" in text and "普通段落" in text
    assert "项目一" in text
    assert len(Document(str(out)).tables) == 1
