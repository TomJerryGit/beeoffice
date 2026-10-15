# tests/fidelity/test_docx_nested_tables.py
"""
R4：**嵌套表格**（单元格里的表格）往返保真 + 保真度报告导出。

## 为什么这是"最后一块内容级丢失"
Markdown 表格语法**表达不了**"单元格里还有一张表"。实测（本文件夹具）：
往返后 `w:tbl` **2 → 1**、内层文字一并消失 —— 这是保真度阶段 A 收尾时
唯一剩下的**内容级**丢失。

## 做法（沿用既有 extras 模式）
合并单元格 / 书签 / 分节都已经走"导入时记语义 → 导出时重建"这条路
（Markdown 表达不了的东西放 `extras`）。嵌套表格同样处理：
`_read_nested_tables()` 记录**结构 + 文字**，`apply_nested_tables()` 在单元格里
`add_table` 重建。

## 保真范围（如实说明，不夸大）
只保 **结构（几行几列）+ 每格文字**；内层的**字符格式/边框/底色不保**
（那需要把整段 OOXML 搬过去，与既有取舍不一致）。只处理**一层**嵌套。
另外：内层内容**在编辑器里看不到**（Markdown 无法表达），
与合并单元格/书签的取舍相同。
"""
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests" / "fidelity"))

from _docx_features import build_feature_docx                     # noqa: E402

from beeoffice.plugins.bee_export import (                         # noqa: E402
    apply_nested_tables, export_to_docx, import_from_docx,
    read_docx_extras, _read_nested_tables,
)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def count_tables(path):
    """文档里 `w:tbl` 的总数（含嵌套）。"""
    import zipfile

    from lxml import etree

    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    return len(root.findall(f".//{W}tbl"))


def nested_texts(path):
    """取出所有**嵌套**表格（表格里的表格）各格文字。"""
    import zipfile

    from lxml import etree

    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    out = []
    for outer in root.findall(f".//{W}tbl"):
        for inner in outer.iter(f"{W}tbl"):
            if inner is outer:
                continue
            rows = []
            for tr in inner.findall(f"{W}tr"):
                rows.append([
                    "".join(t.text or "" for t in tc.iter(f"{W}t")).strip()
                    for tc in tr.findall(f"{W}tc")
                ])
            out.append(rows)
    return out


def roundtrip(src, work, name="out.docx"):
    """走产品真实往返：读 extras → import → export（带 extras）。"""
    out = work / name
    extras = read_docx_extras(src)
    export_to_docx(
        import_from_docx(str(src)), str(out),
        base_dir=str(work),
        header_text=extras.get("header_text"),
        footer_text=extras.get("footer_text"),
        table_merges=extras.get("table_merges"),
        nested_tables=extras.get("nested_tables"),
        sections=extras.get("sections"),
        bookmarks=extras.get("bookmarks"),
        paragraph_styles=extras.get("paragraph_styles"),
        style_definitions=extras.get("style_definitions"),
    )
    return out


# ==================== 1. 读取侧 ====================

def test_read_detects_nested_table(tmp_path):
    src = tmp_path / "nested.docx"
    build_feature_docx(src, only=["nested_table"])

    records = _read_nested_tables(__import__("docx").Document(str(src)))

    assert records, "没有读出嵌套表格"
    record = records[0]
    assert record["table"] == 0
    assert record["row"] == 0 and record["col"] == 0
    flat = [cell for row in record["cells"] for cell in row]
    assert "内层表格内容" in flat, record


def test_read_ignores_plain_tables(tmp_path):
    """普通表格不该被误记成嵌套。"""
    from docx import Document

    src = tmp_path / "plain.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "甲"
    table.cell(1, 1).text = "乙"
    doc.save(str(src))

    assert _read_nested_tables(Document(str(src))) == []


def test_read_ignores_nested_tables_without_text(tmp_path):
    """内层没有任何文字的嵌套表：没有可还原的内容，不记录。"""
    import zipfile

    from lxml import etree

    from docx import Document

    src = tmp_path / "empty_nested.docx"
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "外层"
    doc.save(str(src))

    # 手工塞一个空的嵌套 tbl（只有 tblPr，没有 tr）
    temp = tmp_path / "temp.docx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(temp, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                root = etree.fromstring(data)
                outer_tbl = root.find(f".//{W}tbl")
                tc = outer_tbl.find(f"{W}tr/{W}tc")
                empty = etree.SubElement(tc, f"{W}tbl")
                etree.SubElement(empty, f"{W}tblPr")
                data = etree.tostring(root, xml_declaration=True,
                                      encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    temp.replace(src)

    assert _read_nested_tables(Document(str(src))) == []


# ==================== 2. 往返（核心验收）====================

def test_nested_table_survives_roundtrip(tmp_path):
    """
    **核心验收**：嵌套表格的结构与文字都要回来。

    这是"最后一块内容级丢失"的修复证明 —— 修复前 `w:tbl` 是 2→1。
    """
    src = tmp_path / "nested.docx"
    build_feature_docx(src, only=["nested_table"])
    assert count_tables(src) == 2, "前提：夹具应有外层 + 内层两张表"

    out = roundtrip(src, tmp_path)

    assert count_tables(out) == 2, (
        f"嵌套表格没还原：{count_tables(out)} 张表（应为 2）"
    )
    inner = nested_texts(out)
    assert inner, "输出里没有嵌套表"
    flat = [cell for rows in inner for row in rows for cell in row]
    assert "内层表格内容" in flat, f"内层文字丢了：{inner}"
    # 外层文字也要在
    from docx import Document

    all_text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "外层单元格" in "".join(
        cell for outer in [
            ["".join(c.text for c in row.cells) for row in t.rows]
            for t in Document(str(out)).tables
        ] for cell in outer
    ) or "外层单元格" in all_text, "外层文字丢了"


def test_nested_preserves_shape(tmp_path):
    """结构要保：内层几行几列应一致。"""
    import zipfile

    from lxml import etree
    from docx import Document

    src = tmp_path / "shape.docx"
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "外层"
    inner = outer.cell(0, 0).add_table(rows=2, cols=3)
    for r in range(2):
        for c in range(3):
            inner.cell(r, c).text = f"内{r}{c}"
    doc.save(str(src))

    out = roundtrip(src, tmp_path, "shape_out.docx")

    got = nested_texts(out)
    assert got, "没有还原出嵌套表"
    rows = got[0]
    assert len(rows) == 2, f"行数不对：{rows}"
    assert all(len(r) == 3 for r in rows), f"列数不对：{rows}"
    assert rows[0][2] == "内02", rows


def test_apply_is_idempotent(tmp_path):
    """重复还原不应叠加出多余的表（导出被调用两次也不会翻倍）。"""
    from docx import Document

    src = tmp_path / "nested.docx"
    build_feature_docx(src, only=["nested_table"])
    records = _read_nested_tables(Document(str(src)))

    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "外层"
    first = apply_nested_tables(doc, records)
    second = apply_nested_tables(doc, records)

    assert first == 1 and second == 0, (first, second)
    assert len(doc.tables) == 1, "外层表数量变了"
    assert len(doc.tables[0].cell(0, 0).tables) == 1, "嵌套表被叠加了"


# ==================== 3. 越界与容错 ====================

def test_apply_skips_out_of_range_records():
    """坐标越界的记录要跳过，不能抛异常毁掉整个导出。"""
    from docx import Document

    doc = Document()
    doc.add_table(rows=1, cols=1)

    applied = apply_nested_tables(doc, [
        {"table": 99, "row": 0, "col": 0, "cells": [["x"]]},
        {"table": 0, "row": 5, "col": 0, "cells": [["x"]]},
        {"table": 0, "row": 0, "col": 5, "cells": [["x"]]},
        {"table": 0, "row": 0, "col": 0, "cells": []},
    ])

    assert applied == 0


def test_apply_handles_ragged_rows(tmp_path):
    """内层每行列数不一致时按最大列补齐，不崩。"""
    from docx import Document

    doc = Document()
    doc.add_table(rows=1, cols=1)

    applied = apply_nested_tables(doc, [{
        "table": 0, "row": 0, "col": 0,
        "cells": [["a", "b", "c"], ["d"]],
    }])

    assert applied == 1
    inner = doc.tables[0].cell(0, 0).tables[0]
    assert len(inner.rows) == 2
    assert len(inner.columns) == 3
    assert inner.cell(1, 0).text == "d"


def test_export_without_nested_tables_is_unchanged(tmp_path):
    """不传 nested_tables 时行为与从前一致（向后兼容）。"""
    src = tmp_path / "plain.docx"
    from docx import Document

    doc = Document()
    doc.add_paragraph("只有正文")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "甲"
    doc.save(str(src))

    out = tmp_path / "plain_out.docx"
    export_to_docx(import_from_docx(str(src)), str(out))
    assert count_tables(out) == 1


# ==================== 4. 保真度报告可导出 ====================

def test_fidelity_report_can_be_written_to_file(tmp_path):
    """
    保真度报告要能**导出成文件** —— 竞品普遍不做这件事，
    而它成本极低：用户可以把"这份文件哪些内容保不住"存档或发给同事。
    """
    from beeoffice.plugins.fidelity_report import (
        FidelityReport, LossItem, write_report_file,
    )

    report = FidelityReport(path="有脚注.docx", items=[
        LossItem("footnote", "脚注", 2, "脚注内容与标记都不会保留"),
        LossItem("toc_field", "目录域", 1, "文字保留但不能更新", "medium"),
    ])
    target = tmp_path / "保真度报告.md"

    ok = write_report_file(report, target)

    assert ok is True
    content = target.read_text(encoding="utf-8")
    assert "脚注" in content and "2 处" in content
    assert "目录域" in content
    assert "有脚注.docx" in content


def test_write_report_file_reports_failure(tmp_path, monkeypatch):
    """写不出去要返回 False（调用方据此提示），不能静默成功。"""
    from beeoffice.plugins import fidelity_report as module

    def boom(*args, **kwargs):
        raise OSError("模拟：目录不可写")

    monkeypatch.setattr("pathlib.Path.write_text", boom)

    ok = module.write_report_file(module.FidelityReport(path="x.docx"), tmp_path / "r.md")

    assert ok is False


def test_write_report_file_creates_parent_dirs(tmp_path):
    from beeoffice.plugins.fidelity_report import (
        FidelityReport, write_report_file,
    )

    target = tmp_path / "深" / "层" / "报告.md"
    assert write_report_file(FidelityReport(path="a.docx"), target) is True
    assert target.exists()
