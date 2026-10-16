# tests/fidelity/test_docx_hyperlink_import.py
"""
保真度阶段 A：DOCX **超链接**导入修复。

缺陷（实测）
------------
`import_from_docx()` 把普通段落交给 `_para_to_md()`，而它遍历的是
`para.runs` —— python-docx 的 `Paragraph.runs` 只返回 `w:p` 的**直接子**
`w:r`，超链接的文字位于 `w:hyperlink/w:r`，因此：

  - 一个**只含超链接**的段落转出来是**空字符串**；
  - 链接文字与链接一起**静默消失**。

保真度基线里的两条症状都出自这里：`hyperlinks 1 → 0`、
以及"文字少了 12 字"（少的正是链接文字）。

修复：按**文档顺序**遍历 `w:p` 的子元素，遇到 `w:hyperlink` 时取其文字与
目标（外部 `r:id` → URL；内部 `w:anchor` → `#书签`），写成 `[文字](url)`。
导出路径既有的 `_add_hyperlink()` 会把它还原成真正的 `w:hyperlink`。
"""
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document                     # noqa: E402
from docx.oxml import OxmlElement             # noqa: E402
from docx.oxml.ns import qn                   # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT   # noqa: E402

from beeoffice.plugins.bee_export import (    # noqa: E402
    _para_to_md,
    import_from_docx,
)


def _add_link(para, text, url=None, anchor=None):
    """往段落里追加一个 w:hyperlink（python-docx 没有公开 API）。"""
    if url is not None:
        r_id = para.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
    else:
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), anchor or "书签")

    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    para._p.append(link)


# ==================== 1. _para_to_md：段落级 ====================

def test_paragraph_with_only_hyperlink_is_not_dropped(tmp_path):
    """**核心缺陷**：整段只有超链接时不得变成空字符串。"""
    doc = Document()
    para = doc.add_paragraph()
    _add_link(para, "项目主页", "https://example.com")

    md = _para_to_md(para)

    assert md == "[项目主页](https://example.com)", (
        f"超链接段落被丢掉了或格式不对: {md!r}"
    )


def test_text_and_link_order_is_preserved(tmp_path):
    """混排时的**文档顺序**必须保持。"""
    doc = Document()
    para = doc.add_paragraph("前缀")
    _add_link(para, "链接", "https://a.example")
    para.add_run("后缀")

    md = _para_to_md(para)

    assert md == "前缀[链接](https://a.example)后缀", md


def test_link_between_two_links(tmp_path):
    doc = Document()
    para = doc.add_paragraph()
    _add_link(para, "甲", "https://a.example")
    para.add_run("和")
    _add_link(para, "乙", "https://b.example")

    assert _para_to_md(para) == (
        "[甲](https://a.example)和[乙](https://b.example)"
    )


def test_bold_inside_link(tmp_path):
    doc = Document()
    para = doc.add_paragraph()
    _add_link(para, "加粗链接", "https://example.com")
    # 把链接里的 run 设为粗体
    run_el = para._p.find(qn("w:hyperlink")).find(qn("w:r"))
    rPr = OxmlElement("w:rPr")
    bold = OxmlElement("w:b")
    rPr.append(bold)
    run_el.insert(0, rPr)

    md = _para_to_md(para)

    assert md == "[**加粗链接**](https://example.com)", md


def test_internal_bookmark_uses_anchor(tmp_path):
    """内部跳转（w:anchor）应写成 `#书签名`。"""
    doc = Document()
    para = doc.add_paragraph()
    _add_link(para, "见第一章", anchor="ch1")

    assert _para_to_md(para) == "[见第一章](#ch1)"


def test_plain_paragraph_unchanged():
    """没有超链接的段落行为不变。"""
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("普通")
    bold_run = para.add_run("加粗")
    bold_run.bold = True

    assert _para_to_md(para) == "普通**加粗**"


def test_broken_relationship_does_not_crash():
    """关系解析失败时应退化为纯文字，而不是抛异常。"""
    doc = Document()
    para = doc.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), "rIdDoesNotExist")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "断链"
    run.append(t)
    link.append(run)
    para._p.append(link)

    md = _para_to_md(para)

    assert "断链" in md, f"文字不应丢失: {md!r}"
    assert "](" not in md, "解析不到目标时不应写出空的 Markdown 链接"


# ==================== 2. import_from_docx：文件级 ====================

def _docx_with_link(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("正文")
    para = doc.add_paragraph("详见 ")
    _add_link(para, "官方文档", "https://docs.example.com")
    doc.save(str(path))
    return path


def test_import_from_docx_keeps_hyperlink(tmp_path):
    src = _docx_with_link(tmp_path / "with_link.docx")

    md = import_from_docx(str(src))

    assert "[官方文档](https://docs.example.com)" in md, md
    assert "详见" in md, "链接前后的普通文字也要保留"


# ==================== 3. 端到端：往返后超链接仍在 ====================

def test_hyperlink_survives_docx_roundtrip(tmp_path):
    """
    **用户可见的核心断言**：源文件有 1 个超链接 → 导出的 DOCX 也有 1 个。

    对应保真度基线里的 `hyperlinks 1 → 0`（原为丢 1 个）。
    """
    from tools.ooxml_probe import probe

    src = _docx_with_link(tmp_path / "src.docx")
    assert probe(src)["features"]["hyperlinks"] == 1, "前置条件"

    out = tmp_path / "out.docx"
    from beeoffice.plugins.bee_export import export_to_docx

    assert export_to_docx(
        import_from_docx(str(src)), str(out), base_dir=str(tmp_path)
    ) is True

    assert probe(out)["features"]["hyperlinks"] == 1, (
        "往返后超链接丢了"
    )


def test_hyperlink_target_survives_docx_roundtrip(tmp_path):
    """不只是数量 —— 目标 URL 也不能变。"""
    src = _docx_with_link(tmp_path / "src.docx")
    out = tmp_path / "out.docx"

    from beeoffice.plugins.bee_export import export_to_docx

    export_to_docx(
        import_from_docx(str(src)), str(out), base_dir=str(tmp_path)
    )

    doc = Document(str(out))
    targets = []
    for para in doc.paragraphs:
        for link in para._p.findall(qn("w:hyperlink")):
            r_id = link.get(qn("r:id"))
            if r_id:
                targets.append(para.part.rels[r_id].target_ref)

    assert "https://docs.example.com" in targets, targets
